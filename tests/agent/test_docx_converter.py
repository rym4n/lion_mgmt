# coding=utf-8
import asyncio
import base64
from copy import deepcopy
import io
from pathlib import Path
import threading
import time

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from docx.text.run import Run
from lxml import etree, html as lxml_html
from PIL import Image

from services.agent import docx_converter
from services.agent.docx_converter import (
    DocxConversionError,
    decode_docx_base64,
    docx_to_html,
    html_to_docx,
    replace_paragraph_text_matches,
)


@pytest.fixture
def image_data_uri() -> str:
    image = io.BytesIO()
    Image.new("RGB", (3, 2), (220, 20, 60)).save(image, format="PNG")
    return "data:image/png;base64," + base64.b64encode(image.getvalue()).decode("ascii")


@pytest.fixture
def rich_html(image_data_uri: str) -> str:
    return f"""
    <h1 style="text-align: center; line-height: 1.5">
      <span style="font-family: Noto Sans CJK SC; font-size: 20pt; color: #123456"><strong>运行报告</strong></span>
    </h1>
    <p style="text-align: justify; line-height: 1.8">
      正文 <strong>粗体</strong><em>斜体</em><u>下划线</u><s>删除线</s>
      <sup>上标</sup><sub>下标</sub><mark style="background-color: #00FF00">高亮</mark>
      <a href="https://example.com/report?id=1">详情链接</a><br>
      <img src="{image_data_uri}" alt="告警趋势" style="width: 24pt; height: 16pt">
    </p>
    <ul><li>节点一</li><li>节点二<ul><li>子节点</li></ul></li></ul>
    <ol><li>第一步</li><li>第二步</li></ol>
    <table><tbody>
      <tr><th>指标</th><th>结果</th></tr>
      <tr><td>可用率</td><td><strong>99.99%</strong></td></tr>
    </tbody></table>
    """


@pytest.fixture
def rich_docx_bytes(image_data_uri: str) -> bytes:
    document = Document()
    document.add_heading("既有 Word 报告", level=2)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.line_spacing = 1.25
    formatted = paragraph.add_run("组合格式")
    formatted.bold = True
    formatted.italic = True
    formatted.underline = True
    formatted.font.strike = True
    formatted.font.name = "Microsoft YaHei"
    formatted.font.size = Pt(15)
    formatted.font.color.rgb = RGBColor(0x65, 0x43, 0x21)
    formatted.font.highlight_color = WD_COLOR_INDEX.YELLOW
    superscript = paragraph.add_run("2")
    superscript.font.superscript = True

    relationship_id = paragraph.part.relate_to(
        "https://example.com/original", RT.HYPERLINK, is_external=True
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink_run_element = OxmlElement("w:r")
    hyperlink.append(hyperlink_run_element)
    hyperlink_run = Run(hyperlink_run_element, paragraph)
    hyperlink_run.text = "原始链接"
    paragraph._p.append(hyperlink)

    document.add_paragraph("原始项目", style="List Bullet")
    document.add_paragraph("原始步骤", style="List Number")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"

    raw_image = base64.b64decode(image_data_uri.split(",", 1)[1])
    image_run = document.add_paragraph().add_run()
    shape = image_run.add_picture(io.BytesIO(raw_image), width=Inches(0.5))
    shape._inline.docPr.set("descr", "原始图片")

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _paragraph_with_text(document: Document, text: str):
    return next(
        paragraph for paragraph in document.paragraphs if text in paragraph.text
    )


def _run_with_text(paragraph, text: str):
    return next(run for run in paragraph.runs if text in run.text)


def _document_bytes(document: Document) -> bytes:
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _png_bytes(color: tuple[int, int, int]) -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (3, 2), color).save(output, format="PNG")
    return output.getvalue()


def _assert_valid_png_data_uri(source: str) -> bytes:
    prefix = "data:image/png;base64,"
    assert source.startswith(prefix)
    image_bytes = base64.b64decode(source[len(prefix) :], validate=True)
    with Image.open(io.BytesIO(image_bytes)) as image:
        image.verify()
    return image_bytes


def _picture_run(document: Document, image_bytes: bytes):
    run = document.add_paragraph().add_run()
    run.add_picture(io.BytesIO(image_bytes), width=Inches(0.5))
    drawing = run._r.find(qn("w:drawing"))
    assert drawing is not None
    blip = next(
        drawing.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}blip")
    )
    relationship_id = blip.get(qn("r:embed"))
    assert relationship_id
    return run, drawing, relationship_id


def _vml_picture(relationship_id: str) -> etree._Element:
    vml_namespace = "urn:schemas-microsoft-com:vml"
    picture = OxmlElement("w:pict")
    shape = etree.SubElement(
        picture, f"{{{vml_namespace}}}shape", nsmap={"v": vml_namespace}
    )
    image_data = etree.SubElement(shape, f"{{{vml_namespace}}}imagedata")
    image_data.set(qn("r:id"), relationship_id)
    return picture


def _part_images(part, element) -> list[etree._Element]:
    root = etree.Element("div")
    docx_converter._append_docx_part_images(
        root, part, element, docx_converter._ImageBudget()
    )
    return root.xpath("./img")


def _round_trip_html(html_content: str):
    return lxml_html.fragment_fromstring(
        docx_to_html(html_to_docx(html_content)), create_parent="div"
    )


def _remove_table_look(table) -> None:
    table_look = table._tbl.tblPr.find(qn("w:tblLook"))
    if table_look is not None:
        table._tbl.tblPr.remove(table_look)


def _add_numbered_paragraph(
    document: Document,
    text: str,
    *,
    number_format: str,
    start: int,
):
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    level_start = OxmlElement("w:start")
    level_start.set(qn("w:val"), "1")
    level_format = OxmlElement("w:numFmt")
    level_format.set(qn("w:val"), number_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "%1.")
    level.extend([level_start, level_format, level_text])
    abstract.append(level)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    number.extend([abstract_reference, override])
    numbering.append(number)

    paragraph = document.add_paragraph(text)
    num_properties = OxmlElement("w:numPr")
    level_reference = OxmlElement("w:ilvl")
    level_reference.set(qn("w:val"), "0")
    number_reference = OxmlElement("w:numId")
    number_reference.set(qn("w:val"), str(num_id))
    num_properties.extend([level_reference, number_reference])
    paragraph._p.get_or_add_pPr().append(num_properties)
    return paragraph, num_id


def test_docx_import_drops_only_invalid_style_declarations(caplog):
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing = Pt(300)
    run = paragraph.add_run("越界样式仍保留正文")
    run.font.name = "微软雅黑 (正文)"
    run.font.size = Pt(100)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

    with caplog.at_level("WARNING", logger=docx_converter.__name__):
        converted_html = docx_to_html(_document_bytes(document))

    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    assert root.xpath(".//p[contains(., '越界样式仍保留正文')]")
    assert "color: #123456" in converted_html
    assert "font-size" not in converted_html
    assert "font-family" not in converted_html
    assert "line-height" not in converted_html
    warnings = "\n".join(record.getMessage() for record in caplog.records)
    assert "font-size='100pt'" in warnings
    assert "font-family='微软雅黑 (正文)'" in warnings
    assert "line-height='300pt'" in warnings
    round_tripped = _round_trip_html(converted_html)
    assert round_tripped.xpath(".//p[contains(., '越界样式仍保留正文')]")
    assert "color: #123456" in lxml_html.tostring(round_tripped, encoding="unicode")


def test_numbered_heading_style_takes_precedence_over_list_semantics():
    document = Document()
    heading = document.add_heading("自动编号标题", level=2)
    list_numbering = document.styles["List Number"]._element.pPr.numPr
    heading._p.get_or_add_pPr().append(deepcopy(list_numbering))

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")

    assert root.xpath("./h2[normalize-space()='自动编号标题']")
    assert not root.xpath(".//li[contains(., '自动编号标题')]")
    assert _round_trip_html(converted_html).xpath(
        "./h2[normalize-space()='自动编号标题']"
    )


def test_docx_table_header_and_vertical_merge_are_preserved():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    _remove_table_look(table)
    for row_index in range(3):
        table.cell(row_index, 0).text = f"左{row_index}"
        table.cell(row_index, 1).text = f"右{row_index}"
    table.cell(0, 0).merge(table.cell(2, 0)).text = "纵向合并"
    table_header = OxmlElement("w:tblHeader")
    table._tbl.tr_lst[0].get_or_add_trPr().append(table_header)

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")

    assert len(root.xpath(".//table//th")) == 2
    assert len(root.xpath(".//table//td")) == 2
    assert len(root.xpath(".//table//*[self::th or self::td]")) == 4
    assert root.xpath(".//th[@rowspan='3'][normalize-space()='纵向合并']")
    assert converted_html.count("纵向合并") == 1
    round_tripped = _round_trip_html(converted_html)
    assert len(round_tripped.xpath(".//table//*[self::th or self::td]")) == 4
    assert round_tripped.xpath(".//th[@rowspan='3'][normalize-space()='纵向合并']")


def test_docx_table_look_first_row_emits_headers():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table_look = table._tbl.tblPr.find(qn("w:tblLook"))
    if table_look is None:
        table_look = OxmlElement("w:tblLook")
        table._tbl.tblPr.append(table_look)
    table_look.set(qn("w:firstRow"), "1")
    table.cell(0, 0).text = "表头一"
    table.cell(0, 1).text = "表头二"

    root = lxml_html.fragment_fromstring(
        docx_to_html(_document_bytes(document)), create_parent="div"
    )
    assert len(root.xpath(".//table//th")) == 2
    assert (
        len(_round_trip_html(docx_to_html(_document_bytes(document))).xpath(".//th"))
        == 2
    )


def test_docx_table_without_header_marker_keeps_first_row_as_data_cells():
    document = Document()
    table = document.add_table(rows=2, cols=2)
    _remove_table_look(table)
    table.cell(0, 0).text = "普通一"
    table.cell(0, 1).text = "普通二"

    root = lxml_html.fragment_fromstring(
        docx_to_html(_document_bytes(document)), create_parent="div"
    )
    assert not root.xpath(".//table//th")
    assert len(root.xpath(".//table//td")) == 4
    round_tripped = _round_trip_html(docx_to_html(_document_bytes(document)))
    assert not round_tripped.xpath(".//table//th")
    assert len(round_tripped.xpath(".//table//td")) == 4


def test_docx_paragraph_indentation_round_trip_preserves_word_lengths():
    document = Document()
    paragraph = document.add_paragraph("首行缩进正文")
    paragraph.paragraph_format.left_indent = Pt(18)
    paragraph.paragraph_format.first_line_indent = Pt(24)

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    imported = root.xpath("./p[normalize-space()='首行缩进正文']")[0]
    assert docx_converter._style_map(imported) == {
        "margin-left": "18pt",
        "text-indent": "24pt",
    }

    reexported = Document(io.BytesIO(html_to_docx(converted_html)))
    paragraph = _paragraph_with_text(reexported, "首行缩进正文")
    assert paragraph.paragraph_format.left_indent.pt == pytest.approx(18)
    assert paragraph.paragraph_format.first_line_indent.pt == pytest.approx(24)


def test_docx_paragraph_spacing_round_trip_preserves_word_lengths():
    document = Document()
    paragraph = document.add_paragraph("段落间距正文")
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(9)

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    imported = root.xpath("./p[normalize-space()='段落间距正文']")[0]
    assert docx_converter._style_map(imported) == {
        "margin-top": "6pt",
        "margin-bottom": "9pt",
    }

    reexported = Document(io.BytesIO(html_to_docx(converted_html)))
    paragraph = _paragraph_with_text(reexported, "段落间距正文")
    assert paragraph.paragraph_format.space_before.pt == pytest.approx(6)
    assert paragraph.paragraph_format.space_after.pt == pytest.approx(9)


def test_docx_table_visual_styles_and_widths_round_trip_through_real_package():
    document = Document()
    table = document.add_table(rows=1, cols=2)
    _remove_table_look(table)
    table.cell(0, 0).text = "保真单元格"
    table.cell(0, 1).text = "普通单元格"
    grid_columns = table._tbl.tblGrid.gridCol_lst
    grid_columns[0].set(qn("w:w"), "1800")
    grid_columns[1].set(qn("w:w"), "3600")

    styled_cell = table.cell(0, 0)
    styled_cell.width = Pt(90)
    cell_properties = styled_cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "ABCDEF")
    cell_properties.append(shading)
    borders = OxmlElement("w:tcBorders")
    for edge_name in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{edge_name}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "8")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), "123456")
        borders.append(edge)
    cell_properties.append(borders)

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    columns = root.xpath(".//table/colgroup/col")
    assert [docx_converter._style_map(column)["width"] for column in columns] == [
        "90pt",
        "180pt",
    ]
    imported_cell = root.xpath(".//table//td[normalize-space()='保真单元格']")[0]
    assert docx_converter._style_map(imported_cell) == {
        "background-color": "#ABCDEF",
        "width": "90pt",
        "border": "1pt solid #123456",
    }

    reexported = Document(io.BytesIO(html_to_docx(converted_html)))
    reexported_table = reexported.tables[0]
    assert [
        int(column.get(qn("w:w")))
        for column in reexported_table._tbl.tblGrid.gridCol_lst
    ] == [1800, 3600]
    reexported_properties = reexported_table.cell(0, 0)._tc.tcPr
    cell_width = reexported_properties.find(qn("w:tcW"))
    assert cell_width.get(qn("w:type")) == "dxa"
    assert int(cell_width.get(qn("w:w"))) == 1800
    assert reexported_properties.find(qn("w:shd")).get(qn("w:fill")) == "ABCDEF"
    reexported_borders = reexported_properties.find(qn("w:tcBorders"))
    assert {
        edge.tag.rsplit("}", 1)[-1]: (
            edge.get(qn("w:val")),
            edge.get(qn("w:sz")),
            edge.get(qn("w:color")),
        )
        for edge in reexported_borders
    } == {
        "top": ("single", "8", "123456"),
        "left": ("single", "8", "123456"),
        "bottom": ("single", "8", "123456"),
        "right": ("single", "8", "123456"),
    }
    run_properties = reexported_table.cell(0, 0).paragraphs[0].runs[0]._r.rPr
    assert run_properties is None or run_properties.find(qn("w:shd")) is None


def test_docx_image_dimensions_emit_prosemirror_attributes_and_legacy_style():
    image = io.BytesIO()
    Image.new("RGB", (8, 4), (220, 20, 60)).save(image, format="PNG")
    document = Document()
    document.add_paragraph().add_run().add_picture(
        io.BytesIO(image.getvalue()), width=Inches(2), height=Inches(1)
    )

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    imported_image = root.xpath(".//img")[0]

    assert imported_image.get("width") == "192"
    assert imported_image.get("height") == "96"
    assert imported_image.get("style") == "width: 144pt; height: 72pt"
    round_tripped_image = _round_trip_html(converted_html).xpath(".//img")[0]
    assert round_tripped_image.get("width") == "192"
    assert round_tripped_image.get("height") == "96"


def test_docx_import_extracts_legacy_vml_image_from_real_package():
    image_bytes = _png_bytes((220, 20, 60))
    document = Document()
    run, drawing, relationship_id = _picture_run(document, image_bytes)
    run._r.remove(drawing)
    run._r.append(_vml_picture(relationship_id))

    docx_bytes = _document_bytes(document)
    reloaded = Document(io.BytesIO(docx_bytes))
    assert not tuple(
        reloaded.element.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
        )
    )
    assert (
        len(tuple(reloaded.element.iter("{urn:schemas-microsoft-com:vml}imagedata")))
        == 1
    )

    root = lxml_html.fragment_fromstring(docx_to_html(docx_bytes), create_parent="div")
    images = root.xpath(".//img")
    assert len(images) == 1
    assert _assert_valid_png_data_uri(images[0].get("src")) == image_bytes


def test_docx_import_does_not_dedupe_drawingml_and_vml_nodes_sharing_rid():
    image_bytes = _png_bytes((12, 120, 210))
    document = Document()
    run, _drawing, relationship_id = _picture_run(document, image_bytes)
    run._r.append(_vml_picture(relationship_id))

    root = lxml_html.fragment_fromstring(
        docx_to_html(_document_bytes(document)), create_parent="div"
    )
    images = root.xpath(".//img")
    assert len(images) == 2
    assert all(
        _assert_valid_png_data_uri(image.get("src")) == image_bytes for image in images
    )


def test_docx_import_extracts_textbox_image_from_alternate_content_once():
    image_bytes = _png_bytes((30, 180, 90))
    document = Document()
    run, drawing, relationship_id = _picture_run(document, image_bytes)
    paragraph = run._parent
    run._r.remove(drawing)
    paragraph._p.remove(run._r)

    vml_namespace = "urn:schemas-microsoft-com:vml"
    mc_namespace = "http://schemas.openxmlformats.org/markup-compatibility/2006"
    text_box_picture = OxmlElement("w:pict")
    shape = etree.SubElement(
        text_box_picture,
        f"{{{vml_namespace}}}shape",
        nsmap={"v": vml_namespace},
    )
    text_box = etree.SubElement(shape, f"{{{vml_namespace}}}textbox")
    text_box_content = OxmlElement("w:txbxContent")
    nested_paragraph = OxmlElement("w:p")
    nested_run = OxmlElement("w:r")
    nested_run.append(drawing)
    nested_paragraph.append(nested_run)
    text_box_content.append(nested_paragraph)
    text_box.append(text_box_content)

    alternate = etree.Element(
        f"{{{mc_namespace}}}AlternateContent", nsmap={"mc": mc_namespace}
    )
    choice = etree.SubElement(alternate, f"{{{mc_namespace}}}Choice")
    choice.set("Requires", "wps")
    choice_run = OxmlElement("w:r")
    choice_run.append(text_box_picture)
    choice.append(choice_run)
    fallback = etree.SubElement(alternate, f"{{{mc_namespace}}}Fallback")
    fallback_run = OxmlElement("w:r")
    fallback_run.append(_vml_picture(relationship_id))
    fallback.append(fallback_run)
    paragraph.add_run("图片前")
    paragraph._p.append(alternate)
    paragraph.add_run("图片后")

    docx_bytes = _document_bytes(document)
    reloaded = Document(io.BytesIO(docx_bytes))
    assert len(tuple(reloaded.element.iter(qn("w:txbxContent")))) == 1
    assert reloaded.paragraphs[0]._p[-2].tag == (
        "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
    )
    assert (
        len(
            tuple(
                reloaded.element.iter(
                    "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
                )
            )
        )
        == 1
    )

    root = lxml_html.fragment_fromstring(docx_to_html(docx_bytes), create_parent="div")
    images = root.xpath(".//img")
    assert len(images) == 1
    assert _assert_valid_png_data_uri(images[0].get("src")) == image_bytes
    paragraph_children = root.xpath("./p/*")
    assert [child.tag for child in paragraph_children] == ["span", "img", "span"]
    assert paragraph_children[0].text_content() == "图片前"
    assert paragraph_children[2].text_content() == "图片后"


def test_header_and_footer_parts_can_extract_images_without_body_integration():
    body_image = _png_bytes((220, 20, 60))
    header_image = _png_bytes((30, 90, 180))
    footer_image = _png_bytes((40, 170, 80))
    document = Document()
    document.add_paragraph().add_run().add_picture(io.BytesIO(body_image))
    section = document.sections[0]
    section.header.paragraphs[0].add_run().add_picture(io.BytesIO(header_image))
    section.footer.paragraphs[0].add_run().add_picture(io.BytesIO(footer_image))

    docx_bytes = _document_bytes(document)
    reloaded = Document(io.BytesIO(docx_bytes))
    body_root = lxml_html.fragment_fromstring(
        docx_to_html(docx_bytes), create_parent="div"
    )
    body_images = body_root.xpath(".//img")
    assert len(body_images) == 1
    assert _assert_valid_png_data_uri(body_images[0].get("src")) == body_image

    reloaded_section = reloaded.sections[0]
    header_images = _part_images(
        reloaded_section.header.part, reloaded_section.header._element
    )
    footer_images = _part_images(
        reloaded_section.footer.part, reloaded_section.footer._element
    )
    assert len(header_images) == 1
    assert len(footer_images) == 1
    assert _assert_valid_png_data_uri(header_images[0].get("src")) == header_image
    assert _assert_valid_png_data_uri(footer_images[0].get("src")) == footer_image


@pytest.mark.parametrize("size_pt", [1, 2000])
def test_docx_image_dimension_safety_boundaries_remain_importable(size_pt):
    image = io.BytesIO()
    Image.new("RGB", (1, 1), (220, 20, 60)).save(image, format="PNG")
    document = Document()
    document.add_paragraph().add_run().add_picture(
        io.BytesIO(image.getvalue()), width=Pt(size_pt), height=Pt(size_pt)
    )

    converted_html = docx_to_html(_document_bytes(document))
    imported_image = lxml_html.fragment_fromstring(
        converted_html, create_parent="div"
    ).xpath(".//img")[0]

    assert float(imported_image.get("width")) * 0.75 == pytest.approx(size_pt)
    assert float(imported_image.get("height")) * 0.75 == pytest.approx(size_pt)
    assert imported_image.get("style") == f"width: {size_pt}pt; height: {size_pt}pt"
    round_tripped_image = _round_trip_html(converted_html).xpath(".//img")[0]
    assert float(round_tripped_image.get("width")) * 0.75 == pytest.approx(size_pt)
    assert float(round_tripped_image.get("height")) * 0.75 == pytest.approx(size_pt)


def test_docx_ordered_list_preserves_type_and_start_override():
    document = Document()
    expected = [
        ("lowerLetter", "a", 3),
        ("upperLetter", "A", 4),
        ("lowerRoman", "i", 5),
        ("upperRoman", "I", 6),
    ]
    for index, (number_format, _html_type, start) in enumerate(expected):
        _add_numbered_paragraph(
            document,
            f"编号格式 {index}",
            number_format=number_format,
            start=start,
        )

    converted_html = docx_to_html(_document_bytes(document))
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    ordered_lists = root.xpath("./ol")

    assert [item.get("type") for item in ordered_lists] == [
        item[1] for item in expected
    ]
    assert [item.get("start") for item in ordered_lists] == [
        str(item[2]) for item in expected
    ]
    assert ["".join(item.itertext()).strip() for item in ordered_lists] == [
        f"编号格式 {index}" for index in range(4)
    ]
    round_tripped_lists = _round_trip_html(converted_html).xpath("./ol")
    assert [item.get("type") for item in round_tripped_lists] == [
        item[1] for item in expected
    ]
    assert [item.get("start") for item in round_tripped_lists] == [
        str(item[2]) for item in expected
    ]


def test_rich_html_docx_round_trip_preserves_editor_contract(rich_html: str):
    docx_bytes = html_to_docx(rich_html)
    document = Document(io.BytesIO(docx_bytes))

    heading = _paragraph_with_text(document, "运行报告")
    assert heading.style.name == "Heading 1"
    assert heading.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert heading.paragraph_format.line_spacing == pytest.approx(1.5)
    heading_run = _run_with_text(heading, "运行报告")
    assert heading_run.bold is True
    assert heading_run.font.name == "Noto Sans CJK SC"
    assert heading_run.font.size.pt == pytest.approx(20)
    assert str(heading_run.font.color.rgb) == "123456"

    paragraph = _paragraph_with_text(document, "正文")
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.line_spacing == pytest.approx(1.8)
    assert _run_with_text(paragraph, "粗体").bold is True
    assert _run_with_text(paragraph, "斜体").italic is True
    assert _run_with_text(paragraph, "下划线").underline is True
    assert _run_with_text(paragraph, "删除线").font.strike is True
    assert _run_with_text(paragraph, "上标").font.superscript is True
    assert _run_with_text(paragraph, "下标").font.subscript is True
    highlight_run = _run_with_text(paragraph, "高亮")
    shading = highlight_run._r.get_or_add_rPr().find(qn("w:shd"))
    assert shading is not None
    assert shading.get(qn("w:fill")) == "00FF00"

    external_links = [
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype == RT.HYPERLINK and relationship.is_external
    ]
    assert external_links == ["https://example.com/report?id=1"]
    assert len(document.inline_shapes) == 1
    assert document.inline_shapes[0]._inline.docPr.get("descr") == "告警趋势"
    assert [
        paragraph.style.name
        for paragraph in document.paragraphs
        if "节点" in paragraph.text
    ] == [
        "List Bullet",
        "List Bullet",
        "List Bullet 2",
    ]
    assert [
        paragraph.style.name
        for paragraph in document.paragraphs
        if "步" in paragraph.text
    ] == [
        "List Number",
        "List Number",
    ]
    assert len(document.tables) == 1
    assert document.tables[0].cell(1, 0).text == "可用率"
    assert document.tables[0].cell(1, 1).paragraphs[0].runs[0].bold is True

    converted_html = docx_to_html(docx_bytes)
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    assert root.xpath("./h1[contains(@style, 'text-align: center')]")
    assert root.xpath(".//strong[normalize-space()='运行报告']")
    assert root.xpath(".//em[normalize-space()='斜体']")
    assert root.xpath(".//u[normalize-space()='下划线']")
    assert root.xpath(".//s[normalize-space()='删除线']")
    assert root.xpath(".//sup[normalize-space()='上标']")
    assert root.xpath(".//sub[normalize-space()='下标']")
    assert root.xpath(".//mark[contains(@style, '#00FF00')]")
    assert root.xpath(".//a[@href='https://example.com/report?id=1']")
    assert root.xpath(".//ul/li[normalize-space()='节点一']")
    assert root.xpath(".//ul/li/ul/li[normalize-space()='子节点']")
    assert root.xpath(".//ol/li[normalize-space()='第一步']")
    assert len(root.xpath(".//table/tbody/tr")) == 2
    images = root.xpath(".//img")
    assert len(images) == 1
    assert images[0].get("src", "").startswith("data:image/png;base64,")
    assert images[0].get("alt") == "告警趋势"

    second_docx = html_to_docx(converted_html)
    second_document = Document(io.BytesIO(second_docx))
    assert _paragraph_with_text(second_document, "运行报告").style.name == "Heading 1"
    assert len(second_document.inline_shapes) == 1
    assert len(second_document.tables) == 1


def test_existing_rich_docx_import_and_reexport_preserves_supported_content(
    rich_docx_bytes: bytes,
):
    converted_html = docx_to_html(rich_docx_bytes)
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")

    assert root.xpath("./h2[normalize-space()='既有 Word 报告']")
    assert root.xpath(".//p[contains(@style, 'text-align: right')]")
    assert root.xpath(".//strong//em//u//s[normalize-space()='组合格式']")
    formatted_span = root.xpath(
        ".//span[contains(@style, 'font-family: Microsoft YaHei')]"
    )[0]
    assert "font-size: 15pt" in formatted_span.get("style")
    assert "color: #654321" in formatted_span.get("style")
    assert root.xpath(".//mark[contains(@style, '#FFFF00')]")
    assert root.xpath(".//sup[normalize-space()='2']")
    assert root.xpath(
        ".//a[@href='https://example.com/original'][normalize-space()='原始链接']"
    )
    assert root.xpath(".//ul/li[normalize-space()='原始项目']")
    assert root.xpath(".//ol/li[normalize-space()='原始步骤']")
    assert root.xpath(".//table/tbody/tr/th[normalize-space()='A']")
    assert root.xpath(".//img[@alt='原始图片']")

    reexported = Document(io.BytesIO(html_to_docx(converted_html)))
    assert _paragraph_with_text(reexported, "既有 Word 报告").style.name == "Heading 2"
    assert len(reexported.inline_shapes) == 1
    assert len(reexported.tables) == 1


@pytest.mark.parametrize(
    "content",
    ["", "not-base64", "YWJj\n", base64.b64encode(b"not a docx").decode("ascii")],
)
def test_docx_base64_validation_rejects_invalid_content(content: str):
    with pytest.raises(DocxConversionError):
        decode_docx_base64(content)


@pytest.mark.parametrize(
    "html_content",
    [
        "",
        "<script>alert(1)</script>",
        '<p onclick="alert(1)">x</p>',
        '<a href="javascript:alert(1)">x</a>',
        '<img src="file:///etc/passwd">',
        '<span style="color: url(javascript:alert(1))">x</span>',
        '<img src="data:image/png;base64,YWJj">',
    ],
)
def test_html_validation_rejects_unsafe_or_invalid_content(html_content: str):
    with pytest.raises(DocxConversionError):
        html_to_docx(html_content)


def test_html_size_limit_is_enforced(monkeypatch):
    monkeypatch.setattr(docx_converter, "MAX_HTML_BYTES", 16)
    with pytest.raises(DocxConversionError):
        html_to_docx("<p>12345678901234567890</p>")


def test_actual_tiptap_shapes_round_trip_without_remote_image_download():
    editor_html = """
    <p data-unknown="stripped">
      普通 <code>const answer = 42</code>
      <mark data-color="#22C55E" style="background-color: #22C55E; color: inherit">通过</mark>
      <a href="docs.example.com/report">无协议链接</a>
      <img src="//cdn.example.com/report.png" alt="外链图" style="width: 160px; height: 90px">
    </p>
    <pre><code class="language-python">if ready:\n    run()</code></pre>
    <hr>
    <ul data-type="taskList">
      <li data-type="taskItem" data-checked="true">
        <label><input type="checkbox" checked="checked"><span></span></label>
        <div><p>已完成</p></div>
      </li>
      <li data-type="taskItem" data-checked="false">
        <label><input type="checkbox"><span></span></label>
        <div><p>待处理</p></div>
      </li>
    </ul>
    <table style="min-width: 75px">
      <colgroup>
        <col style="min-width: 25px; width: 120px">
        <col style="min-width: 25px">
      </colgroup>
      <tbody><tr><th>名称</th><th>状态</th></tr><tr><td>任务</td><td>完成</td></tr></tbody>
    </table>
    """

    docx_bytes = html_to_docx(editor_html)
    document = Document(io.BytesIO(docx_bytes))
    ordinary = _paragraph_with_text(document, "普通")
    assert (
        _run_with_text(ordinary, "const answer = 42").style.name == "AIOE Inline Code"
    )
    assert (
        _run_with_text(ordinary, "通过")._r.rPr.find(qn("w:shd")).get(qn("w:fill"))
        == "22C55E"
    )
    assert _paragraph_with_text(document, "if ready:").style.name == "AIOE Code Block"
    assert _paragraph_with_text(document, "已完成").text.startswith("☑ ")
    assert _paragraph_with_text(document, "待处理").text.startswith("☐ ")
    assert document.tables[0].cell(1, 0).text == "任务"

    external_hyperlinks = [
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype == RT.HYPERLINK and relationship.is_external
    ]
    assert external_hyperlinks == ["https://docs.example.com/report"]
    external_images = [
        relationship.target_ref
        for relationship in document.part.rels.values()
        if relationship.reltype == RT.IMAGE and relationship.is_external
    ]
    assert external_images == ["https://cdn.example.com/report.png"]

    converted_html = docx_to_html(docx_bytes)
    root = lxml_html.fragment_fromstring(converted_html, create_parent="div")
    assert root.xpath(".//code[normalize-space()='const answer = 42']")
    assert root.xpath("./pre/code[contains(., 'if ready:')]")
    assert root.xpath("./hr")
    assert root.xpath(".//mark[contains(@style, 'background-color: #22C55E')]")
    assert root.xpath(".//a[@href='https://docs.example.com/report']")
    assert root.xpath(
        ".//img[@src='https://cdn.example.com/report.png'][@alt='外链图']"
    )
    assert root.xpath(
        ".//ul[@data-type='taskList']/li[@data-checked='true']//p[normalize-space()='已完成']"
    )
    assert root.xpath(
        ".//ul[@data-type='taskList']/li[@data-checked='false']//p[normalize-space()='待处理']"
    )
    assert root.xpath(".//table/tbody/tr/td[normalize-space()='任务']")
    assert "data-unknown" not in converted_html
    assert "data-color" not in converted_html


@pytest.mark.parametrize(
    "table_html",
    [
        "<table><tr><td rowspan='2'>越界</td></tr></table>",
        "<table><tr><td>A</td><td rowspan='2'>B</td></tr><tr><td colspan='2'>重叠</td></tr></table>",
        "<table><tr><td colspan='101'>过宽</td></tr></table>",
    ],
)
def test_invalid_table_spans_raise_contract_error(table_html: str):
    with pytest.raises(DocxConversionError):
        html_to_docx(table_html)


def test_table_row_limit_raises_contract_error(monkeypatch):
    monkeypatch.setattr(docx_converter, "MAX_TABLE_ROWS", 2)
    with pytest.raises(DocxConversionError, match="row limit"):
        html_to_docx(
            "<table><tr><td>1</td></tr><tr><td>2</td></tr><tr><td>3</td></tr></table>"
        )


def test_total_table_cell_limit_raises_contract_error(monkeypatch):
    monkeypatch.setattr(docx_converter, "MAX_TOTAL_TABLE_CELLS", 3)
    with pytest.raises(DocxConversionError, match="cell limit"):
        html_to_docx(
            "<table><tr><td>1</td><td>2</td></tr><tr><td>3</td><td>4</td></tr></table>"
        )


def test_docx_structure_limit_rejects_before_python_docx_load(monkeypatch):
    document = Document()
    document.add_paragraph("first")
    document.add_paragraph("second")
    output = io.BytesIO()
    document.save(output)

    monkeypatch.setattr(docx_converter, "MAX_DOCX_PARAGRAPHS", 1)
    monkeypatch.setattr(
        docx_converter,
        "Document",
        lambda *_args, **_kwargs: pytest.fail(
            "python-docx must not load structurally oversized input"
        ),
    )
    with pytest.raises(DocxConversionError, match="structure limit"):
        docx_converter.validate_docx_bytes(output.getvalue())


def test_docx_xml_node_limit_rejects_before_python_docx_load(monkeypatch):
    document = Document()
    document.add_paragraph("node budget")
    output = io.BytesIO()
    document.save(output)

    monkeypatch.setattr(docx_converter, "MAX_DOCX_XML_NODES", 3)
    monkeypatch.setattr(
        docx_converter,
        "Document",
        lambda *_args, **_kwargs: pytest.fail(
            "python-docx must not load node-heavy input"
        ),
    )
    with pytest.raises(DocxConversionError, match="structure limit"):
        docx_converter.validate_docx_bytes(output.getvalue())


def test_conversion_deadline_is_cooperatively_enforced(monkeypatch):
    monkeypatch.setattr(docx_converter, "MAX_CONVERSION_SECONDS", -1)
    with pytest.raises(DocxConversionError, match="timed out"):
        html_to_docx("<p>deadline</p>")


def test_medium_docx_conversion_stays_within_regression_budget():
    document = Document()
    for index in range(400):
        paragraph = document.add_paragraph()
        bold = paragraph.add_run(f"paragraph {index} ")
        bold.bold = True
        italic = paragraph.add_run("value")
        italic.italic = True
    output = io.BytesIO()
    document.save(output)

    started = time.perf_counter()
    converted = docx_to_html(output.getvalue())
    elapsed = time.perf_counter() - started

    assert "paragraph 399" in converted
    assert elapsed < 5.0


def test_report_export_warns_about_dropped_remote_image_without_logging_url(caplog):
    from services.agent import views

    source = "https://private.example.test/report-image.png?token=secret"
    node = views._ReportHTMLNode("img", {"src": source})
    paragraph = Document().add_paragraph()

    with caplog.at_level("WARNING"):
        views._append_data_image(paragraph, node, {})

    messages = [record.getMessage() for record in caplog.records]
    assert any(
        "category=unsupported-src" in message
        and "scheme=https" in message
        and f"src_length={len(source)}" in message
        for message in messages
    )
    assert all(source not in message for message in messages)
    assert not paragraph.runs


def test_report_export_warns_and_drops_malformed_image_url(caplog):
    from services.agent import views

    source = "http://["
    node = views._ReportHTMLNode("img", {"src": source})
    paragraph = Document().add_paragraph()

    with caplog.at_level("WARNING"):
        views._append_data_image(paragraph, node, {})

    messages = [record.getMessage() for record in caplog.records]
    assert any(
        "category=unsupported-src" in message
        and "scheme=invalid" in message
        and f"src_length={len(source)}" in message
        for message in messages
    )
    assert all(source not in message for message in messages)
    assert not paragraph.runs


def test_docx_conversion_runner_has_bounded_capacity(monkeypatch):
    from services.agent import views

    monkeypatch.setattr(views, "_DOCX_CONVERSION_WORKERS", asyncio.BoundedSemaphore(1))
    monkeypatch.setattr(views, "_DOCX_CONVERSION_CAPACITY", asyncio.BoundedSemaphore(1))
    started = threading.Event()
    release = threading.Event()

    def blocking_converter(value):
        started.set()
        release.wait(timeout=2)
        return value

    async def scenario():
        first = asyncio.create_task(
            views._run_docx_conversion(blocking_converter, "first")
        )
        try:
            while not started.is_set():
                await asyncio.sleep(0)
            with pytest.raises(views._DocxConversionBusyError):
                await views._run_docx_conversion(lambda value: value, "second")
        finally:
            release.set()
        assert await first == "first"

    asyncio.run(scenario())


def test_replace_paragraph_text_matches_preserves_surrounding_run_styles():
    document = Document()
    paragraph = document.add_paragraph()
    prefix = paragraph.add_run("前缀 ")
    prefix.bold = True
    placeholder_start = paragraph.add_run("{fetch.")
    placeholder_start.italic = True
    placeholder_end = paragraph.add_run("value}")
    placeholder_end.underline = True
    suffix = paragraph.add_run(" / {fetch.value} 后缀")
    suffix.font.color.rgb = RGBColor(0x12, 0x34, 0x56)

    pattern = docx_converter.re.compile(r"\{([a-z]+)\.([a-z]+)\}")
    assert (
        replace_paragraph_text_matches(paragraph, pattern, lambda match: "替换值")
        is True
    )
    assert paragraph.text == "前缀 替换值 / 替换值 后缀"
    assert prefix.text == "前缀 " and prefix.bold is True
    assert placeholder_start.text == "替换值" and placeholder_start.italic is True
    assert placeholder_end.text == "" and placeholder_end.underline is True
    assert suffix.text == " / 替换值 后缀"
    assert str(suffix.font.color.rgb) == "123456"


def test_replace_paragraph_text_matches_can_skip_unconfigured_split_placeholder():
    document = Document()
    paragraph = document.add_paragraph()
    first = paragraph.add_run("保留 {missing.")
    first.bold = True
    second = paragraph.add_run("value}")
    second.italic = True
    pattern = docx_converter.re.compile(r"\{([a-z]+)\.([a-z]+)\}")

    assert (
        replace_paragraph_text_matches(paragraph, pattern, lambda _match: None) is True
    )
    assert first.text == "保留 {missing." and first.bold is True
    assert second.text == "value}" and second.italic is True


def test_resolve_fetch_placeholders_keeps_run_formatting(monkeypatch):
    from services.agent import report_fetch, views

    document = Document()
    paragraph = document.add_paragraph()
    prefix = paragraph.add_run("前缀 ")
    prefix.bold = True
    start = paragraph.add_run("{fetch.")
    start.italic = True
    end = paragraph.add_run("value}")
    end.underline = True
    suffix = paragraph.add_run(" 后缀")
    suffix.font.color.rgb = RGBColor(0xAB, 0xCD, 0xEF)
    source = io.BytesIO()
    document.save(source)

    async def fake_fetch(*_args, **_kwargs):
        return '{"value":"42"}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_fetch)
    resolved = asyncio.run(
        views._resolve_fetch_placeholders(
            source.getvalue(), {"fetch": {}}, "tenant-1", "user-1"
        )
    )
    result = Document(io.BytesIO(resolved)).paragraphs[0]
    assert result.text == "前缀 42 后缀"
    assert result.runs[0].bold is True and result.runs[0].text == "前缀 "
    assert result.runs[1].italic is True and result.runs[1].text == "42"
    assert result.runs[2].underline is True and result.runs[2].text == ""
    assert str(result.runs[3].font.color.rgb) == "ABCDEF"


def test_docx_conversion_routes_have_fixed_contract_and_permissions():
    source = (
        Path(__file__).resolve().parents[2] / "services/agent/views.py"
    ).read_text()
    to_html = source[source.index('@agent_bp.post("/report/docx/to-html")') :]
    to_html = to_html[: to_html.index('@agent_bp.post("/report/docx/from-html")')]
    from_html = source[source.index('@agent_bp.post("/report/docx/from-html")') :]
    from_html = from_html[: from_html.index('@agent_bp.post("/report/save")')]

    assert '@requires_permission("agent:r")' in to_html
    assert 'await _run_docx_conversion(docx_base64_to_html, data["content"])' in to_html
    assert "ErrorCode.DOCX_CONVERSION_BUSY" in to_html
    assert 'return response_ok({"html": html})' in to_html
    assert '@requires_permission("agent:u")' in from_html
    assert 'await _run_docx_conversion(html_to_docx_base64, data["html"])' in from_html
    assert "ErrorCode.DOCX_CONVERSION_BUSY" in from_html
    assert 'return response_ok({"content": content})' in from_html
