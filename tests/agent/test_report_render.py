# coding=utf-8
import io

import pytest
from docx import Document

from core.error_code import ErrorCode
from libs.error_code.exception import CustomError
from services.agent import report_fetch, report_render


@pytest.fixture(autouse=True)
def _stub_llm_resolver(monkeypatch):
    async def fake_resolve_llm(tenant_id):
        return object()

    monkeypatch.setattr(report_render, "resolve_llm", fake_resolve_llm)


@pytest.mark.asyncio
async def test_resolve_placeholders_replaces_span_text(monkeypatch):
    async def fake_execute(action_cfg, tenant_id, operator_id):
        return '{"data": {"total": 128}}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    html = (
        "<p>本月告警 "
        '<span data-placeholder-id="alarm" data-result-path="total">'
        "{alarm.total}</span> 起</p>"
    )
    result = await report_render.resolve_placeholders_in_html(
        html, {"alarm": {"action": "count"}}, "t1", "u1"
    )
    assert "128" in result
    assert "{alarm.total}" not in result
    assert 'data-placeholder-id="alarm"' in result


@pytest.mark.asyncio
async def test_resolve_placeholders_returns_html_unchanged_without_tokens():
    html = "<p>没有占位符</p>"
    result = await report_render.resolve_placeholders_in_html(html, {}, "t1", "u1")
    assert "没有占位符" in result


@pytest.mark.asyncio
async def test_resolve_placeholders_keeps_attribute_braces(monkeypatch):
    async def fake_execute(action_cfg, tenant_id, operator_id):
        return '{"data": 5}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    html = (
        '<p data-config="{alarm.data}">'
        '<span data-placeholder-id="alarm" data-result-path="">{alarm}</span>'
        "</p>"
    )
    result = await report_render.resolve_placeholders_in_html(
        html, {"alarm": {"action": "count"}}, "t1", "u1"
    )
    assert 'data-config="{alarm.data}"' in result
    assert ">5<" in result


@pytest.mark.asyncio
async def test_resolve_placeholders_fails_on_missing_config():
    html = "<p>{unknown.result}</p>"
    with pytest.raises(CustomError) as excinfo:
        await report_render.resolve_placeholders_in_html(html, {}, "t1", "u1")
    assert excinfo.value.enum == ErrorCode.SERVER_ERROR


@pytest.mark.asyncio
async def test_resolve_placeholders_replaces_token_split_across_nodes(monkeypatch):
    async def fake_execute(action_cfg, tenant_id, operator_id):
        return '{"data": {"total": 42}}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    html = "<p>{alarm<strong></strong>.total}</p>"
    result = await report_render.resolve_placeholders_in_html(
        html, {"alarm": {"action": "count"}}, "t1", "u1"
    )
    assert "42" in result
    assert "{alarm" not in result


def test_collect_annotation_marks_groups_multiple_marks_per_annotation():
    html = (
        '<p><mark data-annotation-ids="a1">前半</mark>'
        '<strong><mark data-annotation-ids="a1">后半</mark></strong></p>'
    )
    root = report_render._parse_document(html)
    marks = report_render.collect_annotation_marks(root)
    assert list(marks.keys()) == ["a1"]
    assert len(marks["a1"]) == 2
    assert report_render.annotation_region_text(marks["a1"]) == "前半后半"


def test_collect_annotation_marks_reads_singular_attribute():
    html = '<p><mark data-annotation-id="a1">文字</mark></p>'
    root = report_render._parse_document(html)
    marks = report_render.collect_annotation_marks(root)
    assert list(marks.keys()) == ["a1"]


def test_collect_annotation_marks_reads_image_anchor_in_document_order():
    html = (
        '<p><img src="data:image/png;base64,image" data-annotation-ids="a1">'
        '<mark data-annotation-ids="a1">图片说明</mark></p>'
    )
    root = report_render._parse_document(html)
    marks = report_render.collect_annotation_marks(root)
    assert [element.tag for element in marks["a1"]] == ["img", "mark"]


def test_collect_annotation_marks_rejects_overlapping_annotation():
    html = '<p><mark data-annotation-ids="a1 a2">重叠</mark></p>'
    root = report_render._parse_document(html)
    with pytest.raises(CustomError) as excinfo:
        report_render.collect_annotation_marks(root)
    assert excinfo.value.enum == ErrorCode.REPORT_ANNOTATION_OVERLAP


def test_collect_annotation_marks_reuses_overlap_check_for_image_anchor():
    html = '<p><img src="x" data-annotation-ids="a1 a2"></p>'
    root = report_render._parse_document(html)
    with pytest.raises(CustomError) as excinfo:
        report_render.collect_annotation_marks(root)
    assert excinfo.value.enum == ErrorCode.REPORT_ANNOTATION_OVERLAP


def test_active_annotations_filters_by_status_and_comment():
    annotations = [
        {"id": "a1", "comment": "补充趋势", "status": "open"},
        {"id": "a2", "comment": "", "status": "open"},
        {"id": "a3", "comment": "已处理", "status": "applied"},
        {"id": "a4", "comment": "已完成", "status": "resolved"},
        {"id": "a5", "comment": "缺省状态"},
    ]
    active = report_render.active_annotations(annotations)
    assert [item["id"] for item in active] == ["a1"]


def test_annotation_mode_defaults_to_overwrite():
    assert report_render.annotation_mode({}) == report_render.ANNOTATION_MODE_OVERWRITE
    assert (
        report_render.annotation_mode({"mode": "append"})
        == report_render.ANNOTATION_MODE_APPEND
    )
    assert (
        report_render.annotation_mode({"mode": "invalid"})
        == report_render.ANNOTATION_MODE_OVERWRITE
    )


def test_annotation_region_text_spans_blocks():
    html = (
        '<p><mark data-annotation-ids="a1">第一段</mark></p>'
        '<p><mark data-annotation-ids="a1">第二段</mark></p>'
    )
    root = report_render._parse_document(html)
    marks = report_render.collect_annotation_marks(root)
    assert report_render.annotation_region_text(marks["a1"]) == "第一段第二段"


def test_sanitize_ai_html_keeps_paragraphs_and_lists():
    blocks = report_render.sanitize_ai_html(
        "<p>本月告警较上月<strong>下降 12%</strong>。</p>"
        "<ul><li>网络类 45 起</li><li>主机类 30 起</li></ul>"
    )
    assert [block.tag for block in blocks] == ["p", "ul"]
    serialized = report_render._serialize_blocks(blocks)
    assert "<strong>下降 12%</strong>" in serialized
    assert serialized.count("<li>") == 2


def test_sanitize_ai_html_strips_script_and_events():
    blocks = report_render.sanitize_ai_html(
        '<p onclick="steal()">正文</p><script>alert(1)</script>'
    )
    serialized = report_render._serialize_blocks(blocks)
    assert "正文" in serialized
    assert "onclick" not in serialized
    assert "script" not in serialized
    assert "alert" not in serialized


def test_sanitize_ai_html_unwraps_disallowed_tags_keeping_text():
    blocks = report_render.sanitize_ai_html(
        "<h2>小标题</h2><table><tr><td>单元格</td></tr></table><p>正文</p>"
    )
    serialized = report_render._serialize_blocks(blocks)
    assert "小标题" in serialized
    assert "单元格" in serialized
    assert "<h2>" not in serialized
    assert "<table" not in serialized


def test_sanitize_ai_html_drops_images_and_links():
    blocks = report_render.sanitize_ai_html(
        '<p>看<img src="http://evil/x.png">这里<a href="javascript:x()">链接</a></p>'
    )
    serialized = report_render._serialize_blocks(blocks)
    assert "<img" not in serialized
    assert "href" not in serialized
    assert "链接" in serialized


def test_sanitize_ai_html_wraps_bare_text_in_paragraph():
    blocks = report_render.sanitize_ai_html("模型只返回了纯文本")
    assert [block.tag for block in blocks] == ["p"]
    assert "模型只返回了纯文本" in report_render._serialize_blocks(blocks)


def test_sanitize_ai_html_returns_empty_for_blank_input():
    assert report_render.sanitize_ai_html("") == []
    assert report_render.sanitize_ai_html("   \n  ") == []
    assert report_render.sanitize_ai_html("<p>  </p>") == []


def test_sanitize_ai_html_strips_all_attributes():
    blocks = report_render.sanitize_ai_html(
        '<p class="x" style="color:red" data-foo="1">正文</p>'
    )
    assert report_render._serialize_blocks(blocks) == "<p>正文</p>"


def test_sanitize_ai_html_drops_nested_annotation_marks():
    blocks = report_render.sanitize_ai_html(
        '<p><mark data-annotation-ids="a1">模型回抄了批注标记</mark></p>'
    )
    serialized = report_render._serialize_blocks(blocks)
    assert "<mark" not in serialized
    assert "data-annotation-ids" not in serialized
    assert "模型回抄了批注标记" in serialized


def _stub_polish(monkeypatch, replies):
    """把 execute_prompt 替换为按调用顺序返回预设结果的假实现。"""
    captured = []

    async def fake_execute_prompt(**kwargs):
        captured.append(kwargs)
        return (
            replies[len(captured) - 1] if len(captured) <= len(replies) else replies[-1]
        )

    monkeypatch.setattr(report_render, "execute_prompt", fake_execute_prompt)
    return captured


@pytest.mark.asyncio
async def test_apply_annotations_resolves_llm_once_before_concurrent_polish(
    monkeypatch,
):
    resolved = object()
    resolve_calls = []
    captured = _stub_polish(monkeypatch, ["<p>一</p>", "<p>二</p>"])

    async def fake_resolve_llm(tenant_id):
        resolve_calls.append(tenant_id)
        return resolved

    monkeypatch.setattr(report_render, "resolve_llm", fake_resolve_llm)
    await report_render.apply_annotations_to_html(
        '<p><mark data-annotation-ids="a1">甲</mark></p>'
        '<p><mark data-annotation-ids="a2">乙</mark></p>',
        [
            {"id": "a1", "comment": "改写一", "status": "open"},
            {"id": "a2", "comment": "改写二", "status": "open"},
        ],
        workspace_id="w1",
        session_id="s1",
        tenant_id="t1",
    )

    assert resolve_calls == ["t1"]
    assert [item["llm"] for item in captured] == [resolved, resolved]


@pytest.mark.asyncio
async def test_apply_annotations_overwrite_replaces_region(monkeypatch):
    _stub_polish(monkeypatch, ["<p>本月告警 128 起，同比下降 12%。</p>"])
    html = '<p>概览：<mark data-annotation-ids="a1">本月告警若干</mark>。</p>'
    annotations = [
        {"id": "a1", "comment": "带上同比", "status": "open", "mode": "overwrite"}
    ]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert "本月告警 128 起，同比下降 12%。" in result
    assert "本月告警若干" not in result
    assert "概览：" in result


@pytest.mark.asyncio
async def test_apply_annotations_append_keeps_original(monkeypatch):
    _stub_polish(monkeypatch, ["<p>其中网络类 45 起。</p>"])
    html = '<p><mark data-annotation-ids="a1">本月告警 128 起。</mark></p>'
    annotations = [
        {"id": "a1", "comment": "补充分布", "status": "open", "mode": "append"}
    ]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert "本月告警 128 起。" in result
    assert "其中网络类 45 起。" in result
    assert result.index("本月告警 128 起。") < result.index("其中网络类 45 起。")


@pytest.mark.asyncio
async def test_apply_annotations_append_inserts_list_after_block(monkeypatch):
    _stub_polish(monkeypatch, ["<ul><li>网络类 45 起</li><li>主机类 30 起</li></ul>"])
    html = '<p><mark data-annotation-ids="a1">告警分布如下。</mark></p>'
    annotations = [
        {"id": "a1", "comment": "列出分布", "status": "open", "mode": "append"}
    ]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert "<ul>" in result
    assert result.index("告警分布如下。") < result.index("<ul>")


@pytest.mark.asyncio
async def test_apply_image_annotation_overwrite_uses_append_semantics(monkeypatch):
    captured = _stub_polish(monkeypatch, ["<p>图片补充说明</p>"])
    html = (
        '<p><img src="data:image/png;base64,image" alt="拓扑图" '
        'data-annotation-id="a1" data-annotation-ids="a1"></p>'
    )
    annotations = [
        {"id": "a1", "comment": "补充图片说明", "status": "open", "mode": "overwrite"}
    ]

    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    root = report_render._parse_document(result)
    image = root.find(".//img")
    assert image is not None
    assert image.get("alt") == "拓扑图"
    assert image.getparent().getnext().text == "图片补充说明"
    assert "写入模式：追加写" in captured[0]["message"]
    assert "写入模式：覆盖写" not in captured[0]["message"]


@pytest.mark.asyncio
async def test_apply_annotations_overwrite_with_list_clears_region(monkeypatch):
    _stub_polish(monkeypatch, ["<ul><li>网络类 45 起</li></ul>"])
    html = '<p><mark data-annotation-ids="a1">旧的分布描述</mark></p>'
    annotations = [
        {"id": "a1", "comment": "改成列表", "status": "open", "mode": "overwrite"}
    ]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert "旧的分布描述" not in result
    assert "<ul>" in result
    assert "网络类 45 起" in result


@pytest.mark.asyncio
async def test_apply_annotations_clears_extra_marks_on_overwrite(monkeypatch):
    _stub_polish(monkeypatch, ["<p>统一后的表述</p>"])
    html = (
        '<p><mark data-annotation-ids="a1">前半</mark>'
        '<strong><mark data-annotation-ids="a1">后半</mark></strong></p>'
    )
    annotations = [{"id": "a1", "comment": "合并", "status": "open"}]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert "统一后的表述" in result
    assert "前半" not in result
    assert "后半" not in result


@pytest.mark.asyncio
async def test_apply_annotations_prompt_carries_context_and_mode(monkeypatch):
    captured = _stub_polish(monkeypatch, ["<p>结果</p>"])
    html = (
        "<p>上一段内容</p>"
        '<p><mark data-annotation-ids="a1">被批注文字</mark></p>'
        "<p>下一段内容</p>"
    )
    annotations = [
        {"id": "a1", "comment": "请补充趋势", "status": "open", "mode": "append"}
    ]
    await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    message = captured[0]["message"]
    assert "被批注文字" in message
    assert "请补充趋势" in message
    assert "上一段内容" in message
    assert "下一段内容" in message
    assert "追加" in message
    assert "【被批注内容】是从段落中截取的片段，其前后文字会原样保留" in message
    assert "若原片段末尾没有标点，你的输出末尾也不要加标点" in message
    assert captured[0]["workspace_id"] == "w1"
    assert captured[0]["timeout"] == report_render.ANNOTATION_POLISH_TIMEOUT


@pytest.mark.asyncio
async def test_apply_annotations_skips_when_no_active_annotation(monkeypatch):
    calls = _stub_polish(monkeypatch, ["<p>不该被调用</p>"])
    html = '<p><mark data-annotation-ids="a1">文字</mark></p>'
    annotations = [{"id": "a1", "comment": "已处理", "status": "applied"}]
    result = await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert calls == []
    assert "文字" in result


@pytest.mark.asyncio
async def test_apply_annotations_fails_on_missing_anchor(monkeypatch):
    _stub_polish(monkeypatch, ["<p>结果</p>"])
    html = "<p>正文里没有任何批注锚点</p>"
    annotations = [{"id": "a1", "comment": "改这里", "status": "open"}]
    with pytest.raises(CustomError) as excinfo:
        await report_render.apply_annotations_to_html(
            html,
            annotations,
            workspace_id="w1",
            session_id="s1",
            tenant_id="t1",
        )
    assert excinfo.value.enum == ErrorCode.REPORT_ANNOTATION_ANCHOR_MISSING


@pytest.mark.asyncio
async def test_apply_annotations_fails_when_llm_raises(monkeypatch):
    async def failing_execute_prompt(**kwargs):
        raise RuntimeError("lion_agent down")

    monkeypatch.setattr(report_render, "execute_prompt", failing_execute_prompt)
    html = '<p><mark data-annotation-ids="a1">文字</mark></p>'
    annotations = [{"id": "a1", "comment": "改这里", "status": "open"}]
    with pytest.raises(CustomError) as excinfo:
        await report_render.apply_annotations_to_html(
            html,
            annotations,
            workspace_id="w1",
            session_id="s1",
            tenant_id="t1",
        )
    assert excinfo.value.enum == ErrorCode.REPORT_ANNOTATION_POLISH_FAILED


@pytest.mark.asyncio
async def test_apply_annotations_propagates_lion_agent_unavailable(monkeypatch):
    async def failing_execute_prompt(**kwargs):
        raise CustomError(ErrorCode.LION_AGENT_UNAVAILABLE)

    monkeypatch.setattr(report_render, "execute_prompt", failing_execute_prompt)
    html = '<p><mark data-annotation-ids="a1">文字</mark></p>'
    annotations = [{"id": "a1", "comment": "改这里", "status": "open"}]
    with pytest.raises(CustomError) as excinfo:
        await report_render.apply_annotations_to_html(
            html,
            annotations,
            workspace_id="w1",
            session_id="s1",
            tenant_id="t1",
        )
    assert excinfo.value.enum == ErrorCode.LION_AGENT_UNAVAILABLE


@pytest.mark.asyncio
async def test_apply_annotations_fails_when_llm_returns_blank(monkeypatch):
    _stub_polish(monkeypatch, ["   "])
    html = '<p><mark data-annotation-ids="a1">文字</mark></p>'
    annotations = [{"id": "a1", "comment": "改这里", "status": "open"}]
    with pytest.raises(CustomError) as excinfo:
        await report_render.apply_annotations_to_html(
            html,
            annotations,
            workspace_id="w1",
            session_id="s1",
            tenant_id="t1",
        )
    assert excinfo.value.enum == ErrorCode.REPORT_ANNOTATION_POLISH_FAILED


@pytest.mark.asyncio
async def test_apply_annotations_runs_concurrently(monkeypatch):
    import asyncio

    order = []

    async def fake_execute_prompt(**kwargs):
        order.append("start")
        await asyncio.sleep(0)
        order.append("end")
        return "<p>结果</p>"

    monkeypatch.setattr(report_render, "execute_prompt", fake_execute_prompt)
    html = (
        '<p><mark data-annotation-ids="a1">一</mark></p>'
        '<p><mark data-annotation-ids="a2">二</mark></p>'
    )
    annotations = [
        {"id": "a1", "comment": "改一", "status": "open"},
        {"id": "a2", "comment": "改二", "status": "open"},
    ]
    await report_render.apply_annotations_to_html(
        html, annotations, workspace_id="w1", session_id="s1", tenant_id="t1"
    )
    assert order[:2] == ["start", "start"], "两条批注必须并发发起，不能串行"


def test_strip_annotation_marks_keeps_text():
    html = (
        '<p>前<mark data-annotation-ids="a1">被批注</mark>后</p>'
        '<p><mark data-annotation-id="a2"><strong>加粗批注</strong></mark></p>'
    )
    result = report_render.strip_annotation_marks(html)
    assert "<mark" not in result
    assert "data-annotation" not in result
    assert "前被批注后" in result
    assert "<strong>加粗批注</strong>" in result


def test_strip_annotation_marks_drops_emptied_marks():
    html = '<p>保留<mark data-annotation-ids="a1"></mark></p>'
    result = report_render.strip_annotation_marks(html)
    assert "<mark" not in result
    assert "保留" in result


def test_strip_annotation_marks_removes_only_image_annotation_attributes():
    html = (
        '<p>前<img src="data:image/png;base64,image" alt="拓扑图" width="320" '
        'data-annotation-id="a1" data-annotation-ids="a1" data-asset-id="asset-1">后</p>'
    )
    result = report_render.strip_annotation_marks(html)
    root = report_render._parse_document(result)
    image = root.find(".//img")
    assert image is not None
    assert image.get("src") == "data:image/png;base64,image"
    assert image.get("alt") == "拓扑图"
    assert image.get("width") == "320"
    assert image.get("data-asset-id") == "asset-1"
    assert image.get("data-annotation-id") is None
    assert image.get("data-annotation-ids") is None
    assert "前" in result and "后" in result


@pytest.mark.asyncio
async def test_render_report_docx_runs_fetch_before_annotations(monkeypatch):
    """核心需求验证：AI 必须看到取数后的真实值，而不是占位符 token。"""
    seen_prompts = []

    async def fake_execute(action_cfg, tenant_id, operator_id):
        return '{"data": {"total": 128}}'

    async def fake_execute_prompt(**kwargs):
        seen_prompts.append(kwargs["message"])
        return "<p>本月告警 128 起，处于高位。</p>"

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    monkeypatch.setattr(report_render, "execute_prompt", fake_execute_prompt)

    content_html = (
        '<p>概览：<mark data-annotation-ids="a1">本月告警 '
        '<span data-placeholder-id="alarm" data-result-path="total">'
        "{alarm.total}</span> 起</mark>。</p>"
    )
    docx_bytes = await report_render.render_report_docx(
        content_html=content_html,
        fetch_config={"alarm": {"action": "count"}},
        annotations=[{"id": "a1", "comment": "评估水位", "status": "open"}],
        tenant_id="t1",
        operator_id="u1",
        workspace_id="w1",
        session_id="s1",
    )

    assert len(seen_prompts) == 1
    assert "128" in seen_prompts[0], "批注优化必须在取数之后，AI 应看到真实值"
    assert "{alarm.total}" not in seen_prompts[0]
    assert docx_bytes[:2] == b"PK"


@pytest.mark.asyncio
async def test_visible_at_label_is_not_a_fetch_placeholder(monkeypatch):
    calls = []

    async def fake_execute(action_cfg, tenant_id, operator_id):
        calls.append(action_cfg)
        return '{"data": {"total": 128}}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    html = (
        '<p><span data-placeholder-id="alarm" data-result-path="total">@test</span></p>'
    )
    result = await report_render.resolve_placeholders_in_html(
        html,
        fetch_config={"alarm": {"action": "count"}},
        tenant_id="t1",
        operator_id="u1",
    )
    assert result == html
    assert calls == []


@pytest.mark.asyncio
async def test_render_report_docx_output_has_no_annotation_trace(monkeypatch):
    from docx import Document as DocxDocument
    import io

    async def fake_execute_prompt(**kwargs):
        return "<p>优化后的表述</p>"

    monkeypatch.setattr(report_render, "execute_prompt", fake_execute_prompt)
    docx_bytes = await report_render.render_report_docx(
        content_html='<p><mark data-annotation-ids="a1">原始表述</mark></p>',
        fetch_config={},
        annotations=[{"id": "a1", "comment": "改写", "status": "open"}],
        tenant_id="t1",
        operator_id="u1",
        workspace_id="w1",
        session_id="s1",
    )
    document = DocxDocument(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "优化后的表述" in text
    assert "原始表述" not in text
    assert "data-annotation" not in text


@pytest.mark.asyncio
async def test_render_report_docx_with_image_annotation_preserves_image(monkeypatch):
    image_data_uri = (
        "data:image/png;base64,"
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAIAAACQd1PeAAAADElEQVR4nGP4/"
        "/8/AAX+Av4N70a4AAAAAElFTkSuQmCC"
    )
    _stub_polish(monkeypatch, ["<p>图片补充说明</p>"])

    docx_bytes = await report_render.render_report_docx(
        content_html=(
            f'<p><img src="{image_data_uri}" alt="拓扑图" '
            'data-annotation-id="a1" data-annotation-ids="a1"></p>'
        ),
        fetch_config={},
        annotations=[
            {
                "id": "a1",
                "comment": "补充图片说明",
                "status": "open",
                "mode": "overwrite",
            }
        ],
        tenant_id="t1",
        operator_id="u1",
        workspace_id="w1",
        session_id="s1",
    )

    document = Document(io.BytesIO(docx_bytes))
    assert len(document.inline_shapes) == 1
    assert "图片补充说明" in "\n".join(p.text for p in document.paragraphs)


@pytest.mark.asyncio
async def test_render_report_docx_without_annotations(monkeypatch):
    calls = []

    async def fake_execute_prompt(**kwargs):
        calls.append(kwargs)
        return "<p>不该被调用</p>"

    monkeypatch.setattr(report_render, "execute_prompt", fake_execute_prompt)
    docx_bytes = await report_render.render_report_docx(
        content_html="<h1>月度报告</h1><p>正文内容</p>",
        fetch_config={},
        annotations=[],
        tenant_id="t1",
        operator_id="u1",
        workspace_id="w1",
        session_id="s1",
    )
    assert calls == []
    assert docx_bytes[:2] == b"PK"


@pytest.mark.asyncio
async def test_render_report_docx_preserves_headings_and_lists(monkeypatch):
    from docx import Document as DocxDocument
    import io

    docx_bytes = await report_render.render_report_docx(
        content_html="<h1>月度报告</h1><p>正文</p><ul><li>要点一</li><li>要点二</li></ul>",
        fetch_config={},
        annotations=[],
        tenant_id="t1",
        operator_id="u1",
        workspace_id="w1",
        session_id="s1",
    )
    document = DocxDocument(io.BytesIO(docx_bytes))
    text = "\n".join(p.text for p in document.paragraphs)
    assert "月度报告" in text
    assert "要点一" in text
    assert "要点二" in text


@pytest.mark.asyncio
@pytest.mark.parametrize("shell_name", ["missing.docx", "corrupt.docx"])
async def test_render_report_docx_falls_back_when_shell_is_unavailable(
    shell_name, tmp_path, caplog
):
    shell_path = tmp_path / shell_name
    if shell_name == "corrupt.docx":
        shell_path.write_bytes(b"not-a-docx")

    with caplog.at_level("WARNING", logger=report_render.__name__):
        docx_bytes = await report_render.render_report_docx(
            content_html="<p>降级后正文</p>",
            fetch_config={},
            annotations=[],
            tenant_id="t1",
            operator_id="u1",
            workspace_id="w1",
            session_id="s1",
            base_docx_path=shell_path,
        )

    document = Document(io.BytesIO(docx_bytes))
    assert "降级后正文" in "\n".join(p.text for p in document.paragraphs)
    assert any("降级为从零生成" in record.getMessage() for record in caplog.records)
