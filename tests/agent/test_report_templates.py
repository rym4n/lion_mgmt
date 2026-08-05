# coding=utf-8
import asyncio
import io
from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from core.error_code import ErrorCode
from libs.error_code.exception import CustomError
from models.report import ReportTemplate
from services.agent import report_fetch
from services.agent import views as agent_views
from services.agent.views import (
    _html_to_docx_bytes,
    _report_template_payload,
    _report_template_request_id,
    _resolve_fetch_placeholders,
    _sync_report_template_document,
)


def test_report_template_model_is_owner_scoped_and_serializes_editor_fields():
    fields = ReportTemplate._meta.fields_map

    assert ReportTemplate._meta.db_table == "report_template"
    assert {"tenant_id", "owner", "document", "name", "description", "content_html"} <= set(fields)
    assert {"fetch_config", "annotations", "source_template_key"} <= set(fields)
    assert ReportTemplate._meta.indexes == (("tenant_id", "owner_id"),)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("name", 123),
        ("description", {}),
        ("content_html", []),
        ("fetch_config", []),
        ("annotations", {}),
        ("source_template_key", 123),
    ],
)
def test_report_template_payload_strictly_rejects_wrong_field_types(field, value):
    data = {
        "name": "周报",
        "description": "说明",
        "content_html": "<p>正文</p>",
        "fetch_config": {},
        "annotations": [],
        "source_template_key": None,
    }
    data[field] = value

    with pytest.raises(CustomError):
        _report_template_payload(data, partial=False)


@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("fetch_config", {"alarm": []}),
        ("annotations", ["不是批注对象"]),
    ],
)
def test_report_template_payload_rejects_invalid_nested_json_types(field, value):
    data = {
        "name": "周报",
        "content_html": "<p>正文</p>",
        "fetch_config": {},
        "annotations": [],
    }
    data[field] = value

    with pytest.raises(CustomError):
        _report_template_payload(data, partial=False)


def test_report_template_payload_accepts_annotation_mode():
    payload = agent_views._report_template_payload(
        {
            "name": "月报",
            "content_html": "<p>正文</p>",
            "annotations": [
                {"id": "a1", "comment": "改写", "status": "open", "mode": "append"},
                {"id": "a2", "comment": "补充", "status": "open", "mode": "overwrite"},
                {"id": "a3", "comment": "无模式字段", "status": "open"},
            ],
        },
        partial=False,
    )
    assert len(payload["annotations"]) == 3


def test_report_template_payload_rejects_invalid_annotation_mode():
    with pytest.raises(CustomError) as excinfo:
        agent_views._report_template_payload(
            {
                "name": "月报",
                "content_html": "<p>正文</p>",
                "annotations": [{"id": "a1", "comment": "改写", "mode": "replace"}],
            },
            partial=False,
        )
    assert excinfo.value.enum == ErrorCode.PARAM_ERROR


def test_report_template_payload_rejects_blank_annotation_id():
    with pytest.raises(CustomError) as excinfo:
        agent_views._report_template_payload(
            {
                "name": "月报",
                "content_html": "<p>正文</p>",
                "annotations": [{"id": "  ", "comment": "改写"}],
            },
            partial=False,
        )
    assert excinfo.value.enum == ErrorCode.PARAM_ERROR


def test_report_template_payload_rejects_invalid_annotation_status():
    with pytest.raises(CustomError) as excinfo:
        agent_views._report_template_payload(
            {
                "name": "月报",
                "content_html": "<p>正文</p>",
                "annotations": [{"id": "a1", "comment": "改写", "status": "pending"}],
            },
            partial=False,
        )
    assert excinfo.value.enum == ErrorCode.PARAM_ERROR


def test_report_template_payload_rejects_non_string_comment():
    with pytest.raises(CustomError) as excinfo:
        agent_views._report_template_payload(
            {
                "name": "月报",
                "content_html": "<p>正文</p>",
                "annotations": [{"id": "a1", "comment": {"text": "改写"}}],
            },
            partial=False,
        )
    assert excinfo.value.enum == ErrorCode.PARAM_ERROR


def test_report_template_payload_keeps_unknown_annotation_keys():
    payload = agent_views._report_template_payload(
        {
            "name": "月报",
            "content_html": "<p>正文</p>",
            "annotations": [
                {
                    "id": "a1",
                    "comment": "改写",
                    "status": "open",
                    "selectedText": "原文",
                    "from": 3,
                    "to": 9,
                    "replies": [],
                }
            ],
        },
        partial=False,
    )
    stored = payload["annotations"][0]
    assert stored["selectedText"] == "原文"
    assert stored["from"] == 3
    assert stored["replies"] == []


def test_report_template_request_id_accepts_frontend_and_compatibility_fields():
    assert _report_template_request_id({"template_id": "template-1"}) == "template-1"
    assert _report_template_request_id({"id": "template-2"}) == "template-2"
    assert _report_template_request_id({"template_id": "same", "id": "same"}) == "same"
    with pytest.raises(CustomError):
        _report_template_request_id({"template_id": "one", "id": "two"})


def test_report_template_html_to_docx_preserves_document_contract():
    content = """
        <h1>安全运营周报</h1>
        <p>本周发现 <strong>高危</strong> 事件，趋势为 <em>上升</em>。</p>
        <ul><li>已处置告警</li><li>待复核事件</li></ul>
        <table>
          <thead><tr><th>指标</th><th>结果</th></tr></thead>
          <tbody><tr><td>告警数</td><td><span data-placeholder-id="alarm" data-result-path="data.total">告警数</span></td></tr></tbody>
        </table>
    """

    document = Document(io.BytesIO(_html_to_docx_bytes(content)))

    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].text == "安全运营周报"
    body = document.paragraphs[1]
    assert any(run.text == "高危" and run.bold for run in body.runs)
    assert any(run.text == "上升" and run.italic for run in body.runs)
    assert [paragraph.style.name for paragraph in document.paragraphs[2:4]] == ["List Bullet", "List Bullet"]
    assert len(document.tables) == 1
    assert document.tables[0].cell(0, 0).text == "指标"
    assert document.tables[0].cell(1, 1).text == "{alarm.data.total}"


def test_report_template_html_to_docx_preserves_rich_styles_and_data_image():
    image = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Y9ZlzcAAAAASUVORK5CYII="
    content = f"""
        <p style="text-align: center; line-height: 1.5">
          <u>下划线</u><s>删除线</s>
          <span style="font-size: 20px; font-family: 'Microsoft YaHei'; color: #336699">多样式</span>
        </p>
        <p><img src="data:image/png;base64,{image}" width="24" /></p>
    """

    document = Document(io.BytesIO(_html_to_docx_bytes(content)))
    paragraph = document.paragraphs[0]
    underline = next(run for run in paragraph.runs if run.text == "下划线")
    strike = next(run for run in paragraph.runs if run.text == "删除线")
    styled = next(run for run in paragraph.runs if run.text == "多样式")

    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert paragraph.paragraph_format.line_spacing == 1.5
    assert underline.underline is True
    assert strike.font.strike is True
    assert styled.font.name == "Microsoft YaHei"
    assert styled.font.size.pt == 15
    assert str(styled.font.color.rgb) == "336699"
    assert len(document.inline_shapes) == 1


def test_template_document_sync_reuses_id_and_renames_canonical_file(monkeypatch, tmp_path):
    class FakeDocument:
        records = {}

        def __init__(self, **kwargs):
            self.id = kwargs.pop("id", "doc-1")
            self.__dict__.update(kwargs)

        @classmethod
        async def get_or_none(cls, **kwargs):
            document = cls.records.get(str(kwargs.get("id")))
            if not document:
                return None
            if any(
                getattr(document, key, None) != value
                for key, value in kwargs.items()
                if key not in {"id", "deleted"}
            ):
                return None
            return document if not kwargs.get("deleted", False) else None

        @classmethod
        async def create(cls, **kwargs):
            document = cls(**kwargs)
            document.deleted = False
            cls.records[document.id] = document
            return document

        async def save(self, **_kwargs):
            self.records[self.id] = self

        async def delete(self):
            self.records.pop(self.id, None)

    class FakeTemplate:
        def __init__(self):
            self.id = "template-12345678"
            self.document_id = None
            self.name = "运营周报"
            self.content_html = (
                '<p>告警 <span data-placeholder-id="alarm" data-result-path="total">数量</span></p>'
            )
            self.fetch_config = {"alarm": {"action": "query"}}
            self.source_template_key = None

        async def save(self, **_kwargs):
            return None

    monkeypatch.setattr(agent_views, "REPORT_USER_BASE", str(tmp_path))
    monkeypatch.setattr(agent_views, "ReportDocument", FakeDocument)
    template = FakeTemplate()

    first = asyncio.run(_sync_report_template_document(template, "tenant-1", "user-1"))
    first_path = tmp_path / "tenant-1" / "user-1" / first.file_path
    first_docx = Document(first_path)

    assert template.document_id == "doc-1"
    assert "{alarm.total}" in first_docx.paragraphs[0].text
    assert first.fetch_config == template.fetch_config

    template.name = "运营月报"
    template.fetch_config = {"alarm": {"action": "query_month"}}
    second = asyncio.run(_sync_report_template_document(template, "tenant-1", "user-1"))
    second_path = tmp_path / "tenant-1" / "user-1" / second.file_path

    assert second is first
    assert second.id == "doc-1"
    assert second_path.exists()
    assert second_path != first_path
    assert not first_path.exists()
    assert second.fetch_config == {"alarm": {"action": "query_month"}}


def test_template_document_sync_write_failure_keeps_existing_document(monkeypatch, tmp_path):
    class FakeDocument:
        async def save(self, **_kwargs):
            raise AssertionError("database row must not change when the file write fails")

    document = FakeDocument()
    document.id = "doc-1"
    document.tenant_id = "tenant-1"
    document.owner_id = "user-1"
    document.file_path = "old.docx"
    document.name = "old.docx"
    document.template_key = "template-1"
    document.fetch_config = {}

    class FakeReportDocument:
        @classmethod
        async def get_or_none(cls, **_kwargs):
            return document

    template = SimpleNamespace(
        id="template-1",
        document_id="doc-1",
        name="新名称",
        content_html="<p>新正文</p>",
        fetch_config={},
        source_template_key=None,
    )
    user_dir = tmp_path / "tenant-1" / "user-1"
    user_dir.mkdir(parents=True)
    old_path = user_dir / "old.docx"
    old_path.write_bytes(b"old-content")

    monkeypatch.setattr(agent_views, "REPORT_USER_BASE", str(tmp_path))
    monkeypatch.setattr(agent_views, "ReportDocument", FakeReportDocument)

    def fail_write(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(agent_views, "_atomic_write_report_file", fail_write)

    with pytest.raises(OSError, match="disk full"):
        asyncio.run(_sync_report_template_document(template, "tenant-1", "user-1"))

    assert old_path.read_bytes() == b"old-content"
    assert document.file_path == "old.docx"


def test_fetch_placeholder_replacement_keeps_surrounding_run_format(monkeypatch):
    document = Document()
    paragraph = document.add_paragraph()
    prefix = paragraph.add_run("告警总数：")
    prefix.bold = True
    paragraph.add_run("{alarm.total}")
    suffix = paragraph.add_run(" 条")
    suffix.italic = True
    buffer = io.BytesIO()
    document.save(buffer)

    async def fake_execute(*_args, **_kwargs):
        return '{"data":{"total":7}}'

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fake_execute)
    resolved = asyncio.run(
        _resolve_fetch_placeholders(
            buffer.getvalue(),
            {"alarm": {}},
            "tenant-1",
            "user-1",
        )
    )
    runs = Document(io.BytesIO(resolved)).paragraphs[0].runs

    assert [run.text for run in runs] == ["告警总数：", "7", " 条"]
    assert runs[0].bold is True
    assert runs[2].italic is True


def test_fetch_placeholder_failure_keeps_tokens_and_returns_failure(monkeypatch):
    original = _html_to_docx_bytes(
        '<p>告警 <span data-placeholder-id="alarm" data-result-path="total">数量</span></p>'
    )

    async def fail_execute(*_args, **_kwargs):
        raise RuntimeError("action failed")

    monkeypatch.setattr(report_fetch, "execute_fetch_action", fail_execute)

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            _resolve_fetch_placeholders(
                original,
                {"alarm": {}},
                "tenant-1",
                "user-1",
            )
        )

    assert exc_info.value.enum == ErrorCode.SERVER_ERROR
    assert "{alarm.total}" in Document(io.BytesIO(original)).paragraphs[0].text


def test_fetch_placeholder_timeout_is_propagated_without_partial_report(monkeypatch):
    original = _html_to_docx_bytes(
        '<p><span data-placeholder-id="alarm" data-result-path="total">数量</span></p>'
    )

    async def timeout_execute(*_args, **_kwargs):
        raise CustomError(ErrorCode.ACTION_TIMEOUT)

    monkeypatch.setattr(report_fetch, "execute_fetch_action", timeout_execute)

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            _resolve_fetch_placeholders(
                original,
                {"alarm": {}},
                "tenant-1",
                "user-1",
            )
        )

    assert exc_info.value.enum == ErrorCode.ACTION_TIMEOUT


def test_fetch_placeholder_missing_result_path_is_explicit_failure(monkeypatch):
    original = _html_to_docx_bytes(
        '<p><span data-placeholder-id="alarm" data-result-path="missing.total">数量</span></p>'
    )

    async def successful_action_with_wrong_shape(*_args, **_kwargs):
        return '{"data":{"total":7}}'

    monkeypatch.setattr(
        report_fetch,
        "execute_fetch_action",
        successful_action_with_wrong_shape,
    )

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            _resolve_fetch_placeholders(
                original,
                {"alarm": {}},
                "tenant-1",
                "user-1",
            )
        )

    assert exc_info.value.enum == ErrorCode.SERVER_ERROR
    assert "{alarm.missing.total}" in Document(io.BytesIO(original)).paragraphs[0].text


def test_fetch_placeholder_without_action_config_is_explicit_failure():
    original = _html_to_docx_bytes(
        '<p><span data-placeholder-id="alarm" data-result-path="total">数量</span></p>'
    )

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            _resolve_fetch_placeholders(
                original,
                {},
                "tenant-1",
                "user-1",
            )
        )

    assert exc_info.value.enum == ErrorCode.SERVER_ERROR
    assert "{alarm.total}" in Document(io.BytesIO(original)).paragraphs[0].text


def test_fetch_action_checks_server_permission_before_kafka(monkeypatch):
    permission_calls = []
    published = []

    class FakeAccount:
        def __init__(self):
            self.id = "user-1"
            self.roles = []

    async def deny_permission(
        account,
        app_name,
        action_name,
        resource,
        tenant_id,
        maintenance_mode,
        tool_version,
    ):
        permission_calls.append(
            (
                account.id,
                app_name,
                action_name,
                resource,
                tenant_id,
                maintenance_mode,
                tool_version,
            )
        )
        return False

    async def fake_publish(*args, **kwargs):
        published.append((args, kwargs))

    monkeypatch.setattr(report_fetch, "Account", FakeAccount)
    monkeypatch.setattr(report_fetch, "check_action_execute_permission", deny_permission)
    monkeypatch.setattr(report_fetch, "aproduce", fake_publish)

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            report_fetch.execute_fetch_action(
                {
                    "app_name": "security_tool",
                    "action": "query_alarm",
                    "resource": "resource-1",
                },
                "tenant-1",
                FakeAccount(),
            )
        )

    assert exc_info.value.enum == ErrorCode.UNAUTH
    assert permission_calls == [
        (
            "user-1",
            "security_tool",
            "query_alarm",
            "resource-1",
            "tenant-1",
            False,
            None,
        )
    ]
    assert published == []


def test_fetch_action_failed_runlog_is_not_treated_as_success(monkeypatch):
    filter_calls = []

    class FakeAccount:
        id = "user-1"
        roles = []

    class FakeQuery:
        def order_by(self, *_args):
            return self

        async def first(self):
            return SimpleNamespace(status=-1, result="动作失败")

    class FakeRunlog:
        @classmethod
        def filter(cls, **kwargs):
            filter_calls.append(kwargs)
            return FakeQuery()

    async def allow_permission(*_args, **_kwargs):
        return True

    async def fake_publish(*_args, **_kwargs):
        return None

    monkeypatch.setattr(report_fetch, "Account", FakeAccount)
    monkeypatch.setattr(report_fetch, "Runlog", FakeRunlog)
    monkeypatch.setattr(report_fetch, "check_action_execute_permission", allow_permission)
    monkeypatch.setattr(report_fetch, "aproduce", fake_publish)

    with pytest.raises(RuntimeError, match="report fetch action failed"):
        asyncio.run(
            report_fetch.execute_fetch_action(
                {"app_name": "tool", "action": "query", "timeout": 1},
                "tenant-1",
                FakeAccount(),
            )
        )

    assert len(filter_calls) == 1
    assert filter_calls[0]["tenant_id"] == "tenant-1"
    assert filter_calls[0]["deleted"] is False
    assert filter_calls[0]["run_id"]


def test_fetch_action_timeout_is_explicit_failure(monkeypatch):
    class FakeAccount:
        id = "user-1"
        roles = []

    class FakeQuery:
        def order_by(self, *_args):
            return self

        async def first(self):
            return None

    class FakeRunlog:
        @classmethod
        def filter(cls, **_kwargs):
            return FakeQuery()

    async def allow_permission(*_args, **_kwargs):
        return True

    async def fake_async(*_args, **_kwargs):
        return None

    monkeypatch.setattr(report_fetch, "Account", FakeAccount)
    monkeypatch.setattr(report_fetch, "Runlog", FakeRunlog)
    monkeypatch.setattr(report_fetch, "check_action_execute_permission", allow_permission)
    monkeypatch.setattr(report_fetch, "aproduce", fake_async)
    monkeypatch.setattr(report_fetch.asyncio, "sleep", fake_async)

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            report_fetch.execute_fetch_action(
                {"app_name": "tool", "action": "query", "timeout": 1},
                "tenant-1",
                FakeAccount(),
            )
        )

    assert exc_info.value.enum == ErrorCode.ACTION_TIMEOUT


def test_report_template_routes_permissions_and_generation_contract():
    views_path = Path(__file__).resolve().parents[2] / "services/agent/views.py"
    source = views_path.read_text()
    routes = source[source.index('@agent_bp.post("/report/template/list")') : source.index('@agent_bp.post("/report/templates")')]

    assert '@agent_bp.post("/report/template/detail")' in routes
    assert '@agent_bp.post("/report/template/save")' in routes
    assert '@agent_bp.post("/report/template/delete")' in routes
    assert '@agent_bp.post("/report/template/generate")' in routes
    assert routes.count('@requires_permission("agent:r")') == 2
    assert '@requires_permission("agent:c", "agent:u")' in routes
    assert 'await check_permissions(request, "agent:u")' in routes
    assert 'await check_permissions(request, "agent:c")' in routes
    assert '@requires_permission("agent:d")' in routes
    assert '@requires_permission("agent:u")' in routes
    assert '"template_id": {"type": "string", "description": "模板ID"}' in routes
    assert '"list": [_report_template_summary(template) for template in templates]' in routes
    assert "template.deleted = True" in routes
    assert "copy.deepcopy(template.fetch_config or {})" in routes
    assert "await _sync_report_template_document(" in routes
    assert "await render_report_docx(" in routes
    assert "canonical_path.read_bytes()" not in routes
    generate_route = routes[routes.index('@agent_bp.post("/report/template/generate")') :]
    assert "ReportDocument.create(" not in generate_route
    assert '"doc": document.json' in routes
    assert 'base64.b64encode(generated).decode("utf-8")' in routes


def test_report_template_generate_uses_render_pipeline():
    """generate 必须走 report_render 流水线，而不是直接读 canonical docx 字节。"""
    import inspect

    source = inspect.getsource(agent_views.report_template_generate)
    assert "render_report_docx" in source
    assert "read_bytes()" not in source, "generate 不应再直接读取 canonical docx 字节"
    assert "_get_visible_agent" in source, "generate 需解析 report_assistant 智能体"
    assert "workspace_id_for" in source


def test_report_template_generate_route_and_permission():
    import inspect

    source = inspect.getsource(agent_views)
    marker = '@agent_bp.post("/report/template/generate")'
    assert marker in source
    block = source[source.index(marker) :]
    block = block[: block.index("async def")]
    assert '@requires_permission("agent:u")' in block


def test_report_export_requires_update_permission_and_resolves_canonical_document():
    views_path = Path(__file__).resolve().parents[2] / "services/agent/views.py"
    source = views_path.read_text()
    export_route = source[
        source.index('@agent_bp.post("/report/export")') : source.index('@agent_bp.post("/report/notify/save")')
    ]
    export_helper = source[
        source.index("async def export_report_document(") : source.index("async def send_report_notify_for_document(")
    ]

    assert '@requires_permission("agent:u")' in export_route
    assert "需要 agent:u 权限" in export_route
    assert "resolved_bytes = await _resolve_fetch_placeholders(" in export_helper


def test_report_export_does_not_upload_partial_file_when_fetch_fails(monkeypatch, tmp_path):
    report_dir = tmp_path / "tenant-1" / "user-1"
    report_dir.mkdir(parents=True)
    (report_dir / "canonical.docx").write_bytes(_html_to_docx_bytes("<p>{alarm.total}</p>"))
    document = SimpleNamespace(
        file_path="canonical.docx",
        name="周报.docx",
        fetch_config={"alarm": {}},
    )
    uploads = []

    async def fail_fetch(*_args, **_kwargs):
        raise CustomError(ErrorCode.SERVER_ERROR)

    async def no_export_template(*_args, **_kwargs):
        return None

    async def fake_upload(*args, **kwargs):
        uploads.append((args, kwargs))
        return "unexpected"

    monkeypatch.setattr(agent_views, "REPORT_USER_BASE", str(tmp_path))
    monkeypatch.setattr(agent_views, "_resolve_export_template", no_export_template)
    monkeypatch.setattr(agent_views, "_resolve_fetch_placeholders", fail_fetch)
    monkeypatch.setattr(agent_views, "upload_to_s3", fake_upload)

    with pytest.raises(CustomError) as exc_info:
        asyncio.run(
            agent_views.export_report_document(
                document,
                tenant_id="tenant-1",
                owner_id="user-1",
                operator_id="user-1",
            )
        )

    assert exc_info.value.enum == ErrorCode.SERVER_ERROR
    assert uploads == []


def test_report_template_migration_creates_json_fields_and_owner_index():
    migration = next((Path(__file__).resolve().parents[2] / "migrations/models").glob("73_*_add_report_template.py"))
    source = migration.read_text()

    assert 'CREATE TABLE IF NOT EXISTS "report_template"' in source
    assert '"fetch_config" JSONB NOT NULL' in source
    assert '"annotations" JSONB NOT NULL' in source
    assert '"document_id" VARCHAR(36)' in source
    assert 'REFERENCES "report_document" ("id") ON DELETE SET NULL' in source
    assert 'ON "report_template" ("tenant_id", "owner_id")' in source
    assert 'DROP TABLE IF EXISTS "report_template"' in source
