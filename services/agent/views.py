# coding=utf-8
import asyncio
import base64
import copy
from contextlib import suppress
import io
import os
import re
import tempfile
import uuid
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path

from croniter import croniter
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from sanic import Blueprint, Request
from sanic.log import logger
from sanic_ext import openapi
from tortoise.expressions import Q

from core.error_code import ErrorCode
from core.response import response_ok
from libs.error_code.exception import CustomError
from libs.security.authorization import check_permissions, requires_permission
from libs.utils import gen_run_id, upload_to_s3
from libs.kafka import aproduce
from models import Account
from models.command_center import CommandAgent, CommandTask, CommandTaskEvent
from models.report import ReportDocument, ReportSchedule, ReportTemplate
from services.agent.docx_converter import (
    DocxConversionError,
    docx_base64_to_html,
    html_to_docx_base64,
    replace_paragraph_text_matches,
)
from services.agent.report_fetch import (
    FETCH_PLACEHOLDER_RE,
    build_placeholder_replacer,
    resolve_fetch_values,
)
from services.agent.report_render import render_report_docx
from services.agent.schedule_runner import delete_report_schedule_timer, sync_report_schedule_timer
from services.digital_employee.lion_agent_adapter import execute_prompt, workspace_id_for
from settings import REPORT_TEMPLATE_BASE, REPORT_USER_BASE, RY_SERVICE_URL, TASK_TOPIC, TZ

agent_bp = Blueprint("agent", url_prefix="/agent")

_DOCX_CONVERSION_MAX_CONCURRENT = 2
_DOCX_CONVERSION_MAX_QUEUED = 4
_DOCX_CONVERSION_QUEUE_WAIT_SECONDS = 5.0
_DOCX_CONVERSION_WORKERS = asyncio.BoundedSemaphore(
    _DOCX_CONVERSION_MAX_CONCURRENT
)
_DOCX_CONVERSION_CAPACITY = asyncio.BoundedSemaphore(
    _DOCX_CONVERSION_MAX_CONCURRENT + _DOCX_CONVERSION_MAX_QUEUED
)


class _DocxConversionBusyError(RuntimeError):
    pass


async def _run_docx_conversion(converter, value):
    try:
        await asyncio.wait_for(_DOCX_CONVERSION_CAPACITY.acquire(), timeout=0.05)
    except TimeoutError:
        raise _DocxConversionBusyError from None
    worker_acquired = False
    try:
        try:
            await asyncio.wait_for(
                _DOCX_CONVERSION_WORKERS.acquire(),
                timeout=_DOCX_CONVERSION_QUEUE_WAIT_SECONDS,
            )
            worker_acquired = True
        except TimeoutError:
            raise _DocxConversionBusyError from None

        conversion_task = asyncio.create_task(asyncio.to_thread(converter, value))
        try:
            return await asyncio.shield(conversion_task)
        except asyncio.CancelledError:
            with suppress(Exception):
                await conversion_task
            raise
    finally:
        if worker_acquired:
            _DOCX_CONVERSION_WORKERS.release()
        _DOCX_CONVERSION_CAPACITY.release()


__all__ = ["agent_bp"]

AGENT_SCHEMA = {
    "type": "object",
    "properties": {
        "id": {"type": "string", "description": "智能体ID"},
        "ctime": {"type": "string", "description": "创建时间"},
        "utime": {"type": "string", "description": "更新时间"},
        "agent_key": {"type": "string", "description": "智能体逻辑标识"},
        "name": {"type": "string", "description": "智能体名称"},
        "description": {"type": "string", "nullable": True, "description": "智能体说明"},
        "capabilities": {"type": "array", "items": {"type": "string"}, "description": "能力标签列表"},
        "status": {"type": "string", "description": "状态: idle,busy,offline,error"},
        "load": {"type": "integer", "description": "当前任务负载"},
        "max_concurrent": {"type": "integer", "description": "最大并发任务数"},
        "current_task_id": {"type": "string", "nullable": True, "description": "当前主任务ID"},
        "heartbeat_at": {"type": "string", "nullable": True, "description": "最近心跳时间"},
        "metadata": {"type": "object", "description": "扩展元数据"},
        "runtime": {"type": "string", "description": "运行时类型，默认 lion_agent"},
        "prompt_config": {"type": "object", "description": "提示词扩展配置"},
        "is_builtin": {"type": "boolean", "description": "是否为内置智能体"},
    },
}

ERROR_RESPONSE_SCHEMA = {
    "type": "object",
    "properties": {
        "code": {"type": "integer", "description": "错误码"},
        "message": {"type": "string", "description": "错误信息"},
        "data": {"type": "object", "nullable": True, "description": "错误详情"},
    },
}

EXECUTE_RESULT_SCHEMA = {
    "type": "object",
    "properties": {
        "run_id": {"type": "string", "description": "执行追踪ID"},
        "task_id": {"type": "string", "description": "CommandTask ID"},
        "status": {
            "type": "string",
            "description": "执行状态",
            "enum": ["running", "completed", "failed"],
        },
        "result": {"type": "string", "nullable": True, "description": "执行结果内容"},
    },
}


def _json_body(schema: dict) -> dict:
    return {"application/json": {"schema": schema}}


def _response(description: str, data_schema: dict, errors: dict[str, str] | None = None) -> dict:
    response = {
        "200": {
            "description": description,
            "content": {
                "application/json": {
                    "schema": {
                        "type": "object",
                        "properties": {
                            "code": {"type": "integer", "description": "业务状态码"},
                            "message": {"type": "string", "description": "响应信息"},
                            "data": data_schema,
                        },
                    }
                }
            },
        }
    }
    for status, error_description in (errors or {}).items():
        response[str(status)] = {
            "description": error_description,
            "content": {"application/json": {"schema": ERROR_RESPONSE_SCHEMA}},
        }
    return response


def _require_json(request: Request) -> dict:
    return request.json or {}


def _session_user_id(request: Request) -> str | None:
    account_info = request.ctx.session.get("account_info")
    if isinstance(account_info, dict):
        return account_info.get("id")
    return getattr(account_info, "id", None)


def _normalize_string_list(value, field_name: str) -> list[str]:
    if value in (None, ""):
        return []
    if not isinstance(value, list):
        raise CustomError(ErrorCode.PARAM_ERROR)
    result = []
    for item in value:
        normalized = str(item).strip()
        if normalized:
            result.append(normalized)
    return result


def _normalize_dict(value, field_name: str) -> dict:
    if value in (None, ""):
        return {}
    if not isinstance(value, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)
    return value


async def _get_agent(tenant_id: str, agent_id: str | None = None, agent_key: str | None = None) -> CommandAgent | None:
    if agent_id:
        return await CommandAgent.get_or_none(id=agent_id, tenant_id=tenant_id, deleted=False)
    if agent_key:
        return await CommandAgent.get_or_none(agent_key=agent_key, tenant_id=tenant_id, deleted=False)
    return None


async def _get_visible_agent(
    tenant_id: str,
    agent_key: str | None = None,
) -> CommandAgent | None:
    tenant_agent = await _get_agent(tenant_id, agent_key=agent_key)
    if tenant_agent:
        return tenant_agent
    if agent_key:
        return await CommandAgent.get_or_none(agent_key=agent_key, is_builtin=True, deleted=False)
    return None


async def disable_command_agent(agent: CommandAgent) -> None:
    agent.status = "offline"
    agent.load = 0
    agent.current_task_id = None
    agent.heartbeat_at = datetime.now(TZ)
    await agent.save(update_fields=["status", "load", "current_task_id", "heartbeat_at", "utime"])


def _agent_payload(data: dict, *, partial: bool = False) -> dict:
    payload = {}
    string_fields = ["name", "description", "status", "runtime"]
    for field in string_fields:
        if field in data:
            payload[field] = data[field]

    if "capabilities" in data:
        payload["capabilities"] = _normalize_string_list(data.get("capabilities"), "capabilities")
    elif not partial:
        payload["capabilities"] = []

    if "metadata" in data:
        payload["metadata"] = _normalize_dict(data.get("metadata"), "metadata")
    elif not partial:
        payload["metadata"] = {}

    if "prompt_config" in data:
        payload["prompt_config"] = _normalize_dict(data.get("prompt_config"), "prompt_config")
    elif not partial:
        payload["prompt_config"] = {}

    if "max_concurrent" in data:
        payload["max_concurrent"] = max(1, int(data.get("max_concurrent") or 1))
    elif not partial:
        payload["max_concurrent"] = 1

    if "is_builtin" in data:
        payload["is_builtin"] = bool(data.get("is_builtin"))
    elif not partial:
        payload["is_builtin"] = False

    if not partial:
        payload.setdefault("status", "idle")
        payload.setdefault("runtime", "lion_agent")

    return payload


CREATE_AGENT_BODY = {
    "type": "object",
    "required": ["agent_key", "name"],
    "properties": {
        "agent_key": {"type": "string", "description": "智能体逻辑标识，同一租户内唯一"},
        "name": {"type": "string", "description": "智能体名称"},
        "description": {"type": "string", "description": "智能体说明"},
        "capabilities": {"type": "array", "items": {"type": "string"}, "description": "能力标签列表，指挥中心按此调度"},
        "prompt_config": {"type": "object", "description": "提示词扩展配置，如变量、模板来源等"},
        "is_builtin": {"type": "boolean", "description": "是否为内置智能体，默认false"},
        "runtime": {"type": "string", "description": "运行时类型，默认 lion_agent"},
        "metadata": {"type": "object", "description": "扩展元数据"},
        "status": {"type": "string", "description": "状态: idle,busy,offline,error", "default": "idle"},
        "max_concurrent": {"type": "integer", "minimum": 1, "description": "最大并发任务数"},
    },
}


@agent_bp.post("/create")
@openapi.definition(
    summary="创建智能体",
    description="创建一个可被指挥中心调度的智能体配置。数据落在command_center_agent表。需要agent:c权限。",
    body=_json_body(CREATE_AGENT_BODY),
    response=_response("创建成功", AGENT_SCHEMA, {"400": "参数错误", "409": "智能体已存在"}),
)
@requires_permission("agent:c")
async def create_agent(request: Request):
    data = _require_json(request)
    agent_key = str(data.get("agent_key") or "").strip()
    name = str(data.get("name") or "").strip()
    if not agent_key or not name:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    exists = await CommandAgent.get_or_none(tenant_id=tenant_id, agent_key=agent_key, deleted=False)
    if exists:
        raise CustomError(ErrorCode.DATA_ALREADY_EXISTS)

    payload = _agent_payload(data)
    payload["agent_key"] = agent_key
    payload["name"] = name
    agent = await CommandAgent.create(tenant_id=tenant_id, **payload)
    return response_ok(agent.json)


@agent_bp.post("/list")
@openapi.definition(
    summary="查询智能体列表",
    description="分页查询智能体配置，可按状态、内置标识、运行时和关键字过滤。需要agent:r权限。",
    body=_json_body(
        {
            "type": "object",
            "properties": {
                "page": {"type": "integer", "minimum": 1, "default": 1},
                "size": {"type": "integer", "minimum": 1, "default": 15},
                "status": {"type": "string", "description": "状态过滤"},
                "runtime": {"type": "string", "description": "运行时过滤"},
                "is_builtin": {"type": "boolean", "description": "是否内置过滤"},
                "keyword": {"type": "string", "description": "按名称、标识、说明模糊搜索"},
            },
        }
    ),
    response=_response(
        "查询成功",
        {
            "type": "object",
            "properties": {
                "total": {"type": "integer"},
                "page": {"type": "integer"},
                "size": {"type": "integer"},
                "data": {"type": "array", "items": AGENT_SCHEMA},
            },
        },
    ),
)
@requires_permission("agent:r")
async def list_agents(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    page = max(1, int(data.get("page", 1) or 1))
    size = max(1, min(int(data.get("size", 15) or 15), 100))

    # 内置 agent 不分租户，所有租户都能看到；非内置 agent 按租户隔离
    base_condition = Q(deleted=False) & (Q(tenant_id=tenant_id) | Q(is_builtin=True))
    extra = Q()
    if data.get("status"):
        extra &= Q(status=data["status"])
    if data.get("runtime"):
        extra &= Q(runtime=data["runtime"])
    if "is_builtin" in data:
        extra &= Q(is_builtin=bool(data.get("is_builtin")))
    if data.get("keyword"):
        keyword = data["keyword"]
        extra &= Q(agent_key__icontains=keyword) | Q(name__icontains=keyword) | Q(description__icontains=keyword)
    condition = base_condition & extra

    queryset = CommandAgent.filter(condition).distinct()
    total = await queryset.count()
    agents = await queryset.order_by("-is_builtin", "status", "agent_key").offset((page - 1) * size).limit(size).all()
    return response_ok({"total": total, "page": page, "size": size, "data": [item.json for item in agents]})


@agent_bp.post("/detail")
@openapi.definition(
    summary="查询智能体详情",
    description="按id或agent_key查询智能体详情。需要agent:r权限。",
    body=_json_body(
        {
            "type": "object",
            "oneOf": [{"required": ["id"]}, {"required": ["agent_key"]}],
            "properties": {
                "id": {"type": "string", "description": "智能体ID"},
                "agent_key": {"type": "string", "description": "智能体逻辑标识"},
            },
        }
    ),
    response=_response("查询成功", AGENT_SCHEMA, {"400": "参数错误", "404": "智能体不存在"}),
)
@requires_permission("agent:r")
async def detail_agent(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    agent = await _get_agent(tenant_id, data.get("id"), data.get("agent_key"))
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)
    return response_ok(agent.json)


@agent_bp.post("/update")
@openapi.definition(
    summary="更新智能体",
    description="更新智能体基础信息、能力标签、lion_agent 运行时和提示词配置。内置智能体不可改为非内置。需要agent:u权限。",
    body=_json_body(
        {
            "type": "object",
            "oneOf": [{"required": ["id"]}, {"required": ["agent_key"]}],
            "properties": {**CREATE_AGENT_BODY["properties"], "id": {"type": "string", "description": "智能体ID"}},
        }
    ),
    response=_response("更新成功", AGENT_SCHEMA, {"400": "参数错误", "404": "智能体不存在"}),
)
@requires_permission("agent:u")
async def update_agent(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    agent = await _get_agent(tenant_id, data.get("id"), data.get("agent_key"))
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)

    payload = _agent_payload(data, partial=True)
    if agent.is_builtin and payload.get("is_builtin") is False:
        raise CustomError(ErrorCode.INTERNAL_DELETE_ERROR)
    if "agent_key" in data and data.get("agent_key") and data.get("agent_key") != agent.agent_key:
        new_key = str(data["agent_key"]).strip()
        exists = await CommandAgent.get_or_none(tenant_id=tenant_id, agent_key=new_key, deleted=False)
        if exists and exists.id != agent.id:
            raise CustomError(ErrorCode.DATA_ALREADY_EXISTS)
        payload["agent_key"] = new_key

    for field, value in payload.items():
        setattr(agent, field, value)
    await agent.save()
    return response_ok(agent.json)


@agent_bp.post("/delete")
@openapi.definition(
    summary="删除智能体",
    description="软删除人工创建的智能体，并同步将指挥中心调度状态置为offline。内置智能体禁止删除。需要agent:d权限。",
    body=_json_body(
        {
            "type": "object",
            "oneOf": [{"required": ["id"]}, {"required": ["agent_key"]}],
            "properties": {
                "id": {"type": "string", "description": "智能体ID"},
                "agent_key": {"type": "string", "description": "智能体逻辑标识"},
            },
        }
    ),
    response=_response("删除成功", {"type": "object"}, {"400": "参数错误", "404": "智能体不存在"}),
)
@requires_permission("agent:d")
async def delete_agent(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    agent = await _get_agent(tenant_id, data.get("id"), data.get("agent_key"))
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)
    if agent.is_builtin:
        raise CustomError(ErrorCode.INTERNAL_DELETE_ERROR)

    await disable_command_agent(agent)
    agent.deleted = True
    await agent.save(update_fields=["deleted", "utime"])
    return response_ok({"id": agent.id, "deleted": True})


async def _run_agent_task(task_id: str, tenant_id: str, agent: CommandAgent, prompt: str, run_id: str, operator_id: str) -> None:
    """后台异步执行智能体任务，结果写回 CommandTask"""
    task = await CommandTask.get_or_none(id=task_id, tenant_id=tenant_id, deleted=False)
    if not task:
        return

    try:
        result = await execute_prompt(
            workspace_id=workspace_id_for(tenant_id, agent.agent_key),
            session_id=run_id,
            request_id=task_id,
            message=prompt,
        )
        task.status = "completed"
        task.progress = 100
        task.result = result
        task.finished_at = datetime.now(TZ)
        await task.save(update_fields=["status", "progress", "result", "finished_at", "utime"])
        await CommandTaskEvent.create(
            tenant_id=tenant_id,
            task_id=task_id,
            event_type="task_completed",
            status="completed",
            progress=100,
            agent_key=agent.agent_key,
            message="任务执行完成",
            metadata={"run_id": run_id},
            operator_id=operator_id,
        )
    except Exception as exc:
        task.status = "failed"
        task.finished_at = datetime.now(TZ)
        await task.save(update_fields=["status", "finished_at", "utime"])
        await CommandTaskEvent.create(
            tenant_id=tenant_id,
            task_id=task_id,
            event_type="task_failed",
            status="failed",
            agent_key=agent.agent_key,
            message=f"任务执行失败: {exc}",
            metadata={"run_id": run_id},
            operator_id=operator_id,
        )


@agent_bp.post("/execute")
@openapi.definition(
    summary="向智能体下发执行任务",
    description="指定智能体和任务内容，后台异步调用 lion_agent 执行任务，立即返回 run_id 和 task_id 用于后续查询结果。需要 agent:u 权限。",
    body=_json_body(
        {
            "type": "object",
            "required": ["prompt"],
            "properties": {
                "agent_id": {"type": "string", "description": "智能体 ID（与 agent_key 二选一）"},
                "agent_key": {"type": "string", "description": "智能体逻辑标识（与 agent_id 二选一）"},
                "prompt": {"type": "string", "description": "任务内容/提示词"},
                "title": {"type": "string", "description": "兼容字段，任务展示统一使用 prompt"},
                "room_id": {"type": "string", "description": "可选，关联作战室 ID"},
            },
        }
    ),
    response=_response(
        "任务已提交",
        EXECUTE_RESULT_SCHEMA,
        {"400": "参数错误", "404": "智能体不存在"},
    ),
)
@requires_permission("agent:u")
async def execute_agent(request: Request):
    data = _require_json(request)
    prompt = str(data.get("prompt") or "").strip()
    if not prompt:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    agent = await _get_agent(tenant_id, data.get("agent_id"), data.get("agent_key"))
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)

    operator_id = _session_user_id(request)
    run_id = gen_run_id()

    task = await CommandTask.create(
        tenant_id=tenant_id,
        title=prompt[:255],
        description=prompt,
        source="agent_execute",
        status="running",
        progress=0,
        payload={"prompt": prompt},
        assigned_agent_id=agent.id,
        assigned_agent_key=agent.agent_key,
        dispatch_status="user_specified",
        run_id=run_id,
        room_id=data.get("room_id"),
        creator_id=operator_id,
        started_at=datetime.now(TZ),
    )
    await CommandTaskEvent.create(
        tenant_id=tenant_id,
        task_id=task.id,
        event_type="task_created",
        status="running",
        progress=0,
        agent_key=agent.agent_key,
        message="任务已创建并开始执行",
        metadata={"run_id": run_id},
        operator_id=operator_id,
    )

    asyncio.create_task(_run_agent_task(str(task.id), tenant_id, agent, prompt, run_id, operator_id or ""))

    return response_ok({"run_id": run_id, "task_id": str(task.id), "status": "running"})


@agent_bp.post("/execute/result")
@openapi.definition(
    summary="查询智能体任务执行结果",
    description="按 run_id 或 task_id 查询任务执行状态和结果。需要 agent:r 权限。",
    body=_json_body(
        {
            "type": "object",
            "oneOf": [{"required": ["run_id"]}, {"required": ["task_id"]}],
            "properties": {
                "run_id": {"type": "string", "description": "执行追踪ID"},
                "task_id": {"type": "string", "description": "任务ID"},
            },
        }
    ),
    response=_response(
        "查询成功",
        EXECUTE_RESULT_SCHEMA,
        {"400": "参数错误", "404": "任务不存在"},
    ),
)
@requires_permission("agent:r")
async def execute_result(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")

    run_id = data.get("run_id")
    task_id = data.get("task_id")
    if not run_id and not task_id:
        raise CustomError(ErrorCode.PARAM_ERROR)

    if task_id:
        task = await CommandTask.get_or_none(id=task_id, tenant_id=tenant_id, deleted=False)
    else:
        task = await CommandTask.filter(run_id=run_id, tenant_id=tenant_id, deleted=False).first()

    if not task:
        raise CustomError(ErrorCode.NOT_FOUND)

    return response_ok({
        "run_id": task.run_id,
        "task_id": str(task.id),
        "status": task.status,
        "result": task.result,
    })


# ---------------------------------------------------------------------------
# 报告模板 & 文档接口
# ---------------------------------------------------------------------------

REPORT_TEMPLATE_SCHEMA = {
    "type": "object",
    "properties": {
        "key": {"type": "string", "description": "模板标识（文件名不含扩展名）"},
        "name": {"type": "string", "description": "模板显示名称"},
        "filename": {"type": "string", "description": "文件名"},
    },
}

REPORT_DOC_SCHEMA = {
    "type": "object",
    "properties": {
        "id": {"type": "string", "description": "文档ID"},
        "name": {"type": "string", "description": "文件名"},
        "template_key": {"type": "string", "nullable": True, "description": "来源模板标识"},
        "ctime": {"type": "string", "description": "创建时间"},
        "utime": {"type": "string", "description": "更新时间"},
    },
}

REPORT_EDITABLE_TEMPLATE_SCHEMA = {
    "type": "object",
    "properties": {
        "id": {"type": "string", "description": "模板ID"},
        "document_id": {"type": "string", "nullable": True, "description": "定时任务复用的报告文档ID"},
        "name": {"type": "string", "description": "模板名称"},
        "description": {"type": "string", "description": "模板说明"},
        "content_html": {"type": "string", "description": "在线编辑器 HTML 内容"},
        "fetch_config": {"type": "object", "description": "取数占位符动作配置"},
        "annotations": {"type": "array", "description": "文档批注"},
        "source_template_key": {"type": "string", "nullable": True, "description": "来源内置模板标识"},
        "ctime": {"type": "string", "description": "创建时间"},
        "utime": {"type": "string", "description": "更新时间"},
    },
}

REPORT_EDITABLE_TEMPLATE_SUMMARY_SCHEMA = {
    "type": "object",
    "properties": {
        "id": {"type": "string", "description": "模板ID"},
        "name": {"type": "string", "description": "模板名称"},
        "description": {"type": "string", "description": "模板说明"},
        "source_template_key": {"type": "string", "nullable": True, "description": "来源内置模板标识"},
        "placeholder_count": {"type": "integer", "description": "取数占位符数量"},
        "annotation_count": {"type": "integer", "description": "批注数量"},
        "ctime": {"type": "string", "description": "创建时间"},
        "utime": {"type": "string", "description": "更新时间"},
    },
}


class _ReportHTMLNode:
    __slots__ = ("tag", "attrs", "children", "text")

    def __init__(self, tag: str | None, attrs: dict | None = None, text: str = ""):
        self.tag = tag
        self.attrs = attrs or {}
        self.children: list[_ReportHTMLNode] = []
        self.text = text


class _ReportHTMLParser(HTMLParser):
    _VOID_TAGS = {"area", "base", "br", "col", "embed", "hr", "img", "input", "link", "meta", "param", "source", "track", "wbr"}
    _IGNORED_TAGS = {"head", "script", "style"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = _ReportHTMLNode("root")
        self.stack = [self.root]

    def handle_starttag(self, tag: str, attrs):
        tag = tag.lower()
        node = _ReportHTMLNode(tag, {str(key).lower(): value for key, value in attrs})
        self.stack[-1].children.append(node)
        if tag not in self._VOID_TAGS:
            self.stack.append(node)

    def handle_startendtag(self, tag: str, attrs):
        self.handle_starttag(tag, attrs)
        if tag.lower() not in self._VOID_TAGS:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str):
        tag = tag.lower()
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                return

    def handle_data(self, data: str):
        if data and not any(node.tag in self._IGNORED_TAGS for node in self.stack):
            self.stack[-1].children.append(_ReportHTMLNode(None, text=data))


_HTML_BLOCK_TAGS = {
    "address", "article", "aside", "blockquote", "div", "dl", "fieldset", "figure", "footer",
    "form", "h1", "h2", "h3", "h4", "h5", "h6", "header", "hr", "li", "main", "nav",
    "ol", "p", "pre", "section", "table", "ul",
}


def _placeholder_token(node: _ReportHTMLNode) -> str | None:
    if node.tag != "span" or "data-placeholder-id" not in node.attrs or "data-result-path" not in node.attrs:
        return None
    placeholder_id = str(node.attrs.get("data-placeholder-id") or "").strip()
    if not placeholder_id:
        return None
    result_path = str(node.attrs.get("data-result-path") or "").strip().strip(".")
    return "{" + placeholder_id + (f".{result_path}" if result_path else "") + "}"


def _css_declarations(node: _ReportHTMLNode) -> dict[str, str]:
    style = node.attrs.get("style") if node.tag is not None else None
    if not isinstance(style, str):
        return {}
    declarations = {}
    for item in style.split(";"):
        if ":" not in item:
            continue
        key, value = item.split(":", 1)
        if key.strip() and value.strip():
            declarations[key.strip().lower()] = value.strip()
    return declarations


def _css_font_size(value: str | None):
    if not value:
        return None
    normalized = value.strip().lower()
    match = re.fullmatch(r"(-?\d+(?:\.\d+)?)(px|pt|em|rem)?", normalized)
    if not match:
        return None
    amount = float(match.group(1))
    if amount <= 0:
        return None
    unit = match.group(2) or "pt"
    if unit == "px":
        amount *= 0.75
    elif unit in {"em", "rem"}:
        amount *= 12
    return Pt(amount)


def _css_color(value: str | None):
    if not value:
        return None
    normalized = value.strip().lower()
    named_colors = {
        "black": "000000",
        "white": "FFFFFF",
        "red": "FF0000",
        "green": "008000",
        "blue": "0000FF",
        "yellow": "FFFF00",
        "gray": "808080",
        "grey": "808080",
        "orange": "FFA500",
        "purple": "800080",
    }
    normalized = named_colors.get(normalized, normalized)
    if normalized.startswith("#"):
        normalized = normalized[1:]
    if re.fullmatch(r"[0-9a-f]{3}", normalized):
        normalized = "".join(char * 2 for char in normalized)
    if re.fullmatch(r"[0-9a-f]{6}", normalized):
        return RGBColor.from_string(normalized.upper())
    match = re.fullmatch(
        r"rgba?\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})(?:\s*,[^)]*)?\)",
        normalized,
    )
    if match:
        channels = [min(255, int(channel)) for channel in match.groups()]
        return RGBColor(*channels)
    return None


def _inline_format(node: _ReportHTMLNode, inherited: dict | None = None) -> dict:
    formatting = dict(inherited or {})
    if node.tag is None:
        return formatting
    styles = _css_declarations(node)
    if node.tag in {"b", "strong"}:
        formatting["bold"] = True
    if node.tag in {"em", "i"}:
        formatting["italic"] = True
    if node.tag == "u":
        formatting["underline"] = True
    if node.tag in {"del", "s", "strike"}:
        formatting["strike"] = True

    font_weight = styles.get("font-weight", "").lower()
    if font_weight:
        formatting["bold"] = font_weight in {"bold", "bolder"} or (
            font_weight.isdigit() and int(font_weight) >= 600
        )
    font_style = styles.get("font-style", "").lower()
    if font_style:
        formatting["italic"] = font_style in {"italic", "oblique"}
    decoration = " ".join(
        value.lower()
        for key, value in styles.items()
        if key in {"text-decoration", "text-decoration-line"}
    )
    if decoration:
        if "none" in decoration:
            formatting["underline"] = False
            formatting["strike"] = False
        else:
            if "underline" in decoration:
                formatting["underline"] = True
            if "line-through" in decoration:
                formatting["strike"] = True

    font_name = styles.get("font-family") or node.attrs.get("face")
    if isinstance(font_name, str) and font_name.strip():
        formatting["font_name"] = font_name.split(",", 1)[0].strip().strip("'\"")
    font_size = _css_font_size(styles.get("font-size"))
    if font_size is None and node.tag == "font" and node.attrs.get("size"):
        html_sizes = {"1": 8, "2": 10, "3": 12, "4": 14, "5": 18, "6": 24, "7": 36}
        size_value = html_sizes.get(str(node.attrs.get("size")).strip())
        font_size = Pt(size_value) if size_value else None
    if font_size is not None:
        formatting["font_size"] = font_size
    color = _css_color(styles.get("color") or node.attrs.get("color"))
    if color is not None:
        formatting["color"] = color
    return formatting


def _apply_run_format(run, formatting: dict) -> None:
    if "bold" in formatting:
        run.bold = bool(formatting["bold"])
    if "italic" in formatting:
        run.italic = bool(formatting["italic"])
    if "underline" in formatting:
        run.underline = bool(formatting["underline"])
    if "strike" in formatting:
        run.font.strike = bool(formatting["strike"])
    if formatting.get("font_name"):
        run.font.name = formatting["font_name"]
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), formatting["font_name"])
    if formatting.get("font_size") is not None:
        run.font.size = formatting["font_size"]
    if formatting.get("color") is not None:
        run.font.color.rgb = formatting["color"]


def _image_width(node: _ReportHTMLNode):
    styles = _css_declarations(node)
    value = styles.get("width") or node.attrs.get("width")
    if not isinstance(value, str):
        return None
    match = re.fullmatch(r"(\d+(?:\.\d+)?)(px|pt|in|cm|mm)?", value.strip().lower())
    if not match or float(match.group(1)) <= 0:
        return None
    amount = float(match.group(1))
    unit = match.group(2) or "px"
    if unit == "px":
        return Inches(amount / 96)
    if unit == "pt":
        return Pt(amount)
    if unit == "cm":
        return Inches(amount / 2.54)
    if unit == "mm":
        return Inches(amount / 25.4)
    return Inches(amount)


def _append_data_image(paragraph, node: _ReportHTMLNode, formatting: dict) -> None:
    source = node.attrs.get("src")
    if not isinstance(source, str) or not source.lower().startswith("data:image/") or "," not in source:
        return
    header, encoded = source.split(",", 1)
    if ";base64" not in header.lower():
        return
    try:
        image_bytes = base64.b64decode(re.sub(r"\s+", "", encoded), validate=True)
        run = paragraph.add_run()
        _apply_run_format(run, formatting)
        run.add_picture(io.BytesIO(image_bytes), width=_image_width(node))
    except Exception as exc:
        logger.warning("忽略无法写入 DOCX 的 data URL 图片: %s", exc)


def _append_html_inline(paragraph, node: _ReportHTMLNode, *, formatting: dict | None = None) -> None:
    current_format = _inline_format(node, formatting)
    if node.tag is None:
        text = re.sub(r"[\r\n\t]+", " ", node.text)
        if text:
            run = paragraph.add_run(text)
            _apply_run_format(run, current_format)
        return

    if node.tag in _ReportHTMLParser._IGNORED_TAGS:
        return
    token = _placeholder_token(node)
    if token is not None:
        run = paragraph.add_run(token)
        _apply_run_format(run, current_format)
        return
    if node.tag == "br":
        paragraph.add_run("\n")
        return
    if node.tag == "img":
        _append_data_image(paragraph, node, current_format)
        return

    for child in node.children:
        _append_html_inline(paragraph, child, formatting=current_format)


def _node_has_content(node: _ReportHTMLNode) -> bool:
    if node.tag is None:
        return bool(node.text.strip())
    if _placeholder_token(node) is not None or node.tag in {"br", "hr", "img", "table"}:
        return True
    return any(_node_has_content(child) for child in node.children)


def _apply_paragraph_format(paragraph, node: _ReportHTMLNode | None) -> None:
    if node is None:
        return
    styles = _css_declarations(node)
    alignment = styles.get("text-align") or node.attrs.get("align")
    if isinstance(alignment, str):
        alignment_map = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "start": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "end": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }
        if alignment.strip().lower() in alignment_map:
            paragraph.alignment = alignment_map[alignment.strip().lower()]

    line_height = styles.get("line-height")
    if not line_height:
        return
    normalized = line_height.strip().lower()
    try:
        if normalized.endswith("%"):
            paragraph.paragraph_format.line_spacing = float(normalized[:-1]) / 100
        elif re.fullmatch(r"\d+(?:\.\d+)?", normalized):
            paragraph.paragraph_format.line_spacing = float(normalized)
        else:
            spacing = _css_font_size(normalized)
            if spacing is not None:
                paragraph.paragraph_format.line_spacing = spacing
    except ValueError:
        return


def _add_docx_paragraph(container, *, style: str | None = None, node: _ReportHTMLNode | None = None):
    if style:
        try:
            paragraph = container.add_paragraph(style=style)
        except KeyError:
            paragraph = container.add_paragraph()
    else:
        paragraph = container.add_paragraph()
    _apply_paragraph_format(paragraph, node)
    return paragraph


def _append_list_item_inline(paragraph, node: _ReportHTMLNode) -> None:
    for child in node.children:
        if child.tag not in {"ul", "ol"}:
            _append_html_inline(paragraph, child)


def _add_html_list(container, node: _ReportHTMLNode, *, ordered: bool) -> None:
    style = "List Number" if ordered else "List Bullet"
    for item in (child for child in node.children if child.tag == "li"):
        paragraph = _add_docx_paragraph(container, style=style, node=item)
        _append_list_item_inline(paragraph, item)
        for nested in (child for child in item.children if child.tag in {"ul", "ol"}):
            _add_html_list(container, nested, ordered=nested.tag == "ol")


def _table_rows(node: _ReportHTMLNode) -> list[_ReportHTMLNode]:
    rows = []
    for child in node.children:
        if child.tag == "tr":
            rows.append(child)
        elif child.tag in {"thead", "tbody", "tfoot"}:
            rows.extend(_table_rows(child))
    return rows


def _remove_empty_first_cell_paragraph(cell) -> None:
    if len(cell.paragraphs) <= 1 or cell.paragraphs[0].text:
        return
    paragraph = cell.paragraphs[0]._element
    paragraph.getparent().remove(paragraph)


def _add_html_table(container, node: _ReportHTMLNode) -> None:
    rows = _table_rows(node)
    cell_rows = [[child for child in row.children if child.tag in {"td", "th"}] for row in rows]
    column_count = max((len(cells) for cells in cell_rows), default=0)
    if not column_count:
        return
    table = container.add_table(rows=0, cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for source_cells in cell_rows:
        target_cells = table.add_row().cells
        for index, source_cell in enumerate(source_cells):
            _add_html_blocks(target_cells[index], source_cell.children)
            _remove_empty_first_cell_paragraph(target_cells[index])
            for paragraph in target_cells[index].paragraphs:
                _apply_paragraph_format(paragraph, source_cell)


def _add_html_block(container, node: _ReportHTMLNode) -> None:
    if node.tag in {"html", "body", "main", "article", "section", "header", "footer", "aside", "nav", "div", "figure", "form", "fieldset", "dl"}:
        _add_html_blocks(container, node.children)
        return
    if node.tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
        paragraph = _add_docx_paragraph(container, style=f"Heading {node.tag[1:]}", node=node)
        _append_html_inline(paragraph, node)
        return
    if node.tag in {"p", "pre", "blockquote", "address"}:
        paragraph = _add_docx_paragraph(
            container,
            style="Quote" if node.tag == "blockquote" else None,
            node=node,
        )
        _append_html_inline(paragraph, node)
        return
    if node.tag in {"ul", "ol"}:
        _add_html_list(container, node, ordered=node.tag == "ol")
        return
    if node.tag == "table":
        _add_html_table(container, node)
        return
    if node.tag == "hr":
        _add_docx_paragraph(container, node=node).add_run("---")
        return
    paragraph = _add_docx_paragraph(container, node=node)
    _append_html_inline(paragraph, node)


def _add_html_blocks(container, nodes: list[_ReportHTMLNode]) -> None:
    inline_nodes: list[_ReportHTMLNode] = []

    def flush_inline() -> None:
        if not any(_node_has_content(item) for item in inline_nodes):
            inline_nodes.clear()
            return
        paragraph = _add_docx_paragraph(container)
        for item in inline_nodes:
            _append_html_inline(paragraph, item)
        inline_nodes.clear()

    for node in nodes:
        if node.tag in _HTML_BLOCK_TAGS or node.tag in {"html", "body"}:
            flush_inline()
            _add_html_block(container, node)
        else:
            inline_nodes.append(node)
    flush_inline()


def _html_to_docx_bytes(content_html: str) -> bytes:
    parser = _ReportHTMLParser()
    parser.feed(content_html)
    parser.close()
    document = Document()
    _add_html_blocks(document, parser.root.children)
    if not document.paragraphs and not document.tables:
        document.add_paragraph()
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


_REPORT_ANNOTATION_MODES = frozenset({"overwrite", "append"})
_REPORT_ANNOTATION_STATUSES = frozenset({"open", "applied", "resolved"})


def _report_template_payload(data: dict, *, partial: bool) -> dict:
    if not isinstance(data, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)
    payload = {}
    required = set() if partial else {"name", "content_html"}
    if any(field not in data for field in required):
        raise CustomError(ErrorCode.PARAM_ERROR)

    if "name" in data:
        if not isinstance(data["name"], str):
            raise CustomError(ErrorCode.PARAM_ERROR)
        name = data["name"].strip()
        if not name or len(name) > 255:
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["name"] = name
    if "description" in data:
        if not isinstance(data["description"], str):
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["description"] = data["description"]
    elif not partial:
        payload["description"] = ""
    if "content_html" in data:
        if not isinstance(data["content_html"], str):
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["content_html"] = data["content_html"]
    if "fetch_config" in data:
        if not isinstance(data["fetch_config"], dict):
            raise CustomError(ErrorCode.PARAM_ERROR)
        if any(
            not isinstance(key, str) or not key.strip() or not isinstance(value, dict)
            for key, value in data["fetch_config"].items()
        ):
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["fetch_config"] = copy.deepcopy(data["fetch_config"])
    elif not partial:
        payload["fetch_config"] = {}
    if "annotations" in data:
        if not isinstance(data["annotations"], list):
            raise CustomError(ErrorCode.PARAM_ERROR)
        for item in data["annotations"]:
            if not isinstance(item, dict):
                raise CustomError(ErrorCode.PARAM_ERROR)
            if not isinstance(item.get("id"), str) or not item["id"].strip():
                raise CustomError(ErrorCode.PARAM_ERROR)
            if "comment" in item and not isinstance(item["comment"], str):
                raise CustomError(ErrorCode.PARAM_ERROR)
            if "mode" in item and item["mode"] not in _REPORT_ANNOTATION_MODES:
                raise CustomError(ErrorCode.PARAM_ERROR)
            if "status" in item and item["status"] not in _REPORT_ANNOTATION_STATUSES:
                raise CustomError(ErrorCode.PARAM_ERROR)
        payload["annotations"] = copy.deepcopy(data["annotations"])
    elif not partial:
        payload["annotations"] = []
    if "source_template_key" in data:
        source_key = data["source_template_key"]
        if source_key is not None and not isinstance(source_key, str):
            raise CustomError(ErrorCode.PARAM_ERROR)
        source_key = source_key.strip() if isinstance(source_key, str) else None
        if source_key and len(source_key) > 128:
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["source_template_key"] = source_key or None
    elif not partial:
        payload["source_template_key"] = None
    return payload


def _report_template_request_id(data: dict) -> str:
    if not isinstance(data, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)
    raw_template_id = data.get("template_id")
    raw_id = data.get("id")
    if raw_template_id is not None and not isinstance(raw_template_id, str):
        raise CustomError(ErrorCode.PARAM_ERROR)
    if raw_id is not None and not isinstance(raw_id, str):
        raise CustomError(ErrorCode.PARAM_ERROR)
    if raw_template_id is not None and raw_id is not None and raw_template_id.strip() != raw_id.strip():
        raise CustomError(ErrorCode.PARAM_ERROR)
    template_id = (raw_template_id if raw_template_id is not None else raw_id or "").strip()
    if not template_id:
        raise CustomError(ErrorCode.PARAM_ERROR)
    return template_id


def _report_template_output_name(name: str, *, suffix: str = "") -> str:
    clean_name = re.sub(r'[\\/:*?"<>|]+', "_", name).strip().strip(".") or "报告"
    stem = clean_name[:-5] if clean_name.lower().endswith(".docx") else clean_name
    stem_byte_limit = max(1, 250 - len(suffix.encode("utf-8")))
    safe_stem = stem.encode("utf-8")[:stem_byte_limit].decode("utf-8", errors="ignore") or "report"
    return f"{safe_stem}{suffix}.docx"


def _report_template_summary(template: ReportTemplate) -> dict:
    fetch_config = template.fetch_config if isinstance(template.fetch_config, dict) else {}
    annotations = template.annotations if isinstance(template.annotations, list) else []
    return {
        "id": template.id,
        "ctime": template.json["ctime"],
        "utime": template.json["utime"],
        "name": template.name,
        "description": template.description,
        "source_template_key": template.source_template_key,
        "placeholder_count": len(fetch_config),
        "annotation_count": len(annotations),
    }


async def _get_owned_report_template(template_id: str, tenant_id: str, owner_id: str) -> ReportTemplate:
    template = await ReportTemplate.get_or_none(
        id=template_id,
        tenant_id=tenant_id,
        owner_id=owner_id,
        deleted=False,
    )
    if not template:
        raise CustomError(ErrorCode.NOT_FOUND)
    return template


def _template_key(filename: str) -> str:
    """去掉 .docx 后缀作为 key"""
    return Path(filename).stem


def _user_report_dir(tenant_id: str, user_id: str) -> Path:
    p = Path(REPORT_USER_BASE) / tenant_id / user_id
    p.mkdir(parents=True, exist_ok=True)
    return p


def _owned_report_path(user_dir: Path, rel_path: str | None) -> Path | None:
    if not rel_path:
        return None
    root = user_dir.resolve()
    candidate = (user_dir / rel_path).resolve()
    if candidate == root or root not in candidate.parents:
        return None
    return candidate


def _atomic_write_report_file(path: Path, content: bytes) -> None:
    temporary = path.with_name(f".report-{uuid.uuid4().hex}.tmp")
    try:
        temporary.write_bytes(content)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def _canonical_report_filename(template: ReportTemplate, user_dir: Path, current_rel_path: str | None) -> str:
    template_suffix = str(template.id).replace("-", "")[:8]
    base_name = _report_template_output_name(template.name, suffix=f"_{template_suffix}")
    if base_name == current_rel_path or not (user_dir / base_name).exists():
        return base_name
    stem = Path(base_name).stem
    index = 1
    while True:
        candidate = _report_template_output_name(stem, suffix=f"_{index}")
        if candidate == current_rel_path or not (user_dir / candidate).exists():
            return candidate
        index += 1


async def _sync_report_template_document(
    template: ReportTemplate,
    tenant_id: str,
    owner_id: str,
    *,
    canonical_bytes: bytes | None = None,
) -> ReportDocument:
    """Create or update the stable canonical document used by report schedules."""
    content = canonical_bytes if canonical_bytes is not None else _html_to_docx_bytes(template.content_html)
    fetch_config = copy.deepcopy(template.fetch_config or {})
    if not isinstance(fetch_config, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)

    document = None
    if template.document_id:
        document = await ReportDocument.get_or_none(
            id=template.document_id,
            tenant_id=tenant_id,
            owner_id=owner_id,
            deleted=False,
        )

    user_dir = _user_report_dir(tenant_id, owner_id)
    old_rel_path = document.file_path if document else None
    old_path = _owned_report_path(user_dir, old_rel_path)
    rel_path = _canonical_report_filename(template, user_dir, old_rel_path)
    document_name = _report_template_output_name(template.name)
    target_path = _owned_report_path(user_dir, rel_path)
    if target_path is None:
        raise CustomError(ErrorCode.PARAM_ERROR)

    previous_target_content = None
    if target_path.exists():
        previous_target_content = target_path.read_bytes()
    _atomic_write_report_file(target_path, content)

    document_created = False
    document_snapshot = None
    previous_template_document_id = template.document_id
    try:
        if document:
            document_snapshot = {
                "name": document.name,
                "template_key": document.template_key,
                "file_path": document.file_path,
                "fetch_config": copy.deepcopy(document.fetch_config),
            }
            document.name = document_name
            document.template_key = template.source_template_key or str(template.id)
            document.file_path = rel_path
            document.fetch_config = copy.deepcopy(fetch_config)
            await document.save(update_fields=["name", "template_key", "file_path", "fetch_config", "utime"])
        else:
            document = await ReportDocument.create(
                tenant_id=tenant_id,
                owner_id=owner_id,
                name=document_name,
                template_key=template.source_template_key or str(template.id),
                file_path=rel_path,
                fetch_config=copy.deepcopy(fetch_config),
            )
            document_created = True

        if str(template.document_id or "") != str(document.id):
            template.document_id = document.id
            await template.save(update_fields=["document_id", "utime"])
    except Exception:
        if previous_target_content is None:
            target_path.unlink(missing_ok=True)
        else:
            try:
                _atomic_write_report_file(target_path, previous_target_content)
            except Exception as restore_exc:
                logger.exception("恢复报告模板 canonical 文件失败: %s", restore_exc)
        if document_created and document:
            try:
                await document.delete()
            except Exception as cleanup_exc:
                logger.exception("清理未关联的报告文档失败: %s", cleanup_exc)
        elif document and document_snapshot:
            for field, value in document_snapshot.items():
                setattr(document, field, value)
            try:
                await document.save(update_fields=["name", "template_key", "file_path", "fetch_config", "utime"])
            except Exception as restore_exc:
                logger.exception("恢复报告文档记录失败: %s", restore_exc)
        template.document_id = previous_template_document_id
        raise

    if old_path and old_path != target_path:
        try:
            old_path.unlink(missing_ok=True)
        except OSError as exc:
            logger.warning("清理已重命名的旧报告模板文件失败 path=%s error=%s", old_path, exc)
    return document


@agent_bp.post("/report/template/list")
@openapi.definition(
    summary="分页查询报告模板",
    description="查询当前用户管理的在线报告模板。需要 agent:r 权限。",
    body=_json_body({
        "type": "object",
        "properties": {
            "page": {"type": "integer", "minimum": 1},
            "size": {"type": "integer", "minimum": 1, "maximum": 100},
            "name": {"type": "string", "description": "模板名称模糊搜索"},
        },
    }),
    response=_response("模板列表", {
        "type": "object",
        "properties": {
            "list": {"type": "array", "items": REPORT_EDITABLE_TEMPLATE_SUMMARY_SCHEMA},
            "total": {"type": "integer"},
            "page": {"type": "integer"},
            "size": {"type": "integer"},
        },
    }),
)
@requires_permission("agent:r")
async def report_template_list(request: Request):
    data = _require_json(request)
    if not isinstance(data, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)
    page = data.get("page", 1)
    size = data.get("size", 20)
    name = data.get("name", "")
    if isinstance(page, bool) or not isinstance(page, int) or page < 1:
        raise CustomError(ErrorCode.PARAM_ERROR)
    if isinstance(size, bool) or not isinstance(size, int) or size < 1 or size > 100:
        raise CustomError(ErrorCode.PARAM_ERROR)
    if not isinstance(name, str):
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    if not user_id:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)
    condition = Q(tenant_id=tenant_id, owner_id=user_id, deleted=False)
    if name.strip():
        condition &= Q(name__icontains=name.strip())
    query = ReportTemplate.filter(condition)
    total = await query.count()
    templates = await query.order_by("-utime").offset((page - 1) * size).limit(size)
    return response_ok({
        "list": [_report_template_summary(template) for template in templates],
        "total": total,
        "page": page,
        "size": size,
    })


@agent_bp.post("/report/template/detail")
@openapi.definition(
    summary="查询报告模板详情",
    description="查询当前用户的一份在线报告模板。需要 agent:r 权限。",
    body=_json_body({
        "type": "object",
        "anyOf": [{"required": ["template_id"]}, {"required": ["id"]}],
        "properties": {
            "template_id": {"type": "string", "description": "模板ID"},
            "id": {"type": "string", "description": "兼容字段，与 template_id 等价"},
        },
    }),
    response=_response("模板详情", REPORT_EDITABLE_TEMPLATE_SCHEMA, {"404": "模板不存在"}),
)
@requires_permission("agent:r")
async def report_template_detail(request: Request):
    data = _require_json(request)
    template_id = _report_template_request_id(data)
    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    template = await _get_owned_report_template(template_id, tenant_id, user_id)
    return response_ok(template.json)


@agent_bp.post("/report/template/save")
@openapi.definition(
    summary="保存报告模板",
    description="无 id 时新增，有 id 时更新当前用户的模板。根据操作校验 agent:c 或 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["name", "content_html"],
        "properties": {
            "id": {"type": "string", "nullable": True},
            "name": {"type": "string"},
            "description": {"type": "string"},
            "content_html": {"type": "string"},
            "fetch_config": {"type": "object"},
            "annotations": {"type": "array"},
            "source_template_key": {"type": "string", "nullable": True},
        },
    }),
    response=_response("保存成功", REPORT_EDITABLE_TEMPLATE_SCHEMA),
)
@requires_permission("agent:c", "agent:u")
async def report_template_save(request: Request):
    data = _require_json(request)
    if not isinstance(data, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)
    raw_id = data.get("id")
    if raw_id is not None and not isinstance(raw_id, str):
        raise CustomError(ErrorCode.PARAM_ERROR)
    template_id = raw_id.strip() if isinstance(raw_id, str) else ""

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    if not user_id:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)
    if template_id:
        await check_permissions(request, "agent:u")
        template = await _get_owned_report_template(template_id, tenant_id, user_id)
        payload = _report_template_payload(data, partial=True)
        if not payload:
            raise CustomError(ErrorCode.PARAM_ERROR)
        snapshot = {field: copy.deepcopy(getattr(template, field)) for field in payload}
        for field, value in payload.items():
            setattr(template, field, value)
        canonical_bytes = _html_to_docx_bytes(template.content_html)
        template_saved = False
        try:
            await template.save()
            template_saved = True
            await _sync_report_template_document(
                template,
                tenant_id,
                user_id,
                canonical_bytes=canonical_bytes,
            )
        except Exception:
            if template_saved:
                for field, value in snapshot.items():
                    setattr(template, field, value)
                try:
                    await template.save()
                except Exception as restore_exc:
                    logger.exception("恢复报告模板记录失败: %s", restore_exc)
            raise
    else:
        await check_permissions(request, "agent:c")
        payload = _report_template_payload(data, partial=False)
        canonical_bytes = _html_to_docx_bytes(payload["content_html"])
        template = await ReportTemplate.create(
            tenant_id=tenant_id,
            owner_id=user_id,
            **payload,
        )
        try:
            await _sync_report_template_document(
                template,
                tenant_id,
                user_id,
                canonical_bytes=canonical_bytes,
            )
        except Exception:
            try:
                await template.delete()
            except Exception as cleanup_exc:
                logger.exception("清理未完成保存的报告模板失败: %s", cleanup_exc)
            raise
    return response_ok(template.json)


@agent_bp.post("/report/template/delete")
@openapi.definition(
    summary="删除报告模板",
    description="软删除当前用户的一份在线报告模板。需要 agent:d 权限。",
    body=_json_body({
        "type": "object",
        "anyOf": [{"required": ["template_id"]}, {"required": ["id"]}],
        "properties": {
            "template_id": {"type": "string", "description": "模板ID"},
            "id": {"type": "string", "description": "兼容字段，与 template_id 等价"},
        },
    }),
    response=_response("删除成功", {"type": "object", "properties": {"id": {"type": "string"}}}),
)
@requires_permission("agent:d")
async def report_template_delete(request: Request):
    data = _require_json(request)
    template_id = _report_template_request_id(data)
    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    template = await _get_owned_report_template(template_id, tenant_id, user_id)
    template.deleted = True
    await template.save(update_fields=["deleted", "utime"])
    return response_ok({"id": str(template.id)})


@agent_bp.post("/report/template/generate")
@openapi.definition(
    summary="从在线模板生成报告",
    description="执行模板中的取数占位符并生成用户 DOCX 文档。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "anyOf": [{"required": ["template_id"]}, {"required": ["id"]}],
        "properties": {
            "template_id": {"type": "string", "description": "模板ID"},
            "id": {"type": "string", "description": "兼容字段，与 template_id 等价"},
            "name": {"type": "string", "description": "可选的生成文件名"},
        },
    }),
    response=_response("生成成功", {
        "type": "object",
        "properties": {
            "doc": REPORT_DOC_SCHEMA,
            "content": {"type": "string", "description": "生成后 DOCX 的 base64 内容"},
        },
    }),
)
@requires_permission("agent:u")
async def report_template_generate(request: Request):
    data = _require_json(request)
    template_id = _report_template_request_id(data)
    output_name = data.get("name")
    if output_name is not None and (not isinstance(output_name, str) or not output_name.strip()):
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    if not user_id:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)
    template = await _get_owned_report_template(template_id, tenant_id, user_id)
    fetch_config = copy.deepcopy(template.fetch_config or {})
    if not isinstance(fetch_config, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)

    document = await _sync_report_template_document(template, tenant_id, user_id)
    annotations = template.annotations if isinstance(template.annotations, list) else []
    agent = await _get_visible_agent(tenant_id, agent_key="report_assistant")
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)
    run_id = gen_run_id()
    generated = await render_report_docx(
        content_html=template.content_html or "",
        fetch_config=fetch_config,
        annotations=annotations,
        tenant_id=tenant_id,
        operator_id=user_id,
        workspace_id=workspace_id_for(tenant_id, agent.agent_key),
        session_id=run_id,
    )
    return response_ok({
        "doc": document.json,
        "content": base64.b64encode(generated).decode("utf-8"),
    })


@agent_bp.post("/report/templates")
@openapi.definition(
    summary="获取内置报告模板列表",
    description="扫描 resources/report-templates/ 目录，返回所有 .docx 模板。需要 agent:r 权限。",
    response=_response("模板列表", {"type": "array", "items": REPORT_TEMPLATE_SCHEMA}),
)
@requires_permission("agent:r")
async def report_templates(request: Request):
    base = Path(REPORT_TEMPLATE_BASE)
    templates = []
    if base.exists():
        for f in sorted(base.iterdir()):
            if f.is_file() and f.suffix.lower() == ".docx":
                templates.append({
                    "key": _template_key(f.name),
                    "name": _template_key(f.name),
                    "filename": f.name,
                })
    return response_ok(templates)


@agent_bp.post("/report/template/content")
@openapi.definition(
    summary="获取内置报告模板内容（base64）",
    description="按 template_key 返回模板 .docx 文件的 base64 编码内容。需要 agent:r 权限。",
    body=_json_body({
        "type": "object",
        "required": ["template_key"],
        "properties": {
            "template_key": {"type": "string", "description": "模板标识"},
        },
    }),
    response=_response("模板内容", {
        "type": "object",
        "properties": {
            "template_key": {"type": "string"},
            "filename": {"type": "string"},
            "content": {"type": "string", "description": "base64 编码的 .docx 文件内容"},
        },
    }, {"404": "模板不存在"}),
)
@requires_permission("agent:r")
async def report_template_content(request: Request):
    data = _require_json(request)
    key = str(data.get("template_key") or "").strip()
    if not key:
        raise CustomError(ErrorCode.PARAM_ERROR)

    path = Path(REPORT_TEMPLATE_BASE) / f"{key}.docx"
    if not path.exists() or not path.is_file():
        raise CustomError(ErrorCode.NOT_FOUND)

    content = base64.b64encode(path.read_bytes()).decode("utf-8")
    return response_ok({"template_key": key, "filename": path.name, "content": content})


@agent_bp.post("/report/docx/to-html")
@openapi.definition(
    summary="将 DOCX 转换为富文本 HTML",
    description=(
        "将 base64 编码的 .docx 正文转换为汇报编辑器支持的安全 HTML。"
        "保留标题、字符与段落格式、列表、链接、表格和图片。需要 agent:r 权限。"
    ),
    body=_json_body({
        "type": "object",
        "required": ["content"],
        "properties": {
            "content": {"type": "string", "description": "严格 base64 编码的 .docx 文件内容"},
        },
    }),
    response=_response("转换成功", {
        "type": "object",
        "properties": {"html": {"type": "string", "description": "安全的富文本 HTML"}},
    }, {"400": "DOCX 内容或大小不符合要求"}),
)
@requires_permission("agent:r")
async def report_docx_to_html(request: Request):
    data = _require_json(request)
    if not isinstance(data, dict) or not isinstance(data.get("content"), str):
        raise CustomError(ErrorCode.PARAM_ERROR)
    try:
        html = await _run_docx_conversion(docx_base64_to_html, data["content"])
    except _DocxConversionBusyError:
        raise CustomError(ErrorCode.DOCX_CONVERSION_BUSY) from None
    except DocxConversionError:
        raise CustomError(ErrorCode.PARAM_ERROR) from None
    return response_ok({"html": html})


@agent_bp.post("/report/docx/from-html")
@openapi.definition(
    summary="将富文本 HTML 转换为 DOCX",
    description=(
        "将汇报编辑器支持的安全 HTML 转换为 base64 编码的 .docx 正文。"
        "图片仅接受安全的 http/https 外链或内嵌栅格图，并拒绝脚本、危险链接"
        "和超出限制的内容。需要 agent:u 权限。"
    ),
    body=_json_body({
        "type": "object",
        "required": ["html"],
        "properties": {
            "html": {"type": "string", "description": "汇报编辑器富文本 HTML"},
        },
    }),
    response=_response("转换成功", {
        "type": "object",
        "properties": {"content": {"type": "string", "description": "base64 编码的 .docx 内容"}},
    }, {"400": "HTML 内容、图片或大小不符合要求"}),
)
@requires_permission("agent:u")
async def report_docx_from_html(request: Request):
    data = _require_json(request)
    if not isinstance(data, dict) or not isinstance(data.get("html"), str):
        raise CustomError(ErrorCode.PARAM_ERROR)
    try:
        content = await _run_docx_conversion(html_to_docx_base64, data["html"])
    except _DocxConversionBusyError:
        raise CustomError(ErrorCode.DOCX_CONVERSION_BUSY) from None
    except DocxConversionError:
        raise CustomError(ErrorCode.PARAM_ERROR) from None
    return response_ok({"content": content})


@agent_bp.post("/report/save")
@openapi.definition(
    summary="另存报告文档",
    description="将前端传来的 .docx base64 内容写入用户目录，创建 ReportDocument 记录。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["name", "content"],
        "properties": {
            "name": {"type": "string", "description": "文件名（不含 .docx 也可，自动补全）"},
            "template_key": {"type": "string", "nullable": True, "description": "来源模板标识"},
            "content": {"type": "string", "description": "base64 编码的 .docx 文件内容"},
            "doc_id": {"type": "string", "nullable": True, "description": "若传入则覆盖已有文档"},
        },
    }),
    response=_response("保存成功", REPORT_DOC_SCHEMA, {"400": "参数错误"}),
)
@requires_permission("agent:u")
async def report_save(request: Request):
    data = _require_json(request)
    name = str(data.get("name") or "").strip()
    content_b64 = str(data.get("content") or "").strip()
    if not name or not content_b64:
        raise CustomError(ErrorCode.PARAM_ERROR)

    # 解码内容
    try:
        file_bytes = base64.b64decode(content_b64)
    except Exception:
        raise CustomError(ErrorCode.PARAM_ERROR)

    if not name.lower().endswith(".docx"):
        name = name + ".docx"

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    template_key = data.get("template_key") or None
    doc_id = data.get("doc_id") or None

    user_dir = _user_report_dir(tenant_id, user_id)

    if doc_id:
        # 覆盖已有文档
        doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, deleted=False)
        if not doc:
            raise CustomError(ErrorCode.NOT_FOUND)
        old_path = user_dir / doc.file_path
        if old_path.exists():
            old_path.unlink()
        rel_path = name
        (user_dir / rel_path).write_bytes(file_bytes)
        doc.name = name
        doc.template_key = template_key
        doc.file_path = rel_path
        await doc.save(update_fields=["name", "template_key", "file_path", "utime"])
    else:
        rel_path = name
        # 同名文件加序号避免冲突
        stem = Path(name).stem
        idx = 1
        while (user_dir / rel_path).exists():
            rel_path = f"{stem}_{idx}.docx"
            idx += 1
        (user_dir / rel_path).write_bytes(file_bytes)
        doc = await ReportDocument.create(
            tenant_id=tenant_id,
            owner_id=user_id if user_id else None,
            name=rel_path,
            template_key=template_key,
            file_path=rel_path,
        )

    return response_ok(doc.json)


@agent_bp.post("/report/list")
@openapi.definition(
    summary="获取当前用户的报告文档列表",
    description="返回当前登录用户保存的所有报告文档。需要 agent:r 权限。",
    response=_response("文档列表", {"type": "array", "items": REPORT_DOC_SCHEMA}),
)
@requires_permission("agent:r")
async def report_list(request: Request):
    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    docs = await ReportDocument.filter(
        tenant_id=tenant_id, owner_id=user_id, deleted=False
    ).order_by("-utime")
    return response_ok([d.json for d in docs])


@agent_bp.post("/report/content")
@openapi.definition(
    summary="获取用户报告文档内容（base64）",
    description="按 doc_id 返回用户已保存报告的 .docx base64 内容。需要 agent:r 权限。",
    body=_json_body({
        "type": "object",
        "required": ["doc_id"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
        },
    }),
    response=_response("文档内容", {
        "type": "object",
        "properties": {
            "doc_id": {"type": "string"},
            "name": {"type": "string"},
            "content": {"type": "string", "description": "base64 编码的 .docx 文件内容"},
        },
    }, {"404": "文档不存在"}),
)
@requires_permission("agent:r")
async def report_content(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    if not doc_id:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, owner_id=user_id, deleted=False)
    if not doc:
        raise CustomError(ErrorCode.NOT_FOUND)

    user_dir = _user_report_dir(tenant_id, user_id)
    file_path = user_dir / doc.file_path
    if not file_path.exists():
        raise CustomError(ErrorCode.NOT_FOUND)

    content = base64.b64encode(file_path.read_bytes()).decode("utf-8")
    return response_ok({"doc_id": str(doc.id), "name": doc.name, "content": content})


@agent_bp.post("/report/polish")
@openapi.definition(
    summary="使用汇报助手润色报告段落",
    description="同步调用汇报助手（report_assistant）对报告段落进行润色，直接返回润色结果。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["prompt"],
        "properties": {
            "prompt": {"type": "string", "description": "报告原文段落 + 批注内容的组合文本"},
        },
    }),
    response=_response("润色结果", {
        "type": "object",
        "properties": {
            "result": {"type": "string", "description": "润色后的段落内容"},
        },
    }, {"404": "汇报助手智能体未初始化", "500": "润色执行失败"}),
)
@requires_permission("agent:u")
async def report_polish(request: Request):
    data = _require_json(request)
    prompt = str(data.get("prompt") or "").strip()
    if not prompt:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    agent = await _get_visible_agent(tenant_id, agent_key="report_assistant")
    if not agent:
        raise CustomError(ErrorCode.NOT_FOUND)

    run_id = gen_run_id()

    try:
        result = await execute_prompt(
            workspace_id=workspace_id_for(tenant_id, agent.agent_key),
            session_id=run_id,
            request_id=f"report-polish:{run_id}",
            message=prompt,
        )
    except Exception:
        raise CustomError(ErrorCode.INTERNAL_ERROR)

    return response_ok({"result": result})


# ---------------------------------------------------------------------------
# 报告导出 & 通知接口
# ---------------------------------------------------------------------------

async def _docx_to_pdf(docx_path: Path) -> bytes:
    """调用 libreoffice 将 docx 转换为 pdf，返回 pdf 字节。"""
    with tempfile.TemporaryDirectory() as tmpdir:
        proc = await asyncio.create_subprocess_exec(
            "libreoffice", "--headless", "--convert-to", "pdf",
            "--outdir", tmpdir, str(docx_path),
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
        )
        _, stderr = await proc.communicate()
        if proc.returncode != 0:
            raise RuntimeError(f"libreoffice 转换失败: {stderr.decode()}")
        pdf_path = Path(tmpdir) / (docx_path.stem + ".pdf")
        if not pdf_path.exists():
            raise RuntimeError("libreoffice 未生成 pdf 文件")
        return pdf_path.read_bytes()


async def _docx_to_pdf_bytes(docx_bytes: bytes) -> bytes:
    """将 docx 字节写入临时文件后调用 _docx_to_pdf 转换，返回 pdf 字节。"""
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(docx_bytes)
        tmp_path = Path(tmp.name)
    try:
        return await _docx_to_pdf(tmp_path)
    finally:
        tmp_path.unlink(missing_ok=True)


async def _get_owned_report_document(doc_id: str, tenant_id: str, owner_id: str) -> ReportDocument:
    doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, owner_id=owner_id, deleted=False)
    if not doc:
        raise CustomError(ErrorCode.NOT_FOUND)
    return doc


@agent_bp.post("/report/export")
@openapi.definition(
    summary="导出报告文件并上传 S3 返回下载链接",
    description="将用户保存的 canonical 报告执行取数后导出为 docx 或 pdf，上传到对象存储并返回下载链接。pdf 转换依赖服务器端 libreoffice。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["doc_id", "format"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
            "format": {"type": "string", "enum": ["docx", "pdf"], "description": "导出格式"},
        },
    }),
    response=_response("导出成功", {
        "type": "object",
        "properties": {
            "url": {"type": "string", "description": "文件下载链接"},
            "format": {"type": "string"},
            "name": {"type": "string"},
        },
    }, {"400": "参数错误", "404": "文档不存在", "500": "pdf 转换失败"}),
)
@requires_permission("agent:u")
async def report_export(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    fmt = str(data.get("format") or "docx").strip().lower()
    if not doc_id or fmt not in ("docx", "pdf"):
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await _get_owned_report_document(doc_id, tenant_id, user_id)
    return response_ok(await export_report_document(doc, tenant_id=tenant_id, owner_id=user_id, operator_id=user_id, fmt=fmt))


@agent_bp.post("/report/notify/save")
@openapi.definition(
    summary="保存报告通知配置",
    description="为报告文档配置通知动作（覆盖式）。notify_config 结构同巡检通知配置。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["doc_id", "notify_enabled", "notify_config"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
            "notify_enabled": {"type": "boolean", "description": "是否启用通知"},
            "notify_config": {"type": "object", "description": "通知动作配置"},
        },
    }),
    response=_response("保存成功", REPORT_DOC_SCHEMA, {"404": "文档不存在"}),
)
@requires_permission("agent:u")
async def report_notify_save(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    notify_enabled = data.get("notify_enabled")
    notify_config = data.get("notify_config")
    if not doc_id or notify_enabled is None or notify_config is None:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, owner_id=user_id, deleted=False)
    if not doc:
        raise CustomError(ErrorCode.NOT_FOUND)

    doc.notify_enabled = bool(notify_enabled)
    doc.notify_config = notify_config
    await doc.save(update_fields=["notify_enabled", "notify_config", "utime"])
    return response_ok(doc.json)


def _extract_report_notify_actions(config) -> list:
    """兼容多种通知配置结构，提取动作列表。"""
    if not config:
        return []
    if isinstance(config, list):
        return [item for item in config if isinstance(item, dict)]
    if isinstance(config, dict):
        actions = config.get("callback_data") or config.get("actions")
        if isinstance(actions, list):
            return [item for item in actions if isinstance(item, dict)]
        if any(k in config for k in ("func", "action", "tool_name", "app_name")):
            return [config]
    return []


@agent_bp.post("/report/notify/send")
@openapi.definition(
    summary="发送报告通知",
    description="将报告文件上传 S3 后，按通知配置发送 Kafka 动作消息。支持指定导出格式，默认 docx。需要 agent:u 权限。",
    body=_json_body({
        "type": "object",
        "required": ["doc_id"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
            "format": {"type": "string", "enum": ["docx", "pdf"], "description": "上传格式，默认 docx"},
        },
    }),
    response=_response("发送成功", {
        "type": "object",
        "properties": {
            "sent_count": {"type": "integer", "description": "已发送通知动作数量"},
            "url": {"type": "string", "description": "上传后的文件下载链接"},
        },
    }, {"400": "未启用通知或配置为空", "404": "文档不存在"}),
)
@requires_permission("agent:u")
async def report_notify_send(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    fmt = str(data.get("format") or "docx").strip().lower()
    if not doc_id or fmt not in ("docx", "pdf"):
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await _get_owned_report_document(doc_id, tenant_id, user_id)
    return response_ok(await send_report_notify_for_document(
        doc,
        tenant_id=tenant_id,
        owner_id=user_id,
        operator_id=str(_session_user_id(request) or ""),
        fmt=fmt,
        base_url=f"{request.scheme}://{request.host}",
    ))


# ---------------------------------------------------------------------------
# 取数占位符功能
# ---------------------------------------------------------------------------

async def _resolve_fetch_placeholders(docx_bytes: bytes, fetch_config: dict, tenant_id: str, operator_id) -> bytes:
    """扫描 docx 中所有 {action_key.field.path} 占位符，并发执行取数动作后替换。"""
    doc = Document(io.BytesIO(docx_bytes))

    def _iter_paragraphs():
        yield from doc.paragraphs
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    yield from cell.paragraphs

    all_text = " ".join(p.text for p in _iter_paragraphs())
    results = await resolve_fetch_values(all_text, fetch_config, tenant_id, operator_id)
    if not results:
        return docx_bytes

    replacer = build_placeholder_replacer(results)
    for para in _iter_paragraphs():
        replace_paragraph_text_matches(para, FETCH_PLACEHOLDER_RE, replacer)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


@agent_bp.post("/report/fetch/save")
@openapi.definition(
    summary="保存取数动作配置",
    description=(
        "为报告文档配置取数占位符对应的动作（覆盖式）。"
        "fetch_config 格式：{\"a\": {动作配置}}，"
        "文档中用 {a.data.userate} 形式引用。需要 agent:u 权限。"
    ),
    body=_json_body({
        "type": "object",
        "required": ["doc_id", "fetch_config"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
            "fetch_config": {
                "type": "object",
                "description": "key 为占位符前缀，value 为动作配置",
            },
        },
    }),
    response=_response("保存成功", REPORT_DOC_SCHEMA, {"404": "文档不存在"}),
)
@requires_permission("agent:u")
async def report_fetch_save(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    fetch_config = data.get("fetch_config")
    if not doc_id or not isinstance(fetch_config, dict):
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, owner_id=user_id, deleted=False)
    if not doc:
        raise CustomError(ErrorCode.NOT_FOUND)

    doc.fetch_config = fetch_config
    await doc.save(update_fields=["fetch_config", "utime"])
    return response_ok(doc.json)


@agent_bp.post("/report/fetch/run")
@openapi.definition(
    summary="手动触发取数替换",
    description=(
        "执行文档中所有取数占位符对应的动作，将结果替换后以 base64 返回。"
        "不修改服务器文件，仅返回替换结果供前端预览/另存。需要 agent:u 权限。"
    ),
    body=_json_body({
        "type": "object",
        "required": ["doc_id"],
        "properties": {
            "doc_id": {"type": "string", "description": "文档ID"},
        },
    }),
    response=_response("执行成功", {
        "type": "object",
        "properties": {
            "doc_id": {"type": "string"},
            "name": {"type": "string"},
            "content": {"type": "string", "description": "替换后的 docx base64 内容"},
        },
    }, {"404": "文档不存在"}),
)
@requires_permission("agent:u")
async def report_fetch_run(request: Request):
    data = _require_json(request)
    doc_id = str(data.get("doc_id") or "").strip()
    if not doc_id:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    doc = await ReportDocument.get_or_none(id=doc_id, tenant_id=tenant_id, owner_id=user_id, deleted=False)
    if not doc:
        raise CustomError(ErrorCode.NOT_FOUND)

    user_dir = _user_report_dir(tenant_id, user_id)
    docx_path = _owned_report_path(user_dir, doc.file_path)
    if docx_path is None or not docx_path.exists():
        raise CustomError(ErrorCode.NOT_FOUND)

    resolved = await _resolve_fetch_placeholders(
        docx_path.read_bytes(), doc.fetch_config or {}, tenant_id, _session_user_id(request)
    )
    return response_ok({
        "doc_id": str(doc.id),
        "name": doc.name,
        "content": base64.b64encode(resolved).decode("utf-8"),
    })


async def export_report_document(
    doc: ReportDocument,
    *,
    tenant_id: str,
    owner_id: str,
    operator_id: str | None,
    fmt: str = "docx",
) -> dict:
    fmt = str(fmt or "docx").strip().lower()
    if fmt not in ("docx", "pdf"):
        raise CustomError(ErrorCode.PARAM_ERROR)

    user_dir = _user_report_dir(tenant_id, owner_id)
    docx_path = _owned_report_path(user_dir, doc.file_path)
    if docx_path is None or not docx_path.exists():
        raise CustomError(ErrorCode.NOT_FOUND)

    from settings import RY_USER_BUCKET

    export_name = Path(doc.name).stem
    try:
        resolved_bytes = await _resolve_fetch_placeholders(
            docx_path.read_bytes(), doc.fetch_config or {}, tenant_id, operator_id or owner_id
        )
        if fmt == "docx":
            file_bytes = resolved_bytes
            content_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        else:
            file_bytes = await _docx_to_pdf_bytes(resolved_bytes)
            content_type = "application/pdf"
    except CustomError:
        raise
    except Exception:
        raise CustomError(ErrorCode.INTERNAL_ERROR)

    s3_name = f"{uuid.uuid4()}.{fmt}"
    url = await upload_to_s3(
        file_bytes,
        RY_USER_BUCKET,
        s3_name,
        prefix=f"reports/{tenant_id}",
        content_type=content_type,
    )
    return {"url": url, "format": fmt, "name": f"{export_name}.{fmt}"}


def _absolute_report_url(url: str, base_url: str | None = None) -> str:
    if not url:
        return ""
    if url.startswith(("http://", "https://")):
        return url
    base = (base_url or RY_SERVICE_URL or "").rstrip("/")
    return f"{base}{url}" if base else url


async def send_report_notify_for_document(
    doc: ReportDocument,
    *,
    tenant_id: str,
    owner_id: str,
    operator_id: str | None,
    fmt: str = "docx",
    base_url: str | None = None,
) -> dict:
    if not doc.notify_enabled or not doc.notify_config:
        raise CustomError(ErrorCode.PARAM_ERROR)

    export_result = await export_report_document(
        doc,
        tenant_id=tenant_id,
        owner_id=owner_id,
        operator_id=operator_id,
        fmt=fmt,
    )
    download_url = _absolute_report_url(export_result["url"], base_url)
    notify_actions = _extract_report_notify_actions(doc.notify_config)
    sent_count = 0

    for action in notify_actions:
        action_input = action.get("action_input") or {}
        if isinstance(action_input, dict):
            action_input = {**action_input, "report_url": download_url}
        else:
            action_input = {"report_url": download_url}

        params = {
            "resource": action.get("resource", ""),
            "node_name": action.get("node_name") or action.get("description", ""),
            "action": action.get("action") or action.get("func", ""),
            "action_name": action.get("action_name") or action.get("name", ""),
            "app_name": action.get("app_name") or action.get("tool_name", ""),
            "name": action.get("name") or action.get("description", ""),
            "app_ver": action.get("app_ver") or action.get("version", ""),
            "main": action.get("main") or action.get("tool_main", ""),
            "action_input": action_input,
            "retry_times": int(action.get("retry_times") or 3),
            "timeout": int(action.get("timeout") or 300),
            "executor": operator_id,
        }
        run_id = gen_run_id()
        await aproduce(TASK_TOPIC, {"run_type": "action", "run_id": run_id, "run_args": params, "tenant_id": tenant_id})
        sent_count += 1

    return {"sent_count": sent_count, "url": export_result["url"]}


def _check_internal_agent_request(request: Request) -> None:
    from settings import LION_ACTION_INTERNAL_TOKEN, LION_AGENT_INTERNAL_TOKEN

    token = request.headers.get("X-Internal-Token") or request.headers.get("Authorization", "").removeprefix("Bearer ").strip()
    allowed_tokens = {item for item in [LION_AGENT_INTERNAL_TOKEN, LION_ACTION_INTERNAL_TOKEN] if item}
    if allowed_tokens and token not in allowed_tokens:
        raise CustomError(ErrorCode.UNAUTH)


@agent_bp.post("/report/schedule/internal/run")
async def report_schedule_internal_run(request: Request):
    _check_internal_agent_request(request)
    data = _require_json(request)
    schedule_id = str(data.get("schedule_id") or "").strip()
    tenant_id = str(data.get("tenant_id") or "").strip() or None
    if not schedule_id:
        raise CustomError(ErrorCode.PARAM_ERROR)

    from services.agent.schedule_runner import execute_report_schedule

    return response_ok(await execute_report_schedule(schedule_id, tenant_id))


def _normalize_report_schedule_payload(data: dict, *, partial: bool = False) -> dict:
    payload = {}
    if "name" in data:
        name = str(data.get("name") or "").strip()
        if not name:
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["name"] = name
    elif not partial:
        raise CustomError(ErrorCode.PARAM_ERROR)

    if "cron" in data:
        cron_expr = str(data.get("cron") or "").strip()
        if not cron_expr:
            raise CustomError(ErrorCode.PARAM_ERROR)
        if not croniter.is_valid(cron_expr):
            raise CustomError(ErrorCode.REPORT_SCHEDULE_CRON_INVALID)
        payload["cron"] = cron_expr
    elif not partial:
        raise CustomError(ErrorCode.PARAM_ERROR)

    if "export_format" in data:
        export_format = str(data.get("export_format") or "docx").strip().lower()
        if export_format not in ("docx", "pdf"):
            raise CustomError(ErrorCode.PARAM_ERROR)
        payload["export_format"] = export_format
    elif not partial:
        payload["export_format"] = "docx"

    if "enabled" in data:
        payload["enabled"] = _parse_report_schedule_bool(data.get("enabled"))
    elif not partial:
        payload["enabled"] = True

    return payload


def _parse_report_schedule_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and value in (0, 1):
        return bool(value)
    if isinstance(value, str):
        normalized = value.strip().lower()
        if normalized in ("true", "1", "yes", "on"):
            return True
        if normalized in ("false", "0", "no", "off"):
            return False
    raise CustomError(ErrorCode.PARAM_ERROR)


async def _get_schedule_doc(doc_id: str, tenant_id: str, owner_id: str) -> ReportDocument:
    return await _get_owned_report_document(doc_id, tenant_id, owner_id)


async def _get_report_schedule(schedule_id: str, tenant_id: str, owner_id: str) -> ReportSchedule:
    schedule = await ReportSchedule.get_or_none(id=schedule_id, tenant_id=tenant_id, deleted=False)
    if not schedule:
        raise CustomError(ErrorCode.REPORT_SCHEDULE_NOT_FOUND)
    document = await ReportDocument.get_or_none(
        id=schedule.document_id,
        tenant_id=tenant_id,
        owner_id=owner_id,
        deleted=False,
    )
    if not document:
        raise CustomError(ErrorCode.REPORT_SCHEDULE_NOT_FOUND)
    return schedule


async def _ensure_report_schedule_name_available(
    *,
    tenant_id: str,
    document_id: str,
    name: str,
    exclude_id: str | None = None,
) -> None:
    query = ReportSchedule.filter(
        tenant_id=tenant_id,
        document_id=document_id,
        name=name,
        deleted=False,
    )
    if exclude_id:
        query = query.exclude(id=exclude_id)
    if await query.exists():
        raise CustomError(ErrorCode.REPORT_SCHEDULE_NAME_EXISTS)


@agent_bp.post("/report/schedule/list")
@requires_permission("agent_report_schedule:r")
async def report_schedule_list(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    page = max(1, int(data.get("page") or 1))
    size = max(1, min(int(data.get("size") or 20), 100))
    doc_id = str(data.get("doc_id") or "").strip()

    condition = Q(tenant_id=tenant_id, deleted=False) & Q(document__tenant_id=tenant_id) & Q(document__owner_id=user_id) & Q(document__deleted=False)
    if doc_id:
        condition &= Q(document_id=doc_id)

    query = ReportSchedule.filter(condition).prefetch_related("document", "creator", "timed")
    total = await query.count()
    schedules = await query.order_by("-utime").offset((page - 1) * size).limit(size)
    return response_ok({"list": [item.json for item in schedules], "total": total, "page": page, "size": size})


@agent_bp.post("/report/schedule/save")
@requires_permission("agent_report_schedule:c", "agent_report_schedule:u")
async def report_schedule_save(request: Request):
    data = _require_json(request)
    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    operator = request.ctx.session.get("account_info")
    if not isinstance(operator, Account):
        operator = await Account.get_or_none(id=user_id, deleted=False)
    if not operator:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)

    schedule_id = str(data.get("id") or "").strip()
    if schedule_id:
        await check_permissions(request, "agent_report_schedule:u")
        schedule = await _get_report_schedule(schedule_id, tenant_id, user_id)
        doc_id = str(data.get("doc_id") or schedule.document_id).strip()
        if doc_id != str(schedule.document_id):
            doc = await _get_schedule_doc(doc_id, tenant_id, user_id)
            schedule.document_id = doc.id
        payload = _normalize_report_schedule_payload(data, partial=True)
        if not payload:
            raise CustomError(ErrorCode.PARAM_ERROR)
        name = payload.get("name", schedule.name)
        document_id = str(schedule.document_id)
        await _ensure_report_schedule_name_available(
            tenant_id=tenant_id,
            document_id=document_id,
            name=name,
            exclude_id=schedule.id,
        )
        for field, value in payload.items():
            setattr(schedule, field, value)
        await schedule.save()
    else:
        await check_permissions(request, "agent_report_schedule:c")
        doc_id = str(data.get("doc_id") or "").strip()
        if not doc_id:
            raise CustomError(ErrorCode.PARAM_ERROR)
        doc = await _get_schedule_doc(doc_id, tenant_id, user_id)
        payload = _normalize_report_schedule_payload(data)
        await _ensure_report_schedule_name_available(
            tenant_id=tenant_id,
            document_id=doc.id,
            name=payload["name"],
        )
        schedule = await ReportSchedule.create(
            tenant_id=tenant_id,
            document_id=doc.id,
            creator=operator,
            **payload,
        )

    sync_result = await sync_report_schedule_timer(schedule, operator)
    return response_ok({
        "id": schedule.id,
        "timed_id": str(sync_result["timed"].id) if sync_result.get("timed") else None,
        "timer_published": sync_result.get("published", False),
        "timer_publish_supported": sync_result.get("publish_supported", False),
    })


@agent_bp.post("/report/schedule/toggle")
@requires_permission("agent_report_schedule:u")
async def report_schedule_toggle(request: Request):
    data = _require_json(request)
    schedule_id = str(data.get("id") or "").strip()
    if not schedule_id or "enabled" not in data:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    operator = request.ctx.session.get("account_info")
    if not isinstance(operator, Account):
        operator = await Account.get_or_none(id=user_id, deleted=False)
    if not operator:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)

    schedule = await _get_report_schedule(schedule_id, tenant_id, user_id)
    schedule.enabled = _parse_report_schedule_bool(data.get("enabled"))
    await schedule.save(update_fields=["enabled", "utime"])
    sync_result = await sync_report_schedule_timer(schedule, operator)
    return response_ok({
        "id": schedule.id,
        "enabled": schedule.enabled,
        "timed_id": str(sync_result["timed"].id) if sync_result.get("timed") else None,
        "timer_published": sync_result.get("published", False),
        "timer_publish_supported": sync_result.get("publish_supported", False),
    })


@agent_bp.post("/report/schedule/delete")
@requires_permission("agent_report_schedule:d")
async def report_schedule_delete(request: Request):
    data = _require_json(request)
    schedule_id = str(data.get("id") or "").strip()
    if not schedule_id:
        raise CustomError(ErrorCode.PARAM_ERROR)

    tenant_id = request.ctx.session.get("tenant_id")
    user_id = str(_session_user_id(request) or "")
    operator = request.ctx.session.get("account_info")
    if not isinstance(operator, Account):
        operator = await Account.get_or_none(id=user_id, deleted=False)
    if not operator:
        raise CustomError(ErrorCode.ACCOUNT_NOT_FOUND)

    schedule = await _get_report_schedule(schedule_id, tenant_id, user_id)
    delete_result = await delete_report_schedule_timer(schedule, operator)
    schedule.enabled = False
    schedule.deleted = True
    await schedule.save(update_fields=["enabled", "deleted", "utime"])
    return response_ok({
        "id": schedule.id,
        "deleted": True,
        "timer_published": delete_result.get("published", False),
        "timer_publish_supported": delete_result.get("publish_supported", False),
    })
