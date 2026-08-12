# coding=utf-8
from __future__ import annotations

import base64
import binascii
import io
import logging
import re
import time
import zipfile
from dataclasses import dataclass, field, replace
from typing import Callable, Iterable
from urllib.parse import urlparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree, html as lxml_html
from PIL import Image, UnidentifiedImageError


MAX_DOCX_BYTES = 16 * 1024 * 1024
MAX_DOCX_UNCOMPRESSED_BYTES = 96 * 1024 * 1024
MAX_DOCX_PARTS = 4096
MAX_HTML_BYTES = 24 * 1024 * 1024
MAX_HTML_ELEMENTS = 20_000
MAX_IMAGE_BYTES = 10 * 1024 * 1024
MAX_TOTAL_IMAGE_BYTES = 20 * 1024 * 1024
MAX_IMAGE_PIXELS = 40_000_000
MAX_TABLE_ROWS = 500
MAX_TABLE_COLUMNS = 100
MAX_TOTAL_TABLE_CELLS = 10_000
MAX_CONVERSION_SECONDS = 45.0
MAX_DOCX_XML_NODES = 100_000
MAX_DOCX_PARAGRAPHS = 20_000
MAX_BASE64_CHARS = ((MAX_DOCX_BYTES + 2) // 3) * 4

_INLINE_CODE_STYLE = "AIOE Inline Code"
_CODE_BLOCK_STYLE = "AIOE Code Block"
_HORIZONTAL_RULE_STYLE = "AIOE Horizontal Rule"
_TASK_CHECKED = "\u2611"
_TASK_UNCHECKED = "\u2610"
_EXTERNAL_IMAGE_PLACEHOLDER = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR4nGNgYGD4DwABBAEAHnOcQAAAAABJRU5ErkJggg=="
)

_DRAWINGML_BLIP_TAG = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
_VML_IMAGEDATA_TAG = "{urn:schemas-microsoft-com:vml}imagedata"
_MC_ALTERNATE_CONTENT_TAG = (
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
)
_MC_CHOICE_TAG = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Choice"
_MC_FALLBACK_TAG = (
    "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"
)
_IMAGE_CONTAINER_TAGS = {
    qn("wp:inline"),
    qn("wp:anchor"),
    "{urn:schemas-microsoft-com:vml}shape",
}

_DOCX_REQUIRED_PARTS = {"[Content_Types].xml", "word/document.xml"}
_BLOCK_TAGS = {
    "p",
    "div",
    "blockquote",
    "h1",
    "h2",
    "h3",
    "h4",
    "h5",
    "h6",
    "ul",
    "ol",
    "pre",
    "table",
    "hr",
}
_ALLOWED_TAGS = _BLOCK_TAGS | {
    "a",
    "b",
    "br",
    "code",
    "col",
    "colgroup",
    "del",
    "em",
    "i",
    "img",
    "input",
    "li",
    "label",
    "mark",
    "s",
    "span",
    "strike",
    "strong",
    "sub",
    "sup",
    "tbody",
    "td",
    "tfoot",
    "th",
    "thead",
    "tr",
    "u",
}
_GLOBAL_ATTRIBUTES = {"style"}
_TAG_ATTRIBUTES = {
    "a": {"href", "title"},
    "code": {"class"},
    "img": {"src", "alt", "title", "width", "height"},
    "input": {"checked", "type"},
    "li": {"data-checked", "data-type"},
    "ol": {"start", "type", "data-type"},
    "td": {"colspan", "rowspan"},
    "th": {"colspan", "rowspan"},
    "ul": {"data-type"},
}
_ALLOWED_STYLE_PROPERTIES = {
    "background-color",
    "border",
    "color",
    "font-family",
    "font-size",
    "font-style",
    "font-weight",
    "height",
    "line-height",
    "margin-bottom",
    "margin-left",
    "margin-top",
    "text-align",
    "text-decoration",
    "text-decoration-line",
    "text-indent",
    "vertical-align",
    "width",
}
_DATA_IMAGE_RE = re.compile(
    r"\Adata:(?P<mime>image/(?:png|jpe?g|gif|bmp|tiff|webp));base64,(?P<data>[A-Za-z0-9+/]*={0,2})\Z",
    re.IGNORECASE,
)
_RGB_RE = re.compile(
    r"\Argb\(\s*(\d{1,3})\s*,\s*(\d{1,3})\s*,\s*(\d{1,3})\s*\)\Z", re.IGNORECASE
)
_HEX_COLOR_RE = re.compile(r"\A#(?:[0-9a-fA-F]{3}|[0-9a-fA-F]{6})\Z")
_FONT_NAME_RE = re.compile(r"\A[\w \-.,'\u0080-\uffff]+\Z")
_LENGTH_RE = re.compile(r"\A(\d+(?:\.\d+)?)(px|pt|in|cm|mm)?\Z", re.IGNORECASE)
_SIGNED_LENGTH_RE = re.compile(r"\A(-?\d+(?:\.\d+)?)(px|pt|in|cm|mm)?\Z", re.IGNORECASE)
_BORDER_RE = re.compile(
    r"\A(\d+(?:\.\d+)?(?:px|pt|in|cm|mm)?)\s+"
    r"(solid|dashed|dotted|double)\s+(.+)\Z",
    re.IGNORECASE,
)

_NAMED_COLORS = {
    "black": "#000000",
    "blue": "#0000FF",
    "cyan": "#00FFFF",
    "gray": "#808080",
    "green": "#008000",
    "lime": "#00FF00",
    "magenta": "#FF00FF",
    "maroon": "#800000",
    "navy": "#000080",
    "olive": "#808000",
    "orange": "#FFA500",
    "purple": "#800080",
    "red": "#FF0000",
    "silver": "#C0C0C0",
    "teal": "#008080",
    "white": "#FFFFFF",
    "yellow": "#FFFF00",
}
_HIGHLIGHT_TO_HEX = {
    WD_COLOR_INDEX.BLACK: "#000000",
    WD_COLOR_INDEX.BLUE: "#0000FF",
    WD_COLOR_INDEX.BRIGHT_GREEN: "#00FF00",
    WD_COLOR_INDEX.DARK_BLUE: "#000080",
    WD_COLOR_INDEX.DARK_RED: "#800000",
    WD_COLOR_INDEX.DARK_YELLOW: "#808000",
    WD_COLOR_INDEX.GRAY_25: "#C0C0C0",
    WD_COLOR_INDEX.GRAY_50: "#808080",
    WD_COLOR_INDEX.GREEN: "#008000",
    WD_COLOR_INDEX.PINK: "#FF00FF",
    WD_COLOR_INDEX.RED: "#FF0000",
    WD_COLOR_INDEX.TEAL: "#008080",
    WD_COLOR_INDEX.TURQUOISE: "#00FFFF",
    WD_COLOR_INDEX.VIOLET: "#800080",
    WD_COLOR_INDEX.WHITE: "#FFFFFF",
    WD_COLOR_INDEX.YELLOW: "#FFFF00",
}
_IMAGE_FORMAT_MIME = {
    "BMP": "image/bmp",
    "GIF": "image/gif",
    "JPEG": "image/jpeg",
    "PNG": "image/png",
    "TIFF": "image/tiff",
    "WEBP": "image/webp",
}

logger = logging.getLogger(__name__)


class DocxConversionError(ValueError):
    """Raised when conversion input is invalid or outside the supported contract."""


@dataclass(frozen=True)
class _InlineStyle:
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    superscript: bool = False
    subscript: bool = False
    font_name: str | None = None
    font_size_pt: float | None = None
    color: str | None = None
    highlight: str | None = None
    code: bool = False


@dataclass(frozen=True)
class _FontProperties:
    bold: object = None
    italic: object = None
    underline: object = None
    strike: object = None
    superscript: object = None
    subscript: object = None
    font_name: str | None = None
    font_size: object = None
    color: str | None = None
    highlight: str | None = None


@dataclass
class _ImageBudget:
    total_bytes: int = 0
    started_at: float = field(default_factory=time.monotonic)
    checks: int = 0
    style_cache: dict[tuple[str, str | None], object] = field(default_factory=dict)
    style_chain_cache: dict[tuple[object, str], tuple] = field(default_factory=dict)
    style_font_cache: dict[tuple[object, str], _FontProperties] = field(
        default_factory=dict
    )
    paragraph_font_cache: dict[int, _FontProperties] = field(default_factory=dict)

    def check(self, *, force: bool = False) -> None:
        self.checks += 1
        if (force or self.checks % 64 == 0) and (
            time.monotonic() - self.started_at > MAX_CONVERSION_SECONDS
        ):
            raise DocxConversionError("document conversion timed out")

    def add(self, size: int) -> None:
        self.check()
        if size > MAX_IMAGE_BYTES or self.total_bytes + size > MAX_TOTAL_IMAGE_BYTES:
            raise DocxConversionError("embedded image exceeds size limit")
        self.total_bytes += size


@dataclass
class _ListContext:
    level: int
    kind: str
    element: etree._Element
    num_id: int | None = None
    ordered_type: str | None = None
    last_item: etree._Element | None = None


@dataclass(frozen=True)
class _ListInfo:
    kind: str
    level: int
    num_id: int | None = None
    ordered_type: str | None = None
    start: int | None = None


def _fail(message: str) -> DocxConversionError:
    return DocxConversionError(message)


def decode_docx_base64(content: str, *, _budget: _ImageBudget | None = None) -> bytes:
    if not isinstance(content, str) or not content or len(content) > MAX_BASE64_CHARS:
        raise _fail("invalid DOCX base64 content")
    try:
        encoded = content.encode("ascii")
        docx_bytes = base64.b64decode(encoded, validate=True)
    except (UnicodeEncodeError, binascii.Error, ValueError):
        raise _fail("invalid DOCX base64 content") from None
    validate_docx_bytes(docx_bytes, _budget=_budget)
    return docx_bytes


def validate_docx_bytes(
    docx_bytes: bytes, *, _budget: _ImageBudget | None = None
) -> None:
    if (
        not isinstance(docx_bytes, bytes)
        or not docx_bytes
        or len(docx_bytes) > MAX_DOCX_BYTES
    ):
        raise _fail("invalid DOCX size")
    try:
        with zipfile.ZipFile(io.BytesIO(docx_bytes)) as archive:
            infos = archive.infolist()
            names = {info.filename for info in infos}
            if len(infos) > MAX_DOCX_PARTS or not _DOCX_REQUIRED_PARTS.issubset(names):
                raise _fail("invalid DOCX package")
            total_size = 0
            for info in infos:
                if _budget:
                    _budget.check()
                if info.file_size < 0 or info.file_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise _fail("invalid DOCX package")
                total_size += info.file_size
                if total_size > MAX_DOCX_UNCOMPRESSED_BYTES:
                    raise _fail("DOCX package exceeds uncompressed size limit")
            if archive.testzip() is not None:
                raise _fail("corrupt DOCX package")
            content_types = archive.read("[Content_Types].xml")
            parser = etree.XMLParser(
                resolve_entities=False, no_network=True, recover=False, huge_tree=False
            )
            etree.fromstring(content_types, parser=parser)
            document_xml = archive.read("word/document.xml")
            if re.search(rb"<!\s*(?:DOCTYPE|ENTITY)", document_xml, re.IGNORECASE):
                raise _fail("unsafe DOCX document XML")
            node_count = 0
            paragraph_count = 0
            for _, element in etree.iterparse(
                io.BytesIO(document_xml),
                events=("end",),
                resolve_entities=False,
                no_network=True,
                recover=False,
                huge_tree=False,
            ):
                node_count += 1
                if element.tag == qn("w:p"):
                    paragraph_count += 1
                if (
                    node_count > MAX_DOCX_XML_NODES
                    or paragraph_count > MAX_DOCX_PARAGRAPHS
                ):
                    raise _fail("DOCX document structure limit exceeded")
                if _budget:
                    _budget.check()
                element.clear()
        Document(io.BytesIO(docx_bytes))
        if _budget:
            _budget.check(force=True)
    except DocxConversionError:
        raise
    except (zipfile.BadZipFile, KeyError, ValueError, etree.XMLSyntaxError, OSError):
        raise _fail("invalid DOCX package") from None
    except Exception:
        raise _fail("unsupported DOCX package") from None


def docx_base64_to_html(content: str) -> str:
    budget = _ImageBudget()
    return docx_to_html(
        decode_docx_base64(content, _budget=budget),
        _budget=budget,
        _validated=True,
    )


def html_to_docx_base64(html_content: str) -> str:
    return base64.b64encode(html_to_docx(html_content)).decode("ascii")


def _canonical_color(value: str) -> str | None:
    value = value.strip()
    if not value or value.lower() == "transparent":
        return None
    named = _NAMED_COLORS.get(value.lower())
    if named:
        return named
    if _HEX_COLOR_RE.fullmatch(value):
        value = value[1:]
        if len(value) == 3:
            value = "".join(char * 2 for char in value)
        return f"#{value.upper()}"
    match = _RGB_RE.fullmatch(value)
    if match:
        channels = [int(item) for item in match.groups()]
        if all(0 <= item <= 255 for item in channels):
            return "#" + "".join(f"{item:02X}" for item in channels)
    raise _fail("invalid CSS color")


def _canonical_font_name(value: str) -> str:
    value = value.strip().strip("\"'").split(",", 1)[0].strip().strip("\"'")
    if not value or len(value) > 128 or not _FONT_NAME_RE.fullmatch(value):
        raise _fail("invalid font family")
    return value


def _length_to_points(
    value: str,
    *,
    minimum: float = 1,
    maximum: float = 2000,
    signed: bool = False,
) -> float:
    match = (_SIGNED_LENGTH_RE if signed else _LENGTH_RE).fullmatch(value.strip())
    if not match:
        raise _fail("invalid CSS length")
    amount = float(match.group(1))
    unit = (match.group(2) or "px").lower()
    factors = {"px": 0.75, "pt": 1.0, "in": 72.0, "cm": 72.0 / 2.54, "mm": 72.0 / 25.4}
    points = amount * factors[unit]
    if not minimum <= points <= maximum:
        raise _fail("CSS length outside supported range")
    return points


def _canonical_style_value(name: str, value: str) -> str | None:
    lowered = value.lower()
    if name in {"color", "background-color"}:
        if lowered in {"inherit", "initial", "unset"}:
            return None
        return _canonical_color(value)
    if name == "font-family":
        return _canonical_font_name(value)
    if name == "font-size":
        return f"{_length_to_points(value, minimum=6, maximum=96):g}pt"
    if name == "font-style":
        if lowered not in {"normal", "italic", "oblique"}:
            raise _fail("invalid font style")
        return lowered
    if name == "font-weight":
        if lowered not in {
            "normal",
            "bold",
            "bolder",
            "lighter",
        } and not re.fullmatch(r"[1-9]00", lowered):
            raise _fail("invalid font weight")
        return lowered
    if name in {"text-decoration", "text-decoration-line"}:
        tokens = set(lowered.split())
        if not tokens or not tokens <= {"none", "underline", "line-through"}:
            raise _fail("invalid text decoration")
        return " ".join(sorted(tokens))
    if name == "vertical-align":
        if lowered not in {"baseline", "sub", "super"}:
            raise _fail("invalid vertical alignment")
        return lowered
    if name == "text-align":
        if lowered not in {"left", "right", "center", "justify", "start", "end"}:
            raise _fail("invalid text alignment")
        return {"start": "left", "end": "right"}.get(lowered, lowered)
    if name == "line-height":
        if re.fullmatch(r"\d+(?:\.\d+)?", lowered):
            amount = float(lowered)
            if not 0.5 <= amount <= 5:
                raise _fail("invalid line height")
            return f"{amount:g}"
        if lowered.endswith("%"):
            try:
                amount = float(lowered[:-1]) / 100
            except ValueError:
                raise _fail("invalid line height") from None
            if not 0.5 <= amount <= 5:
                raise _fail("invalid line height")
            return f"{amount:g}"
        return f"{_length_to_points(value, minimum=6, maximum=240):g}pt"
    if name in {"margin-left", "text-indent"}:
        points = _length_to_points(value, minimum=-2000, maximum=2000, signed=True)
        return f"{points:g}pt"
    if name in {"margin-top", "margin-bottom"}:
        return f"{_length_to_points(value, minimum=0, maximum=2000):g}pt"
    if name == "border":
        match = _BORDER_RE.fullmatch(value.strip())
        if not match:
            raise _fail("invalid CSS border")
        width = _length_to_points(match.group(1), minimum=0.125, maximum=12)
        color = _canonical_color(match.group(3))
        if color is None:
            raise _fail("invalid CSS border color")
        return f"{width:g}pt {match.group(2).lower()} {color}"
    if name in {"width", "height"}:
        return f"{_length_to_points(value, minimum=1, maximum=2000):g}pt"
    return None


def _canonical_style(style: str) -> str:
    if not isinstance(style, str) or len(style) > 4096:
        raise _fail("invalid style attribute")
    declarations: dict[str, str] = {}
    for declaration in style.split(";"):
        if not declaration.strip():
            continue
        if ":" not in declaration:
            raise _fail("invalid style declaration")
        name, value = declaration.split(":", 1)
        name = name.strip().lower()
        value = value.strip()
        if name not in _ALLOWED_STYLE_PROPERTIES:
            continue
        lowered = value.lower()
        if any(
            token in lowered
            for token in ("url(", "expression(", "javascript:", "@import", "var(")
        ):
            raise _fail("unsafe style value")
        try:
            canonical = _canonical_style_value(name, value)
        except DocxConversionError as exc:
            logger.warning(
                "Dropping invalid style declaration %s=%r: %s", name, value, exc
            )
            continue
        if canonical is not None:
            declarations[name] = canonical
    return "; ".join(f"{name}: {value}" for name, value in declarations.items())


def _safe_network_url(value: str) -> str:
    value = value.strip()
    if (
        not value
        or len(value) > 2048
        or any(char.isspace() or char == "\x00" for char in value)
    ):
        raise _fail("invalid network URL")
    if value.startswith("//"):
        value = f"https:{value}"
    elif not urlparse(value).scheme:
        value = f"https://{value}"
    try:
        parsed = urlparse(value)
        hostname = parsed.hostname
    except ValueError:
        raise _fail("invalid network URL") from None
    if (
        parsed.scheme.lower() not in {"http", "https"}
        or not parsed.netloc
        or not hostname
    ):
        raise _fail("unsupported network URL")
    return value


def _safe_href(value: str) -> str:
    value = value.strip()
    if not value or len(value) > 2048 or any(char in value for char in "\r\n\x00"):
        raise _fail("invalid hyperlink")
    if value.startswith("#"):
        return value
    parsed = urlparse(value)
    if not parsed.scheme or parsed.scheme.lower() in {"http", "https"}:
        return _safe_network_url(value)
    if parsed.scheme.lower() not in {"mailto", "tel"} or not parsed.path:
        raise _fail("unsupported hyperlink scheme")
    return value


def _decode_image_data_uri(
    src: str, budget: _ImageBudget | None = None
) -> tuple[bytes, str]:
    if not isinstance(src, str):
        raise _fail("invalid embedded image")
    match = _DATA_IMAGE_RE.fullmatch(src)
    if not match:
        raise _fail("only base64 data images are supported")
    try:
        image_bytes = base64.b64decode(match.group("data"), validate=True)
    except (binascii.Error, ValueError):
        raise _fail("invalid embedded image base64") from None
    if not image_bytes or len(image_bytes) > MAX_IMAGE_BYTES:
        raise _fail("embedded image exceeds size limit")
    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            image_format = (image.format or "").upper()
            width, height = image.size
            image.verify()
    except (UnidentifiedImageError, Image.DecompressionBombError, OSError, ValueError):
        raise _fail("invalid embedded image") from None
    if width <= 0 or height <= 0 or width * height > MAX_IMAGE_PIXELS:
        raise _fail("embedded image dimensions exceed limit")
    actual_mime = _IMAGE_FORMAT_MIME.get(image_format)
    declared_mime = match.group("mime").lower().replace("image/jpg", "image/jpeg")
    if not actual_mime or declared_mime != actual_mime:
        raise _fail("embedded image type mismatch")
    if image_format == "WEBP":
        with Image.open(io.BytesIO(image_bytes)) as image:
            converted = io.BytesIO()
            mode = "RGBA" if "A" in image.getbands() else "RGB"
            image.convert(mode).save(converted, format="PNG")
            image_bytes = converted.getvalue()
            actual_mime = "image/png"
    if budget:
        budget.add(len(image_bytes))
    return image_bytes, actual_mime


def _parse_html_fragment(
    html_content: str, budget: _ImageBudget | None = None
) -> etree._Element:
    budget = budget or _ImageBudget()
    budget.check(force=True)
    if not isinstance(html_content, str) or not html_content.strip():
        raise _fail("HTML content is required")
    try:
        encoded = html_content.encode("utf-8")
    except UnicodeEncodeError:
        raise _fail("invalid HTML encoding") from None
    if len(encoded) > MAX_HTML_BYTES or re.search(
        r"<!\s*(?:DOCTYPE|ENTITY)", html_content, re.IGNORECASE
    ):
        raise _fail("invalid or oversized HTML content")
    parser = lxml_html.HTMLParser(
        encoding="utf-8", recover=False, remove_comments=True, no_network=True
    )
    try:
        root = lxml_html.fragment_fromstring(
            encoded, create_parent="div", parser=parser
        )
    except (etree.ParserError, etree.XMLSyntaxError, ValueError):
        raise _fail("malformed HTML content") from None
    elements = list(root.iterdescendants())
    if len(elements) > MAX_HTML_ELEMENTS:
        raise _fail("HTML element limit exceeded")
    for element in elements:
        budget.check()
        if not isinstance(element.tag, str) or not element.tag.isascii():
            raise _fail("unsupported HTML node")
        tag = element.tag.lower()
        if tag not in _ALLOWED_TAGS:
            raise _fail(f"unsupported HTML tag: {tag}")
        cleaned: dict[str, str] = {}
        for raw_name, raw_value in element.attrib.items():
            name = raw_name.lower()
            if name.startswith("on"):
                raise _fail("event attributes are not allowed")
            if name not in _GLOBAL_ATTRIBUTES | _TAG_ATTRIBUTES.get(tag, set()):
                continue
            if name == "style":
                style = _canonical_style(raw_value)
                if style:
                    cleaned[name] = style
            elif name == "href":
                cleaned[name] = _safe_href(raw_value)
            elif name == "src":
                if raw_value.strip().lower().startswith("data:"):
                    _decode_image_data_uri(raw_value, budget)
                    cleaned[name] = raw_value
                else:
                    cleaned[name] = _safe_network_url(raw_value)
            elif name in {"colspan", "rowspan", "start"}:
                try:
                    number = int(raw_value)
                except (TypeError, ValueError):
                    raise _fail(f"invalid {name}") from None
                maximum = 100_000 if name == "start" else MAX_TABLE_COLUMNS
                if not 1 <= number <= maximum:
                    raise _fail(f"invalid {name}")
                cleaned[name] = str(number)
            elif name in {"width", "height"}:
                points = _length_to_points(raw_value, minimum=1, maximum=2000)
                cleaned[name] = (
                    f"{points * 4 / 3:.17g}" if tag == "img" else f"{points:g}pt"
                )
            elif name == "class":
                language_class = next(
                    (
                        token
                        for token in raw_value.split()
                        if re.fullmatch(r"language-[A-Za-z0-9_+.-]{1,64}", token)
                    ),
                    "",
                )
                if language_class:
                    cleaned[name] = language_class
            elif name == "data-type":
                expected = "taskList" if tag in {"ul", "ol"} else "taskItem"
                if raw_value == expected:
                    cleaned[name] = expected
            elif name == "data-checked":
                if raw_value.lower() in {"true", "false"}:
                    cleaned[name] = raw_value.lower()
            elif name == "type" and tag == "ol":
                if raw_value not in {"1", "a", "A", "i", "I"}:
                    raise _fail("unsupported ordered list type")
                cleaned[name] = raw_value
            elif name == "type":
                if raw_value.lower() != "checkbox":
                    raise _fail("unsupported input type")
                cleaned[name] = "checkbox"
            elif name == "checked":
                cleaned[name] = "checked"
            else:
                cleaned[name] = raw_value[:1024]
        if tag == "img" and "src" not in cleaned:
            raise _fail("image source is required")
        if tag == "input" and cleaned.get("type") != "checkbox":
            raise _fail("unsupported input type")
        element.attrib.clear()
        element.attrib.update(cleaned)
    _validate_html_tables(root, budget)
    return root


def _style_map(element: etree._Element) -> dict[str, str]:
    result: dict[str, str] = {}
    for declaration in element.get("style", "").split(";"):
        if ":" in declaration:
            name, value = declaration.split(":", 1)
            result[name.strip().lower()] = value.strip()
    return result


def _inline_style_for_element(
    parent: _InlineStyle, element: etree._Element
) -> _InlineStyle:
    tag = element.tag.lower()
    style = _style_map(element)
    changes: dict = {}
    if tag in {"b", "strong", "th"}:
        changes["bold"] = True
    if tag in {"em", "i"}:
        changes["italic"] = True
    if tag == "u":
        changes["underline"] = True
    if tag in {"s", "strike", "del"}:
        changes["strike"] = True
    if tag == "sup":
        changes.update(superscript=True, subscript=False)
    if tag == "sub":
        changes.update(subscript=True, superscript=False)
    if tag == "mark":
        changes["highlight"] = "#FFFF00"
    if tag == "code":
        changes.update(code=True, font_name="Consolas")

    weight = style.get("font-weight")
    if weight:
        changes["bold"] = weight in {"bold", "bolder"} or (
            weight.isdigit() and int(weight) >= 600
        )
    font_style = style.get("font-style")
    if font_style:
        changes["italic"] = font_style in {"italic", "oblique"}
    decoration = f"{style.get('text-decoration', '')} {style.get('text-decoration-line', '')}".split()
    if decoration:
        changes["underline"] = "underline" in decoration
        changes["strike"] = "line-through" in decoration
    vertical_align = style.get("vertical-align")
    if vertical_align == "super":
        changes.update(superscript=True, subscript=False)
    elif vertical_align == "sub":
        changes.update(subscript=True, superscript=False)
    elif vertical_align == "baseline":
        changes.update(subscript=False, superscript=False)
    if "font-family" in style:
        changes["font_name"] = style["font-family"]
    if "font-size" in style:
        changes["font_size_pt"] = _length_to_points(
            style["font-size"], minimum=6, maximum=96
        )
    if "color" in style:
        changes["color"] = _canonical_color(style["color"])
    if "background-color" in style and tag not in {"td", "th"}:
        changes["highlight"] = _canonical_color(style["background-color"])
    return replace(parent, **changes)


def _apply_run_style(run: Run, style: _InlineStyle) -> None:
    if style.code:
        try:
            run.style = _INLINE_CODE_STYLE
        except KeyError:
            pass
    run.bold = style.bold
    run.italic = style.italic
    run.underline = style.underline
    run.font.strike = style.strike
    run.font.superscript = style.superscript
    run.font.subscript = style.subscript
    if style.font_name:
        run.font.name = style.font_name
        r_fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
        r_fonts.set(qn("w:eastAsia"), style.font_name)
    if style.font_size_pt:
        run.font.size = Pt(style.font_size_pt)
    if style.color:
        run.font.color.rgb = RGBColor.from_string(style.color.removeprefix("#"))
    if style.highlight:
        r_pr = run._r.get_or_add_rPr()
        shading = r_pr.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            r_pr.append(shading)
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:color"), "auto")
        shading.set(qn("w:fill"), style.highlight.removeprefix("#"))


def _new_run(
    paragraph: Paragraph, text: str, parent_xml: etree._Element | None = None
) -> Run:
    if parent_xml is None:
        return paragraph.add_run(text)
    run_element = OxmlElement("w:r")
    parent_xml.append(run_element)
    run = Run(run_element, paragraph)
    run.text = text
    return run


def _append_text_run(
    paragraph: Paragraph,
    text: str | None,
    style: _InlineStyle,
    parent_xml: etree._Element | None = None,
) -> None:
    if not text:
        return
    normalized = re.sub(r"[\t\r\n ]+", " ", text)
    if not normalized:
        return
    run = _new_run(paragraph, normalized, parent_xml)
    _apply_run_style(run, style)


def _image_lengths(element: etree._Element) -> tuple[Pt | None, Pt | None]:
    style = _style_map(element)
    width_value = style.get("width") or element.get("width")
    height_value = style.get("height") or element.get("height")
    width = Pt(_length_to_points(width_value)) if width_value else None
    height = Pt(_length_to_points(height_value)) if height_value else None
    return width, height


def _append_image_run(
    paragraph: Paragraph,
    element: etree._Element,
    style: _InlineStyle,
    budget: _ImageBudget,
    parent_xml: etree._Element | None = None,
) -> None:
    src = element.get("src", "")
    width, height = _image_lengths(element)
    run = _new_run(paragraph, "", parent_xml)
    _apply_run_style(run, style)
    if src.lower().startswith("data:"):
        image_bytes, _ = _decode_image_data_uri(src, budget)
        try:
            shape = run.add_picture(io.BytesIO(image_bytes), width=width, height=height)
        except (ValueError, OSError, UnidentifiedImageError):
            raise _fail("unsupported embedded image") from None
    else:
        src = _safe_network_url(src)
        if width is None and height is None:
            width, height = Pt(320), Pt(180)
        elif width is None:
            width = Pt(height.pt * 16 / 9)
        elif height is None:
            height = Pt(width.pt * 9 / 16)
        try:
            shape = run.add_picture(
                io.BytesIO(_EXTERNAL_IMAGE_PLACEHOLDER), width=width, height=height
            )
            blip = shape._inline.xpath(".//a:blip")[0]
            embedded_relationship_id = blip.get(qn("r:embed"))
            external_relationship_id = paragraph.part.relate_to(
                src, RT.IMAGE, is_external=True
            )
            blip.attrib.pop(qn("r:embed"), None)
            blip.set(qn("r:link"), external_relationship_id)
            if embedded_relationship_id:
                paragraph.part.drop_rel(embedded_relationship_id)
        except (IndexError, KeyError, ValueError, OSError):
            raise _fail("unsupported external image") from None
    alt = (element.get("alt") or "").strip()
    if alt:
        shape._inline.docPr.set("descr", alt[:1024])


def _ensure_conversion_styles(document) -> None:
    styles = document.styles
    try:
        inline_code = styles[_INLINE_CODE_STYLE]
    except KeyError:
        inline_code = styles.add_style(_INLINE_CODE_STYLE, WD_STYLE_TYPE.CHARACTER)
    inline_code.font.name = "Consolas"
    inline_code.font.size = Pt(10)

    try:
        code_block = styles[_CODE_BLOCK_STYLE]
    except KeyError:
        code_block = styles.add_style(_CODE_BLOCK_STYLE, WD_STYLE_TYPE.PARAGRAPH)
    code_block.font.name = "Consolas"
    code_block.font.size = Pt(10)

    try:
        horizontal_rule = styles[_HORIZONTAL_RULE_STYLE]
    except KeyError:
        horizontal_rule = styles.add_style(
            _HORIZONTAL_RULE_STYLE, WD_STYLE_TYPE.PARAGRAPH
        )
    horizontal_rule.paragraph_format.space_after = Pt(0)
    horizontal_rule.paragraph_format.space_before = Pt(0)


def _append_hyperlink(
    paragraph: Paragraph,
    element: etree._Element,
    parent_style: _InlineStyle,
    budget: _ImageBudget,
) -> None:
    href = _safe_href(element.get("href", ""))
    hyperlink = OxmlElement("w:hyperlink")
    if href.startswith("#"):
        hyperlink.set(qn("w:anchor"), href[1:])
    else:
        relationship_id = paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True)
        hyperlink.set(qn("r:id"), relationship_id)
    paragraph._p.append(hyperlink)
    _append_inline_contents(paragraph, element, parent_style, budget, hyperlink)


def _append_inline_contents(
    paragraph: Paragraph,
    element: etree._Element,
    parent_style: _InlineStyle,
    budget: _ImageBudget,
    parent_xml: etree._Element | None = None,
) -> None:
    budget.check()
    style = _inline_style_for_element(parent_style, element)
    _append_text_run(paragraph, element.text, style, parent_xml)
    for child in element:
        tag = child.tag.lower()
        if tag == "br":
            run = _new_run(paragraph, "\n", parent_xml)
            _apply_run_style(run, style)
        elif tag == "img":
            _append_image_run(
                paragraph,
                child,
                _inline_style_for_element(style, child),
                budget,
                parent_xml,
            )
        elif tag == "a" and parent_xml is None:
            _append_hyperlink(paragraph, child, style, budget)
        else:
            _append_inline_contents(paragraph, child, style, budget, parent_xml)
        _append_text_run(paragraph, child.tail, style, parent_xml)


def _apply_paragraph_style(paragraph: Paragraph, element: etree._Element) -> None:
    style = _style_map(element)
    alignment = style.get("text-align")
    if alignment:
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
            "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        }[alignment]
    line_height = style.get("line-height")
    if line_height:
        if line_height.endswith("pt"):
            paragraph.paragraph_format.line_spacing = Pt(float(line_height[:-2]))
        else:
            paragraph.paragraph_format.line_spacing = float(line_height)
    paragraph_lengths = {
        "margin-left": "left_indent",
        "text-indent": "first_line_indent",
        "margin-top": "space_before",
        "margin-bottom": "space_after",
    }
    for css_name, format_name in paragraph_lengths.items():
        value = style.get(css_name)
        if value is not None:
            setattr(paragraph.paragraph_format, format_name, Pt(float(value[:-2])))
    if element.tag.lower() == "blockquote" and "margin-left" not in style:
        paragraph.paragraph_format.left_indent = Inches(0.25)


def _add_paragraph_from_element(
    container,
    element: etree._Element,
    budget: _ImageBudget,
    *,
    style_name: str | None = None,
) -> Paragraph:
    paragraph = container.add_paragraph(style=style_name)
    _apply_paragraph_style(paragraph, element)
    initial_style = _InlineStyle(bold=element.tag.lower() == "th")
    _append_inline_contents(paragraph, element, initial_style, budget)
    _trim_paragraph_spaces(paragraph)
    return paragraph


def _add_text_paragraph(container, text: str, budget: _ImageBudget) -> None:
    if not text or not text.strip():
        return
    paragraph = container.add_paragraph()
    _append_text_run(paragraph, text, _InlineStyle())


def _list_style_name(kind: str, level: int) -> str:
    base = "List Number" if kind == "ol" else "List Bullet"
    return base if level <= 0 else f"{base} {min(level + 1, 3)}"


def _append_list_item_contents(
    paragraph: Paragraph,
    item: etree._Element,
    budget: _ImageBudget,
    *,
    task_checked: bool | None = None,
) -> None:
    base_style = _inline_style_for_element(_InlineStyle(), item)
    if task_checked is not None:
        _append_text_run(
            paragraph,
            f"{_TASK_CHECKED if task_checked else _TASK_UNCHECKED} ",
            base_style,
        )
    _append_text_run(paragraph, item.text, base_style)
    has_content = bool(item.text and item.text.strip())
    for child in item:
        if child.tag.lower() in {"ul", "ol"}:
            continue
        if task_checked is not None and child.tag.lower() == "label":
            continue
        if child.tag.lower() in {"p", "div", "blockquote"} and has_content:
            run = paragraph.add_run("\n")
            _apply_run_style(run, base_style)
        _append_inline_contents(paragraph, child, base_style, budget)
        has_content = has_content or bool("".join(child.itertext()).strip())
        _append_text_run(paragraph, child.tail, base_style)
    _trim_paragraph_spaces(paragraph)


def _add_ordered_list_numbering(
    document, *, ordered_type: str | None, start: int, level: int
) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId"), "0"))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"), "0")) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    numbering_level = min(max(level, 0), 8)

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "multilevel")
    abstract.append(multi_level_type)
    level_element = OxmlElement("w:lvl")
    level_element.set(qn("w:ilvl"), str(numbering_level))
    start_element = OxmlElement("w:start")
    start_element.set(qn("w:val"), str(start))
    format_element = OxmlElement("w:numFmt")
    format_element.set(
        qn("w:val"),
        {
            "a": "lowerLetter",
            "A": "upperLetter",
            "i": "lowerRoman",
            "I": "upperRoman",
        }.get(ordered_type, "decimal"),
    )
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), f"%{numbering_level + 1}.")
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    paragraph_properties = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), str(720 * (numbering_level + 1)))
    indentation.set(qn("w:hanging"), "360")
    paragraph_properties.append(indentation)
    level_element.extend(
        [
            start_element,
            format_element,
            level_text,
            level_justification,
            paragraph_properties,
        ]
    )
    abstract.append(level_element)
    numbering.append(abstract)

    number = OxmlElement("w:num")
    number.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    number.append(abstract_reference)
    numbering.append(number)
    return num_id, numbering_level


def _apply_paragraph_numbering(
    paragraph: Paragraph, *, num_id: int, level: int
) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    existing = paragraph_properties.find(qn("w:numPr"))
    if existing is not None:
        paragraph_properties.remove(existing)
    numbering_properties = OxmlElement("w:numPr")
    level_reference = OxmlElement("w:ilvl")
    level_reference.set(qn("w:val"), str(level))
    number_reference = OxmlElement("w:numId")
    number_reference.set(qn("w:val"), str(num_id))
    numbering_properties.extend([level_reference, number_reference])
    paragraph_properties.append(numbering_properties)


def _add_html_list(
    container, element: etree._Element, budget: _ImageBudget, level: int = 0
) -> None:
    kind = element.tag.lower()
    is_task_list = element.get("data-type") == "taskList"
    numbering_info = None
    if kind == "ol" and not is_task_list:
        numbering_info = _add_ordered_list_numbering(
            container.part.document,
            ordered_type=element.get("type"),
            start=int(element.get("start", "1")),
            level=level,
        )
    for item in element:
        if item.tag.lower() != "li":
            continue
        paragraph = container.add_paragraph(style=_list_style_name(kind, level))
        _apply_paragraph_style(paragraph, item)
        if numbering_info is not None:
            _apply_paragraph_numbering(
                paragraph, num_id=numbering_info[0], level=numbering_info[1]
            )
        task_checked = None
        if is_task_list or item.get("data-type") == "taskItem":
            checked_value = item.get("data-checked")
            if checked_value is None:
                checkbox = next(
                    iter(item.xpath("./label/input[@type='checkbox']")), None
                )
                checked_value = (
                    "true"
                    if checkbox is not None and "checked" in checkbox.attrib
                    else "false"
                )
            task_checked = checked_value == "true"
        _append_list_item_contents(paragraph, item, budget, task_checked=task_checked)
        for nested in item:
            if nested.tag.lower() in {"ul", "ol"}:
                _add_html_list(container, nested, budget, level + 1)


def _clear_cell(cell: _Cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for child in list(cell._tc):
        if child is not tc_pr:
            cell._tc.remove(child)


def _set_cell_shading(cell: _Cell, color: str) -> None:
    cell_properties = cell._tc.get_or_add_tcPr()
    for shading in list(cell_properties.findall(qn("w:shd"))):
        cell_properties.remove(shading)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), color.lstrip("#"))
    cell_properties.append(shading)


def _set_cell_border(cell: _Cell, value: str) -> None:
    match = _BORDER_RE.fullmatch(value)
    if match is None:
        return
    width_pt = _length_to_points(match.group(1), minimum=0.125, maximum=12)
    color = _canonical_color(match.group(3))
    if color is None:
        return
    border_style = {
        "solid": "single",
        "dashed": "dashed",
        "dotted": "dotted",
        "double": "double",
    }[match.group(2).lower()]
    cell_properties = cell._tc.get_or_add_tcPr()
    for existing in list(cell_properties.findall(qn("w:tcBorders"))):
        cell_properties.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), border_style)
        border.set(qn("w:sz"), str(max(1, round(width_pt * 8))))
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color.lstrip("#"))
        borders.append(border)
    cell_properties.append(borders)


def _apply_table_cell_style(cell: _Cell, element: etree._Element) -> None:
    style = _style_map(element)
    if background := style.get("background-color"):
        _set_cell_shading(cell, background)
    if width := style.get("width"):
        cell.width = Pt(float(width[:-2]))
    if border := style.get("border"):
        _set_cell_border(cell, border)


def _apply_table_column_widths(table: Table, element: etree._Element) -> None:
    columns = element.xpath("./colgroup/col")
    if not columns:
        return
    table.autofit = False
    for index, column_element in enumerate(columns[: len(table.columns)]):
        width = _style_map(column_element).get("width")
        if width:
            table.columns[index].width = Pt(_length_to_points(width))


def _table_rows(element: etree._Element) -> list[etree._Element]:
    return list(element.xpath("./tr|./thead/tr|./tbody/tr|./tfoot/tr"))


def _row_cells(element: etree._Element) -> list[etree._Element]:
    return list(element.xpath("./th|./td"))


def _table_layout(element: etree._Element, budget: _ImageBudget | None = None):
    rows = _table_rows(element)
    if not rows:
        raise _fail("table must contain a row")
    if len(rows) > MAX_TABLE_ROWS:
        raise _fail("table row limit exceeded")

    occupied: set[tuple[int, int]] = set()
    placements = []
    max_column = 0
    expanded_cells = 0
    for row_index, row in enumerate(rows):
        if budget:
            budget.check()
        column_index = 0
        for cell in _row_cells(row):
            if budget:
                budget.check()
            while (row_index, column_index) in occupied:
                column_index += 1
            colspan = int(cell.get("colspan", "1"))
            rowspan = int(cell.get("rowspan", "1"))
            end_row = row_index + rowspan - 1
            end_column = column_index + colspan - 1
            if (
                end_row >= len(rows)
                or end_column >= MAX_TABLE_COLUMNS
                or colspan < 1
                or rowspan < 1
            ):
                raise _fail("table span outside supported grid")
            coordinates = {
                (occupied_row, occupied_column)
                for occupied_row in range(row_index, end_row + 1)
                for occupied_column in range(column_index, end_column + 1)
            }
            if coordinates & occupied:
                raise _fail("overlapping table span")
            occupied.update(coordinates)
            expanded_cells += len(coordinates)
            if expanded_cells > MAX_TOTAL_TABLE_CELLS:
                raise _fail("table cell limit exceeded")
            placements.append((row_index, column_index, end_row, end_column, cell))
            max_column = max(max_column, end_column + 1)
            column_index = end_column + 1
    if max_column <= 0 or max_column > MAX_TABLE_COLUMNS:
        raise _fail("table column limit exceeded")
    return rows, placements, max_column, expanded_cells


def _validate_html_tables(root: etree._Element, budget: _ImageBudget) -> None:
    total_cells = 0
    for table in root.xpath(".//table"):
        budget.check()
        _, _, _, expanded_cells = _table_layout(table, budget)
        total_cells += expanded_cells
        if total_cells > MAX_TOTAL_TABLE_CELLS:
            raise _fail("total table cell limit exceeded")


def _set_table_first_row_header(table: Table, enabled: bool) -> None:
    if not table._tbl.tr_lst:
        return
    row_properties = table._tbl.tr_lst[0].get_or_add_trPr()
    for header in list(row_properties.findall(qn("w:tblHeader"))):
        row_properties.remove(header)
    if enabled:
        row_properties.append(OxmlElement("w:tblHeader"))

    table_look = table._tbl.tblPr.find(qn("w:tblLook"))
    if table_look is None:
        table_look = OxmlElement("w:tblLook")
        table._tbl.tblPr.append(table_look)
    table_look.set(qn("w:firstRow"), "1" if enabled else "0")
    try:
        look_value = int(table_look.get(qn("w:val"), "0"), 16)
    except ValueError:
        look_value = 0
    look_value = look_value | 0x20 if enabled else look_value & ~0x20
    table_look.set(qn("w:val"), f"{look_value:04X}")


def _add_table(container, element: etree._Element, budget: _ImageBudget) -> None:
    rows, placements, column_count, _ = _table_layout(element, budget)
    if isinstance(container, _Cell):
        table = container.add_table(rows=len(rows), cols=column_count, width=Inches(6))
    else:
        table = container.add_table(rows=len(rows), cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    _apply_table_column_widths(table, element)
    first_row_cells = _row_cells(rows[0])
    _set_table_first_row_header(
        table,
        bool(first_row_cells)
        and all(cell.tag.lower() == "th" for cell in first_row_cells),
    )

    for row_index, column_index, end_row, end_column, cell_element in placements:
        target = table.cell(row_index, column_index)
        if end_row != row_index or end_column != column_index:
            target = target.merge(table.cell(end_row, end_column))
        _apply_table_cell_style(target, cell_element)
        _clear_cell(target)
        block_children = [
            child for child in cell_element if child.tag.lower() in _BLOCK_TAGS
        ]
        if block_children:
            _add_text_paragraph(target, cell_element.text or "", budget)
            for child in block_children:
                _append_html_block(target, child, budget)
        else:
            _add_paragraph_from_element(target, cell_element, budget)
        if not target.paragraphs:
            target.add_paragraph()


def _preformatted_text(element: etree._Element) -> str:
    parts = [element.text or ""]
    for child in element:
        if child.tag.lower() == "br":
            parts.append("\n")
        else:
            parts.append(_preformatted_text(child))
        parts.append(child.tail or "")
    return "".join(parts)


def _add_code_block(container, element: etree._Element) -> None:
    paragraph = container.add_paragraph(style=_CODE_BLOCK_STYLE)
    code = next(iter(element.xpath("./code")), element)
    run = paragraph.add_run(_preformatted_text(code))
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def _add_horizontal_rule(container) -> None:
    paragraph = container.add_paragraph(style=_HORIZONTAL_RULE_STYLE)
    paragraph_property = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    borders.append(bottom)
    paragraph_property.append(borders)


def _append_html_block(
    container, element: etree._Element, budget: _ImageBudget
) -> None:
    tag = element.tag.lower()
    if tag in {"ul", "ol"}:
        _add_html_list(container, element, budget)
        return
    if tag == "table":
        _add_table(container, element, budget)
        return
    if tag == "pre":
        _add_code_block(container, element)
        return
    if tag == "hr":
        _add_horizontal_rule(container)
        return
    if tag == "div" and any(child.tag.lower() in _BLOCK_TAGS for child in element):
        _add_text_paragraph(container, element.text or "", budget)
        for child in element:
            if child.tag.lower() in _BLOCK_TAGS:
                _append_html_block(container, child, budget)
            elif "".join(child.itertext()).strip():
                _add_paragraph_from_element(container, child, budget)
            _add_text_paragraph(container, child.tail or "", budget)
        return
    style_name = None
    if re.fullmatch(r"h[1-6]", tag):
        style_name = f"Heading {tag[1]}"
    _add_paragraph_from_element(container, element, budget, style_name=style_name)


def _clear_document_body(document) -> None:
    body = document._body._element
    section_properties = body.find(qn("w:sectPr"))
    for child in list(body):
        if child is not section_properties:
            body.remove(child)


def html_to_docx(html_content: str, base_document: bytes | None = None) -> bytes:
    budget = _ImageBudget()
    root = _parse_html_fragment(html_content, budget)
    if base_document is None:
        document = Document()
    else:
        validate_docx_bytes(base_document, _budget=budget)
        try:
            document = Document(io.BytesIO(base_document))
        except Exception:
            raise _fail("unsupported DOCX package") from None
        _clear_document_body(document)
    _ensure_conversion_styles(document)
    _add_text_paragraph(document, root.text or "", budget)
    for element in root:
        budget.check()
        if element.tag.lower() in _BLOCK_TAGS:
            _append_html_block(document, element, budget)
        else:
            _add_paragraph_from_element(document, element, budget)
        _add_text_paragraph(document, element.tail or "", budget)
    if not document.paragraphs and not document.tables:
        document.add_paragraph()
    output = io.BytesIO()
    budget.check(force=True)
    document.save(output)
    docx_bytes = output.getvalue()
    validate_docx_bytes(docx_bytes, _budget=budget)
    budget.check(force=True)
    return docx_bytes


def _style_chain(style) -> Iterable:
    seen: set[str] = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        yield style
        style = style.base_style


def _item_style_id(item: Run | Paragraph) -> str | None:
    if isinstance(item, Run):
        properties = item._r.rPr
        style_element = properties.rStyle if properties is not None else None
    else:
        properties = item._p.pPr
        style_element = properties.pStyle if properties is not None else None
    return str(style_element.val) if style_element is not None else None


def _cached_item_style(item: Run | Paragraph, kind: str, budget: _ImageBudget):
    key = (kind, _item_style_id(item))
    if key not in budget.style_cache:
        try:
            budget.style_cache[key] = item.style
        except (AttributeError, KeyError, ValueError):
            budget.style_cache[key] = None
    return budget.style_cache[key]


def _cached_named_style(document, name: str, budget: _ImageBudget):
    key = ("named", name)
    if key not in budget.style_cache:
        try:
            budget.style_cache[key] = document.styles[name]
        except (AttributeError, KeyError, ValueError):
            budget.style_cache[key] = None
    return budget.style_cache[key]


def _cached_style_chain(style, budget: _ImageBudget) -> tuple:
    if style is None:
        return ()
    try:
        key = (style.type, style.style_id)
    except (AttributeError, ValueError):
        return ()
    if key not in budget.style_chain_cache:
        try:
            budget.style_chain_cache[key] = tuple(_style_chain(style))
        except (AttributeError, ValueError):
            budget.style_chain_cache[key] = (style,)
    return budget.style_chain_cache[key]


def _read_font_properties(font) -> _FontProperties:
    def read(name: str):
        try:
            return getattr(font, name)
        except (AttributeError, ValueError):
            return None

    color = None
    highlight = None
    try:
        rgb = font.color.rgb
        if rgb is not None:
            color = f"#{str(rgb).upper()}"
    except (AttributeError, ValueError):
        pass
    try:
        highlight_value = font.highlight_color
        if highlight_value is not None:
            highlight = _HIGHLIGHT_TO_HEX.get(highlight_value)
    except (AttributeError, ValueError):
        pass
    return _FontProperties(
        bold=read("bold"),
        italic=read("italic"),
        underline=read("underline"),
        strike=read("strike"),
        superscript=read("superscript"),
        subscript=read("subscript"),
        font_name=read("name"),
        font_size=read("size"),
        color=color,
        highlight=highlight,
    )


def _merge_font_properties(*sources: _FontProperties) -> _FontProperties:
    def first(name: str):
        return next(
            (
                value
                for source in sources
                if (value := getattr(source, name)) is not None
            ),
            None,
        )

    return _FontProperties(
        bold=first("bold"),
        italic=first("italic"),
        underline=first("underline"),
        strike=first("strike"),
        superscript=first("superscript"),
        subscript=first("subscript"),
        font_name=first("font_name"),
        font_size=first("font_size"),
        color=first("color"),
        highlight=first("highlight"),
    )


def _style_font_properties(style, budget: _ImageBudget) -> _FontProperties:
    if style is None:
        return _FontProperties()
    try:
        key = (style.type, style.style_id)
    except (AttributeError, ValueError):
        return _FontProperties()
    if key not in budget.style_font_cache:
        budget.style_font_cache[key] = _merge_font_properties(
            *(
                _read_font_properties(item.font)
                for item in _cached_style_chain(style, budget)
            )
        )
    return budget.style_font_cache[key]


def _paragraph_font_properties(
    paragraph: Paragraph, budget: _ImageBudget
) -> _FontProperties:
    key = id(paragraph._p)
    if key not in budget.paragraph_font_cache:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        normal_style = _cached_named_style(paragraph.part.document, "Normal", budget)
        budget.paragraph_font_cache[key] = _merge_font_properties(
            _style_font_properties(paragraph_style, budget),
            _style_font_properties(normal_style, budget),
        )
    return budget.paragraph_font_cache[key]


def _run_inline_style(
    run: Run, paragraph: Paragraph, budget: _ImageBudget
) -> _InlineStyle:
    budget.check()
    run_style = _cached_item_style(run, "character", budget)
    properties = _merge_font_properties(
        _read_font_properties(run.font),
        _style_font_properties(run_style, budget),
        _paragraph_font_properties(paragraph, budget),
    )
    r_pr = run._r.rPr
    font_name = properties.font_name
    highlight = properties.highlight
    if r_pr is not None and r_pr.rFonts is not None:
        for attribute in ("w:eastAsia", "w:hAnsi", "w:ascii"):
            value = r_pr.rFonts.get(qn(attribute))
            if value:
                font_name = value
                break
    if r_pr is not None:
        shading = r_pr.find(qn("w:shd"))
        if shading is not None:
            fill = (shading.get(qn("w:fill")) or "").upper()
            if re.fullmatch(r"[0-9A-F]{6}", fill):
                highlight = f"#{fill}"
    try:
        run_style_name = run_style.name if run_style is not None else ""
        run_style_id = run_style.style_id if run_style is not None else ""
    except (AttributeError, KeyError, ValueError):
        run_style_name = run_style_id = ""
    size = properties.font_size
    return _InlineStyle(
        bold=bool(properties.bold),
        italic=bool(properties.italic),
        underline=bool(properties.underline),
        strike=bool(properties.strike),
        superscript=bool(properties.superscript),
        subscript=bool(properties.subscript),
        font_name=font_name,
        font_size_pt=size.pt if size is not None else None,
        color=properties.color,
        highlight=highlight,
        code=_INLINE_CODE_STYLE in {run_style_name, run_style_id},
    )


def _set_html_text(element: etree._Element, text: str) -> None:
    parts = text.split("\n")
    element.text = parts[0]
    for part in parts[1:]:
        line_break = etree.SubElement(element, "br")
        line_break.tail = part


def _iter_docx_image_nodes(element: etree._Element) -> Iterable[etree._Element]:
    if element.tag == _MC_ALTERNATE_CONTENT_TAG:
        choices = [child for child in element if child.tag == _MC_CHOICE_TAG]
        fallback = next(
            (child for child in element if child.tag == _MC_FALLBACK_TAG), None
        )
        for branch in (*choices, fallback):
            if branch is None:
                continue
            images = tuple(_iter_docx_image_nodes(branch))
            if images:
                yield from images
                return
        return
    if element.tag in {_DRAWINGML_BLIP_TAG, _VML_IMAGEDATA_TAG}:
        yield element
        return
    for child in element:
        yield from _iter_docx_image_nodes(child)


def _docx_image_container(
    image_node: etree._Element, scope: etree._Element
) -> etree._Element:
    current = image_node
    while current is not scope:
        parent = current.getparent()
        if parent is None:
            break
        current = parent
        if current.tag in _IMAGE_CONTAINER_TAGS:
            return current
    return scope


def _append_docx_image(
    parent: etree._Element,
    part,
    image_node: etree._Element,
    scope: etree._Element,
    budget: _ImageBudget,
) -> None:
    if image_node.tag == _VML_IMAGEDATA_TAG:
        linked_relationship_id = None
        relationship_id = image_node.get(qn("r:id"))
    else:
        linked_relationship_id = image_node.get(qn("r:link"))
        relationship_id = linked_relationship_id or image_node.get(qn("r:embed"))
    if not relationship_id:
        return
    image = etree.SubElement(parent, "img")
    if linked_relationship_id:
        try:
            relationship = part.rels[linked_relationship_id]
            if relationship.reltype != RT.IMAGE or not relationship.is_external:
                raise _fail("invalid external image relationship")
            image.set("src", _safe_network_url(relationship.target_ref))
        except (AttributeError, KeyError):
            raise _fail("invalid external image relationship") from None
    else:
        try:
            image_part = part.related_parts[relationship_id]
            image_bytes = image_part.blob
            content_type = image_part.content_type.lower().replace(
                "image/jpg", "image/jpeg"
            )
        except (AttributeError, KeyError):
            raise _fail("invalid embedded image relationship") from None
        data_uri = f"data:{content_type};base64,{base64.b64encode(image_bytes).decode('ascii')}"
        normalized_bytes, normalized_mime = _decode_image_data_uri(data_uri, budget)
        image.set(
            "src",
            f"data:{normalized_mime};base64,{base64.b64encode(normalized_bytes).decode('ascii')}",
        )
    container = _docx_image_container(image_node, scope)
    extent = next(container.iter(qn("wp:extent")), None)
    if extent is not None:
        try:
            width_pt = int(extent.get("cx")) / 12700
            height_pt = int(extent.get("cy")) / 12700
            if 1 <= width_pt <= 2000:
                image.set("width", f"{width_pt * 4 / 3:.17g}")
            else:
                logger.warning("Dropping invalid image width=%rpt", width_pt)
            if 1 <= height_pt <= 2000:
                image.set("height", f"{height_pt * 4 / 3:.17g}")
            else:
                logger.warning("Dropping invalid image height=%rpt", height_pt)
            image.set("style", f"width: {width_pt:g}pt; height: {height_pt:g}pt")
        except (TypeError, ValueError):
            pass
    properties = next(container.iter(qn("wp:docPr")), None)
    if properties is not None:
        alt = properties.get("descr") or properties.get("title")
        if alt:
            image.set("alt", alt[:1024])


def _append_docx_part_images(
    parent: etree._Element,
    part,
    element: etree._Element,
    budget: _ImageBudget,
    seen_image_nodes: set[etree._Element] | None = None,
) -> None:
    for image_node in _iter_docx_image_nodes(element):
        if seen_image_nodes is not None:
            if image_node in seen_image_nodes:
                continue
            seen_image_nodes.add(image_node)
        budget.check()
        _append_docx_image(parent, part, image_node, element, budget)


def _append_docx_run(
    parent: etree._Element,
    run: Run,
    paragraph: Paragraph,
    budget: _ImageBudget,
    seen_image_nodes: set[etree._Element] | None = None,
) -> None:
    budget.check()
    if run.text:
        style = _run_inline_style(run, paragraph, budget)
        outer = etree.SubElement(parent, "span")
        current = outer
        for enabled, tag in (
            (style.bold, "strong"),
            (style.italic, "em"),
            (style.underline, "u"),
            (style.strike, "s"),
            (style.superscript, "sup"),
            (style.subscript and not style.superscript, "sub"),
        ):
            if enabled:
                current = etree.SubElement(current, tag)
        if style.code:
            current = etree.SubElement(current, "code")
        if style.highlight:
            current = etree.SubElement(current, "mark")
            current.set("style", f"background-color: {style.highlight}")
        span_styles = []
        if style.font_name:
            span_styles.append(f"font-family: {style.font_name}")
        if style.font_size_pt:
            span_styles.append(f"font-size: {style.font_size_pt:g}pt")
        if style.color:
            span_styles.append(f"color: {style.color}")
        if span_styles:
            current = etree.SubElement(current, "span")
            current.set("style", "; ".join(span_styles))
        _set_html_text(current, run.text)
    _append_docx_part_images(
        parent, run.part, run._r, budget, seen_image_nodes=seen_image_nodes
    )


def _paragraph_alignment(paragraph: Paragraph, budget: _ImageBudget):
    if paragraph.alignment is not None:
        return paragraph.alignment
    try:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        for style in _cached_style_chain(paragraph_style, budget):
            alignment = style.paragraph_format.alignment
            if alignment is not None:
                return alignment
    except (AttributeError, KeyError, ValueError):
        pass
    return None


def _paragraph_format_value(paragraph: Paragraph, name: str, budget: _ImageBudget):
    try:
        value = getattr(paragraph.paragraph_format, name)
        if value is not None:
            return value
    except (AttributeError, ValueError):
        pass
    try:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        for style in _cached_style_chain(paragraph_style, budget):
            value = getattr(style.paragraph_format, name)
            if value is not None:
                return value
    except (AttributeError, KeyError, ValueError):
        pass
    return None


def _paragraph_css(paragraph: Paragraph, budget: _ImageBudget) -> str:
    declarations = []
    alignment = _paragraph_alignment(paragraph, budget)
    alignment_map = {
        WD_ALIGN_PARAGRAPH.LEFT: "left",
        WD_ALIGN_PARAGRAPH.CENTER: "center",
        WD_ALIGN_PARAGRAPH.RIGHT: "right",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "justify",
    }
    if alignment in alignment_map:
        declarations.append(f"text-align: {alignment_map[alignment]}")
    spacing = _paragraph_format_value(paragraph, "line_spacing", budget)
    if spacing is not None:
        if hasattr(spacing, "pt"):
            declarations.append(f"line-height: {spacing.pt:g}pt")
        elif isinstance(spacing, (float, int)):
            declarations.append(f"line-height: {float(spacing):g}")
    for format_name, css_name in (
        ("left_indent", "margin-left"),
        ("first_line_indent", "text-indent"),
        ("space_before", "margin-top"),
        ("space_after", "margin-bottom"),
    ):
        value = _paragraph_format_value(paragraph, format_name, budget)
        if value is not None and hasattr(value, "pt"):
            declarations.append(f"{css_name}: {value.pt:g}pt")
    return "; ".join(declarations)


def _heading_level(paragraph: Paragraph, budget: _ImageBudget) -> int | None:
    try:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        style_name = paragraph_style.name or ""
        style_id = paragraph_style.style_id or ""
    except (AttributeError, KeyError, ValueError):
        return None
    match = re.fullmatch(
        r"Heading\s*([1-6])", style_name, re.IGNORECASE
    ) or re.fullmatch(r"Heading([1-6])", style_id, re.IGNORECASE)
    return int(match.group(1)) if match else None


def _paragraph_has_style(
    paragraph: Paragraph, expected: str, budget: _ImageBudget
) -> bool:
    try:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        return expected in {paragraph_style.name, paragraph_style.style_id}
    except (AttributeError, KeyError, ValueError):
        return False


def _is_horizontal_rule(paragraph: Paragraph, budget: _ImageBudget) -> bool:
    if _paragraph_has_style(paragraph, _HORIZONTAL_RULE_STYLE, budget):
        return True
    if paragraph.text.strip():
        return False
    paragraph_property = paragraph._p.pPr
    if paragraph_property is None:
        return False
    borders = paragraph_property.find(qn("w:pBdr"))
    return borders is not None and borders.find(qn("w:bottom")) is not None


def _style_list_info(paragraph: Paragraph, budget: _ImageBudget) -> _ListInfo | None:
    try:
        paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
        style_name = paragraph_style.name or ""
        style_id = paragraph_style.style_id or ""
    except (AttributeError, KeyError, ValueError):
        return None
    match = re.fullmatch(
        r"List\s+(Bullet|Number)(?:\s+([1-9]\d*))?", style_name, re.IGNORECASE
    ) or re.fullmatch(r"List(Bullet|Number)([1-9]\d*)?", style_id, re.IGNORECASE)
    if not match:
        return None
    return _ListInfo(
        kind="ul" if match.group(1).lower() == "bullet" else "ol",
        level=max(0, int(match.group(2) or "1") - 1),
    )


def _numbering_kind(paragraph: Paragraph, num_id: int, level: int) -> _ListInfo | None:
    try:
        numbering = paragraph.part.numbering_part.element
        number = next(
            node
            for node in numbering.xpath("./w:num")
            if int(node.get(qn("w:numId"), "-1")) == num_id
        )
        abstract_id_element = number.find(qn("w:abstractNumId"))
        abstract_id = int(abstract_id_element.get(qn("w:val")))
        abstract = next(
            node
            for node in numbering.xpath("./w:abstractNum")
            if int(node.get(qn("w:abstractNumId"), "-1")) == abstract_id
        )
        levels = abstract.findall(qn("w:lvl"))
        level_element = next(
            (node for node in levels if int(node.get(qn("w:ilvl"), "0")) == level),
            levels[0] if levels else None,
        )
        if level_element is None:
            return None
        override = next(
            (
                node
                for node in number.findall(qn("w:lvlOverride"))
                if int(node.get(qn("w:ilvl"), "0")) == level
            ),
            None,
        )
        override_level = override.find(qn("w:lvl")) if override is not None else None
        format_element = (
            override_level.find(qn("w:numFmt")) if override_level is not None else None
        )
        if format_element is None:
            format_element = level_element.find(qn("w:numFmt"))
        number_format = (
            format_element.get(qn("w:val"), "") if format_element is not None else ""
        )
        if number_format == "bullet":
            return _ListInfo(kind="ul", level=level, num_id=num_id)

        start_element = (
            override.find(qn("w:startOverride")) if override is not None else None
        )
        if start_element is None and override_level is not None:
            start_element = override_level.find(qn("w:start"))
        if start_element is None:
            start_element = level_element.find(qn("w:start"))
        start = (
            int(start_element.get(qn("w:val"), "1")) if start_element is not None else 1
        )
        if not 1 <= start <= 100_000:
            logger.warning(
                "Dropping invalid ordered-list start=%r for numId=%s level=%s",
                start,
                num_id,
                level,
            )
            start = 1
        ordered_type = {
            "lowerLetter": "a",
            "upperLetter": "A",
            "lowerRoman": "i",
            "upperRoman": "I",
        }.get(number_format)
        return _ListInfo(
            kind="ol",
            level=level,
            num_id=num_id,
            ordered_type=ordered_type,
            start=start,
        )
    except (AttributeError, KeyError, StopIteration, TypeError, ValueError):
        return None


def _list_info(paragraph: Paragraph, budget: _ImageBudget) -> _ListInfo | None:
    style_info = _style_list_info(paragraph, budget)
    num_pr = paragraph._p.pPr.numPr if paragraph._p.pPr is not None else None
    if num_pr is None:
        try:
            paragraph_style = _cached_item_style(paragraph, "paragraph", budget)
            for style in _cached_style_chain(paragraph_style, budget):
                style_p_pr = style._element.pPr
                if style_p_pr is not None and style_p_pr.numPr is not None:
                    num_pr = style_p_pr.numPr
                    break
        except (AttributeError, KeyError, ValueError):
            pass
    if num_pr is None or num_pr.numId is None:
        return style_info
    try:
        num_id = int(num_pr.numId.val)
        level = (
            int(num_pr.ilvl.val)
            if num_pr.ilvl is not None
            else (style_info.level if style_info else 0)
        )
    except (TypeError, ValueError):
        return style_info
    return _numbering_kind(paragraph, num_id, level) or style_info


def _append_paragraph_contents(
    parent: etree._Element, paragraph: Paragraph, budget: _ImageBudget
) -> None:
    seen_image_nodes: set[etree._Element] = set()
    content_by_element = {
        item._element: item for item in paragraph.iter_inner_content()
    }
    for child in paragraph._p:
        item = content_by_element.get(child)
        if item is None:
            _append_docx_part_images(
                parent,
                paragraph.part,
                child,
                budget,
                seen_image_nodes=seen_image_nodes,
            )
            continue
        if isinstance(item, Run):
            _append_docx_run(
                parent,
                item,
                paragraph,
                budget,
                seen_image_nodes=seen_image_nodes,
            )
        elif isinstance(item, Hyperlink):
            href = item.url or (f"#{item.fragment}" if item.fragment else "")
            try:
                href = _safe_href(href)
            except DocxConversionError:
                href = ""
            hyperlink_parent = (
                etree.SubElement(parent, "a", href=href) if href else parent
            )
            for run in item.runs:
                _append_docx_run(
                    hyperlink_parent,
                    run,
                    paragraph,
                    budget,
                    seen_image_nodes=seen_image_nodes,
                )


def _append_docx_paragraph(
    parent: etree._Element, paragraph: Paragraph, budget: _ImageBudget
) -> etree._Element:
    if _paragraph_has_style(paragraph, _CODE_BLOCK_STYLE, budget):
        pre = etree.SubElement(parent, "pre")
        code = etree.SubElement(pre, "code")
        code.text = paragraph.text
        return pre
    if _is_horizontal_rule(paragraph, budget):
        return etree.SubElement(parent, "hr")
    heading_level = _heading_level(paragraph, budget)
    element = etree.SubElement(parent, f"h{heading_level}" if heading_level else "p")
    css = _paragraph_css(paragraph, budget)
    if css:
        element.set("style", css)
    _append_paragraph_contents(element, paragraph, budget)
    return element


def _grid_span(cell: _Cell) -> int:
    tc_pr = cell._tc.tcPr
    if tc_pr is None or tc_pr.gridSpan is None:
        return 1
    try:
        return max(1, int(tc_pr.gridSpan.val))
    except (TypeError, ValueError):
        return 1


def _row_grid_before(row) -> int:
    row_properties = row.trPr
    grid_before = (
        row_properties.find(qn("w:gridBefore")) if row_properties is not None else None
    )
    try:
        return (
            max(0, int(grid_before.get(qn("w:val")))) if grid_before is not None else 0
        )
    except (TypeError, ValueError):
        return 0


def _raw_row_cells(table: Table, row_index: int):
    row = table._tbl.tr_lst[row_index]
    column_index = _row_grid_before(row)
    for cell_xml in row.tc_lst:
        cell = _Cell(cell_xml, table)
        colspan = _grid_span(cell)
        yield column_index, colspan, cell
        column_index += colspan


def _vertical_merge_state(cell: _Cell) -> str | None:
    cell_properties = cell._tc.tcPr
    merge = (
        cell_properties.find(qn("w:vMerge")) if cell_properties is not None else None
    )
    if merge is None:
        return None
    return merge.get(qn("w:val")) or "continue"


def _vertical_span(
    table: Table, row_index: int, column_index: int, colspan: int, cell: _Cell
) -> int:
    if _vertical_merge_state(cell) != "restart":
        return 1
    rowspan = 1
    for following_row in range(row_index + 1, len(table._tbl.tr_lst)):
        continuation = next(
            (
                candidate
                for start, candidate_span, candidate in _raw_row_cells(
                    table, following_row
                )
                if start == column_index and candidate_span == colspan
            ),
            None,
        )
        if continuation is None or _vertical_merge_state(continuation) != "continue":
            break
        rowspan += 1
    return rowspan


def _on_off_enabled(element: etree._Element) -> bool:
    value = element.get(qn("w:val"))
    return value is None or value.lower() not in {"0", "false", "off", "no"}


def _first_row_is_header(table: Table) -> bool:
    if not table._tbl.tr_lst:
        return False
    row_properties = table._tbl.tr_lst[0].trPr
    row_header = (
        row_properties.find(qn("w:tblHeader")) if row_properties is not None else None
    )
    if row_header is not None and _on_off_enabled(row_header):
        return True

    table_properties = table._tbl.tblPr
    table_look = table_properties.find(qn("w:tblLook"))
    if table_look is None:
        return False
    first_row = table_look.get(qn("w:firstRow"))
    if first_row is not None:
        return first_row.lower() not in {"0", "false", "off", "no"}
    try:
        return bool(int(table_look.get(qn("w:val"), "0"), 16) & 0x20)
    except ValueError:
        return False


def _word_color(value: str | None) -> str | None:
    value = (value or "").upper()
    if value == "AUTO":
        return "#000000"
    return f"#{value}" if re.fullmatch(r"[0-9A-F]{6}", value) else None


def _cell_width_points(
    table: Table, cell: _Cell, column_index: int, colspan: int
) -> float | None:
    cell_properties = cell._tc.tcPr
    width = cell_properties.find(qn("w:tcW")) if cell_properties is not None else None
    if width is not None and (width.get(qn("w:type")) or "dxa") == "dxa":
        try:
            twips = int(width.get(qn("w:w"), "0"))
            if twips > 0:
                return twips / 20
        except (TypeError, ValueError):
            pass
    grid_columns = (
        table._tbl.tblGrid.gridCol_lst if table._tbl.tblGrid is not None else []
    )
    if column_index + colspan > len(grid_columns):
        return None
    try:
        twips = sum(
            int(column.get(qn("w:w"), "0"))
            for column in grid_columns[column_index : column_index + colspan]
        )
    except (TypeError, ValueError):
        return None
    return twips / 20 if twips > 0 else None


def _cell_border_css(cell: _Cell) -> str | None:
    cell_properties = cell._tc.tcPr
    borders = (
        cell_properties.find(qn("w:tcBorders")) if cell_properties is not None else None
    )
    if borders is None:
        return None
    signatures = []
    for edge_name in ("top", "right", "bottom", "left"):
        edge = borders.find(qn(f"w:{edge_name}"))
        if edge is None:
            continue
        word_style = edge.get(qn("w:val"), "single")
        css_style = {
            "single": "solid",
            "dashed": "dashed",
            "dashSmallGap": "dashed",
            "dotted": "dotted",
            "double": "double",
        }.get(word_style)
        color = _word_color(edge.get(qn("w:color")))
        if css_style is None or color is None:
            continue
        try:
            width_pt = max(0.125, int(edge.get(qn("w:sz"), "8")) / 8)
        except (TypeError, ValueError):
            continue
        signatures.append((width_pt, css_style, color))
    if not signatures or any(item != signatures[0] for item in signatures[1:]):
        return None
    width_pt, css_style, color = signatures[0]
    return f"{width_pt:g}pt {css_style} {color}"


def _table_cell_css(table: Table, cell: _Cell, column_index: int, colspan: int) -> str:
    declarations = []
    cell_properties = cell._tc.tcPr
    shading = cell_properties.find(qn("w:shd")) if cell_properties is not None else None
    if shading is not None:
        fill = _word_color(shading.get(qn("w:fill")))
        pattern = shading.get(qn("w:val"), "clear")
        if fill is not None and pattern not in {"nil", "none"}:
            declarations.append(f"background-color: {fill}")
    width_pt = _cell_width_points(table, cell, column_index, colspan)
    if width_pt is not None:
        declarations.append(f"width: {width_pt:g}pt")
    border = _cell_border_css(cell)
    if border is not None:
        declarations.append(f"border: {border}")
    return "; ".join(declarations)


def _append_table_grid_columns(table_element: etree._Element, table: Table) -> None:
    grid = table._tbl.tblGrid
    if grid is None or not grid.gridCol_lst:
        return
    widths = []
    for column in grid.gridCol_lst:
        try:
            width_pt = int(column.get(qn("w:w"), "0")) / 20
        except (TypeError, ValueError):
            return
        if not 1 <= width_pt <= 2000:
            return
        widths.append(width_pt)
    column_group = etree.SubElement(table_element, "colgroup")
    for width_pt in widths:
        column = etree.SubElement(column_group, "col")
        column.set("style", f"width: {width_pt:g}pt")


def _append_docx_table(
    parent: etree._Element, table: Table, budget: _ImageBudget
) -> None:
    table_element = etree.SubElement(parent, "table")
    _append_table_grid_columns(table_element, table)
    body = etree.SubElement(table_element, "tbody")
    first_row_is_header = _first_row_is_header(table)
    for row_index, _row in enumerate(table._tbl.tr_lst):
        budget.check()
        row_element = etree.SubElement(body, "tr")
        cell_tag = "th" if row_index == 0 and first_row_is_header else "td"
        for column_index, colspan, cell in _raw_row_cells(table, row_index):
            budget.check()
            if _vertical_merge_state(cell) == "continue":
                continue
            cell_element = etree.SubElement(row_element, cell_tag)
            if colspan > 1:
                cell_element.set("colspan", str(colspan))
            rowspan = _vertical_span(table, row_index, column_index, colspan, cell)
            if rowspan > 1:
                cell_element.set("rowspan", str(rowspan))
            css = _table_cell_css(table, cell, column_index, colspan)
            if css:
                cell_element.set("style", css)
            _append_docx_blocks(cell_element, cell.iter_inner_content(), budget)


def _append_list_paragraph(
    parent: etree._Element,
    paragraph: Paragraph,
    info: _ListInfo,
    stack: list[_ListContext],
    budget: _ImageBudget,
    task_checked: bool | None = None,
) -> None:
    kind = info.kind
    level = info.level
    if not stack:
        level = 0
    elif level > stack[-1].level + 1:
        level = stack[-1].level + 1
    while stack and (
        level < stack[-1].level
        or (
            level == stack[-1].level
            and (
                kind != stack[-1].kind
                or info.num_id != stack[-1].num_id
                or info.ordered_type != stack[-1].ordered_type
            )
        )
    ):
        stack.pop()
    if not stack or level > stack[-1].level:
        list_parent = (
            stack[-1].last_item if stack and stack[-1].last_item is not None else parent
        )
        list_element = etree.SubElement(list_parent, "ul" if kind == "task" else kind)
        if kind == "task":
            list_element.set("data-type", "taskList")
        elif kind == "ol":
            if info.ordered_type:
                list_element.set("type", info.ordered_type)
            if info.start is not None and info.start != 1:
                list_element.set("start", str(info.start))
        stack.append(
            _ListContext(
                level=level,
                kind=kind,
                element=list_element,
                num_id=info.num_id,
                ordered_type=info.ordered_type,
            )
        )
    context = stack[-1]
    item = etree.SubElement(context.element, "li")
    content_parent = item
    if kind == "task":
        item.set("data-type", "taskItem")
        item.set("data-checked", "true" if task_checked else "false")
        content_parent = etree.SubElement(etree.SubElement(item, "div"), "p")
    css = _paragraph_css(paragraph, budget)
    if css:
        content_parent.set("style", css)
    _append_paragraph_contents(content_parent, paragraph, budget)
    if kind == "task":
        _strip_task_prefix(content_parent)
    context.last_item = item


def _task_checked(paragraph: Paragraph) -> bool | None:
    text = paragraph.text.lstrip()
    if text.startswith(_TASK_CHECKED):
        return True
    if text.startswith(_TASK_UNCHECKED):
        return False
    return None


def _strip_task_prefix(element: etree._Element) -> None:
    for text_node in element.xpath(".//text()"):
        value = str(text_node)
        if not value.strip():
            continue
        match = re.match(rf"^(\s*)[{_TASK_CHECKED}{_TASK_UNCHECKED}]\s?", value)
        if not match:
            return
        parent = text_node.getparent()
        replacement = value[: len(match.group(1))] + value[match.end() :]
        if text_node.is_text:
            parent.text = replacement
        elif text_node.is_tail:
            parent.tail = replacement
        return


def _append_docx_blocks(
    parent: etree._Element,
    blocks: Iterable[Paragraph | Table],
    budget: _ImageBudget,
) -> None:
    stack: list[_ListContext] = []
    for block in blocks:
        budget.check()
        if isinstance(block, Paragraph):
            if _heading_level(block, budget) is not None:
                stack.clear()
                _append_docx_paragraph(parent, block, budget)
                continue
            info = _list_info(block, budget)
            task_checked = _task_checked(block)
            if task_checked is not None:
                info = _ListInfo(
                    kind="task", level=info.level if info is not None else 0
                )
            if info:
                _append_list_paragraph(
                    parent,
                    block,
                    info,
                    stack,
                    budget,
                    task_checked=task_checked,
                )
            else:
                stack.clear()
                _append_docx_paragraph(parent, block, budget)
        elif isinstance(block, Table):
            stack.clear()
            _append_docx_table(parent, block, budget)


def docx_to_html(
    docx_bytes: bytes,
    *,
    _budget: _ImageBudget | None = None,
    _validated: bool = False,
) -> str:
    budget = _budget or _ImageBudget()
    if not _validated:
        validate_docx_bytes(docx_bytes, _budget=budget)
    try:
        document = Document(io.BytesIO(docx_bytes))
        root = etree.Element("div")
        _append_docx_blocks(root, document.iter_inner_content(), budget)
        html_content = (
            "".join(
                etree.tostring(
                    child, encoding="unicode", method="html", with_tail=False
                )
                for child in root
            )
            or "<p></p>"
        )
        sanitized_root = _parse_html_fragment(html_content, budget)
        html_content = "".join(
            etree.tostring(child, encoding="unicode", method="html", with_tail=False)
            for child in sanitized_root
        )
    except DocxConversionError:
        raise
    except Exception:
        raise _fail("unsupported DOCX content") from None
    if len(html_content.encode("utf-8")) > MAX_HTML_BYTES:
        raise _fail("converted HTML exceeds size limit")
    budget.check(force=True)
    return html_content or "<p></p>"


def _paragraph_text_runs(paragraph: Paragraph) -> list[Run]:
    runs: list[Run] = []
    for item in paragraph.iter_inner_content():
        if isinstance(item, Run):
            runs.append(item)
        elif isinstance(item, Hyperlink):
            runs.extend(item.runs)
    return runs


def _trim_paragraph_spaces(paragraph: Paragraph) -> None:
    runs = [run for run in _paragraph_text_runs(paragraph) if run.text]
    if not runs:
        return
    runs[0].text = runs[0].text.lstrip(" \t")
    runs[-1].text = runs[-1].text.rstrip(" \t")


def replace_paragraph_text_matches(
    paragraph: Paragraph,
    pattern: re.Pattern,
    replacer: Callable[[re.Match], str | None],
) -> bool:
    """Replace matches split across runs without collapsing surrounding run styles."""
    runs = _paragraph_text_runs(paragraph)
    if not runs:
        return False
    original_texts = [run.text for run in runs]
    full_text = "".join(original_texts)
    matches = list(pattern.finditer(full_text))
    if not matches:
        return False

    starts: list[int] = []
    cursor = 0
    for text in original_texts:
        starts.append(cursor)
        cursor += len(text)

    def locate(offset: int) -> tuple[int, int]:
        for index, (start, text) in enumerate(zip(starts, original_texts)):
            if start <= offset < start + len(text):
                return index, offset - start
        raise _fail("placeholder run mapping failed")

    for match in reversed(matches):
        if match.start() == match.end():
            continue
        replacement = replacer(match)
        if replacement is None:
            continue
        start_index, start_offset = locate(match.start())
        end_index, end_last_offset = locate(match.end() - 1)
        end_offset = end_last_offset + 1
        replacement_text = str(replacement)
        if start_index == end_index:
            current = runs[start_index].text
            runs[start_index].text = (
                current[:start_offset] + replacement_text + current[end_offset:]
            )
            continue
        start_text = runs[start_index].text
        end_text = runs[end_index].text
        runs[start_index].text = start_text[:start_offset] + replacement_text
        for index in range(start_index + 1, end_index):
            runs[index].text = ""
        runs[end_index].text = end_text[end_offset:]
    return True
